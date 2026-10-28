# -*- coding: utf-8 -*-
"""Çekirdek mantık ve güvenlik testleri (arayüz gerektirmez; gerçek veri kullanmaz).

Çalıştırma (proje klasöründe):   python -m unittest discover -s tests -v
"""
import os
import sqlite3
import sys
import tempfile
import unittest
import zlib

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import fatura_takip as f

_TUR = f._PBKDF2_TUR


def setUpModule():
    f._PBKDF2_TUR = 1000            # testler hızlı olsun; gerçek tur sayısı üretimde 600.000


def tearDownModule():
    f._PBKDF2_TUR = _TUR


def yeni_db():
    klasor = tempfile.mkdtemp()
    return f.FaturaDB(os.path.join(klasor, "t.db")), klasor


def fatura(i, **k):
    v = dict(id=i, fatura_tarihi="2026-10-%02d" % (1 + i % 28), fatura_no="F-%d" % i, firma_adi="ALE YÜK A.Ş.",
             yapilan_is="Vinç", tutar_haric=100000 + i, kategori="Genel Alım", talep_no="4035%03d" % i,
             talep_tarihi="2026-09-29", talep_eden_birim="Aliağa İletim Şebekeleri İşletme ve Bakım Müdürlüğü",
             dt_no="26DT500%04d" % i, pdf_yolu="", ek_sayisi=None)
    v.update(k)
    return v


class TutarVeTarih(unittest.TestCase):
    def test_tutar_bicimleri(self):
        for metin, kurus in (("1.234,56", 123456), ("1234,56", 123456), ("1234.56", 123456), ("1.234", 123400),
                             ("1.234,56 TL", 123456), ("0,5", 50), ("", None)):
            self.assertEqual(f.tutar_ayristir(metin), kurus, metin)
        for kotu in ("abc", "-5"):
            with self.assertRaises(ValueError):
                f.tutar_ayristir(kotu)

    def test_tarih(self):
        self.assertEqual(f.tarih_ayristir("04.06.2026"), "2026-06-04")
        with self.assertRaises(ValueError):
            f.tarih_ayristir("32.13.2026")


class Hesaplar(unittest.TestCase):
    def setUp(self):
        self.db, self.klasor = yeni_db()

    def test_parola_en_az_uzunluk(self):
        with self.assertRaises(f.VeritabaniHatasi):
            self.db.ilk_yonetici_olustur("Y", "12345")
        self.assertEqual(self.db.kullanici_sayisi(), 0)

    def test_parola_alt_siniri_6_karakter(self):
        self.assertEqual(f.SIFRE_EN_AZ, 6)
        self.db.ilk_yonetici_olustur("Y", "123456")                       # tam 6 karakter kabul edilir
        with self.assertRaises(f.VeritabaniHatasi):
            self.db.kullanici_ekle("Ayse", "12345")                       # 5 karakter reddedilir
        self.db.kullanici_ekle("Ayse", "abcdef", sifre_degismeli=False)
        a = self.db.giris_dene("ayse", "abcdef")[0]["id"]
        with self.assertRaises(f.VeritabaniHatasi):
            self.db.sifre_degistir(a, "abcde")
        self.db.sifre_degistir(a, "ghijkl")
        self.assertTrue(self.db.giris_dene("ayse", "ghijkl")[0])

    def test_ilk_yonetici_yalniz_bir_kez(self):
        self.db.ilk_yonetici_olustur("Yonetici", "gizli1234")
        with self.assertRaises(f.VeritabaniHatasi):
            self.db.ilk_yonetici_olustur("Baska", "gizli1234")
        self.assertEqual(self.db.kullanici_sayisi(), 1)

    def test_giris_ve_turkce_ad_eslemesi(self):
        self.db.ilk_yonetici_olustur("Yönetici", "gizli1234")
        self.assertTrue(self.db.giris_dene("YÖNETİCİ", "gizli1234")[0])
        self.assertIsNone(self.db.giris_dene("Yönetici", "gizli1235")[0])
        self.assertIsNone(self.db.giris_dene("yok", "gizli1234")[0])

    def test_sifre_duz_metin_saklanmaz(self):
        self.db.ilk_yonetici_olustur("Y", "gizli1234")
        kayit = sqlite3.connect(self.db.yol).execute("SELECT sifre FROM kullanicilar").fetchone()[0]
        self.assertTrue(kayit.startswith("pbkdf2$"))
        self.assertNotIn("gizli1234", kayit)

    def test_eski_hash_giriste_yenilenir(self):
        self.db.ilk_yonetici_olustur("Y", "gizli1234")
        eski = f._PBKDF2_TUR
        f._PBKDF2_TUR = 500                                  # "eski" daha az turlu hash üret
        eski_hash = f.sifre_hashle("gizli1234")
        f._PBKDF2_TUR = eski
        con = sqlite3.connect(self.db.yol); con.execute("UPDATE kullanicilar SET sifre=?", (eski_hash,)); con.commit(); con.close()
        self.assertTrue(self.db.giris_dene("y", "gizli1234")[0])
        yeni = sqlite3.connect(self.db.yol).execute("SELECT sifre FROM kullanicilar").fetchone()[0]
        self.assertEqual(int(yeni.split("$")[1]), eski)      # artık güncel tur sayısıyla
        self.assertTrue(self.db.giris_dene("y", "gizli1234")[0])

    def test_kullanici_adi_sinirlari(self):
        self.db.ilk_yonetici_olustur("Y", "gizli1234")
        for kotu in ("", "a" * 41, "ad\x00soyad", "ad\tsoyad"):
            with self.assertRaises(f.VeritabaniHatasi):
                self.db.kullanici_ekle(kotu, "gecici1234")
        with self.assertRaises(f.MukerrerKayit):
            self.db.kullanici_ekle("y", "gecici1234")

    def test_oturumlar_sifre_degisince_ve_pasifte_duser(self):
        self.db.ilk_yonetici_olustur("Y", "gizli1234")
        self.db.kullanici_ekle("Ayse", "gecici1234", sifre_degismeli=False)
        a = self.db.giris_dene("ayse", "gecici1234")[0]["id"]
        t1, t2 = self.db.oturum_olustur(a), self.db.oturum_olustur(a)
        self.assertTrue(self.db.oturum_dogrula(t1))
        self.db.sifre_degistir(a, "yeni-sifre-9", eski="gecici1234")
        self.assertIsNone(self.db.oturum_dogrula(t1))
        self.assertIsNone(self.db.oturum_dogrula(t2))
        t3 = self.db.oturum_olustur(a)
        self.db.sifre_sifirla(a, "sifir-1234")
        self.assertIsNone(self.db.oturum_dogrula(t3))
        t4 = self.db.oturum_olustur(a)
        self.db.kullanici_durumu(a, aktif=False)
        self.assertIsNone(self.db.oturum_dogrula(t4))
        self.assertIn("pasif", self.db.giris_dene("ayse", "sifir-1234")[1])

    def test_oturum_suresi_dolar(self):
        self.db.ilk_yonetici_olustur("Y", "gizli1234")
        kid = self.db.giris_dene("y", "gizli1234")[0]["id"]
        t = self.db.oturum_olustur(kid)
        con = sqlite3.connect(self.db.yol); con.execute("UPDATE oturumlar SET olusturma_tarihi='2000-01-01 00:00:00'"); con.commit(); con.close()
        self.assertIsNone(self.db.oturum_dogrula(t))

    def test_son_aktif_yonetici_korunur(self):
        self.db.ilk_yonetici_olustur("Y", "gizli1234")
        y = self.db.giris_dene("y", "gizli1234")[0]["id"]
        with self.assertRaises(f.VeritabaniHatasi):
            self.db.kullanici_durumu(y, aktif=False)
        with self.assertRaises(f.VeritabaniHatasi):
            self.db.kullanici_durumu(y, yonetici=False)
        self.assertEqual(self.db.giris_dene("y", "gizli1234")[0]["aktif"], 1)     # işlem geri alındı


class TemaHesabaOzel(unittest.TestCase):
    def test_tema_her_kullanicinin_kendi_hesabinda(self):
        db, _ = yeni_db()
        db.ilk_yonetici_olustur("Y", "gizli1234")
        db.kullanici_ekle("Ayse", "gecici1234", sifre_degismeli=False)
        y = db.giris_dene("y", "gizli1234")[0]
        a = db.giris_dene("ayse", "gecici1234")[0]
        self.assertIsNone(y["tema"])                                   # seçim yoksa varsayılan (config.ini) kullanılır
        db.kullanici_tema_yaz(a["id"], "Mor")
        self.assertEqual(db.giris_dene("ayse", "gecici1234")[0]["tema"], "Mor")
        self.assertIsNone(db.giris_dene("y", "gizli1234")[0]["tema"])   # başkasının seçimi diğerini etkilemez

    def test_eski_veritabanina_tema_sutunu_eklenir(self):
        yol = os.path.join(tempfile.mkdtemp(), "eski.db")
        con = sqlite3.connect(yol)
        con.execute("CREATE TABLE kullanicilar (id INTEGER PRIMARY KEY AUTOINCREMENT, kullanici_adi TEXT NOT NULL, "
                    "kullanici_anahtar TEXT NOT NULL UNIQUE, sifre TEXT NOT NULL, yonetici INTEGER NOT NULL DEFAULT 0, "
                    "aktif INTEGER NOT NULL DEFAULT 1, sifre_degismeli INTEGER NOT NULL DEFAULT 0, olusturma_tarihi TEXT)")
        con.commit(); con.close()
        f.FaturaDB(yol)                                                # açılışta şema yükseltilir
        sutunlar = [x[1] for x in sqlite3.connect(yol).execute("PRAGMA table_info(kullanicilar)")]
        self.assertIn("tema", sutunlar)


class Arama(unittest.TestCase):
    def test_fatura_ara(self):
        tum = [fatura(1, firma_adi="ALE YÜK A.Ş.", fatura_no="ALE-001", tutar_haric=3250000),
               fatura(2, firma_adi="Özer Doğrayan", fatura_no="GIB2026000000015", tutar_haric=2000000),
               fatura(3, firma_adi="ÇİMENTO LTD.", fatura_no="ÇİM 77", tutar_haric=None)]
        ids = lambda l: sorted(x["id"] for x in l)
        self.assertEqual(ids(f.fatura_ara(tum)), [1, 2, 3])
        self.assertEqual(ids(f.fatura_ara(tum, firma="ÖZER")), [2])
        self.assertEqual(ids(f.fatura_ara(tum, firma="çimento")), [3])
        self.assertEqual(ids(f.fatura_ara(tum, fatura_no="çim77")), [3])               # boşluk önemsiz
        self.assertEqual(ids(f.fatura_ara(tum, en_az=3250000, en_cok=3250000)), [1])
        self.assertEqual(ids(f.fatura_ara(tum, en_az=1)), [1, 2])                      # tutarsız kayıt sınırla çıkmaz
        self.assertEqual(ids(f.fatura_ara(tum, firma="ale", en_az=9999999)), [])


class PdfKurallari(unittest.TestCase):
    ust = "T.C.\nTÜRKİYE ELEKTRİK İLETİM ANONİM ŞİRKETİ GENEL MÜDÜRLÜĞÜ\n3.Bölge Müdürlüğü (İzmir)\nXYZ Birimi\nSayı : E-1-3570050\n"

    def test_birim_dort_belge_turu(self):
        for baslik in ("TALEP FORMU", "Malzeme Talep Formu", "Hizmet   İhtiyaç  Formu", "Malzeme İhtiyaç Formu"):
            self.assertEqual(f.birim_coz(self.ust + baslik), ("XYZ Birimi", ""), baslik)

    def test_birim_yanlis_belge_okunmaz(self):
        deger, uyari = f.birim_coz(self.ust + "Teklif Mektubu")
        self.assertIsNone(deger)
        self.assertTrue(uyari)

    def test_talep_no_tarih_sadece_talep_formundan(self):
        metin = "23.01.2026\nSayı : E-60994212-934.01.99-3570050\nMal Alımı Talep Formu"
        self.assertEqual(f.talep_coz(metin)[:2], ("3570050", "2026-01-23"))
        self.assertEqual(f.talep_coz(metin.replace("Mal Alımı Talep Formu", "Hizmet İhtiyaç Formu"))[:2], (None, None))

    def test_is_adinin_sonundaki_parantez(self):
        a = f.is_adi_parantez_at
        self.assertEqual(a("Sepetli Vinç Kiralama (1 Adet) (Gölmarmara)"), "Sepetli Vinç Kiralama")
        self.assertEqual(a("Vinç (Sepetli) Kiralama"), "Vinç (Sepetli) Kiralama")
        self.assertEqual(a("(Sadece parantez)"), "(Sadece parantez)")


def _docx(paragraflar):
    """Verilen paragraflardan geçici bir .docx üretir (her paragraf ayrı satır)."""
    import zipfile
    ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
    govde = "".join("<w:p><w:r><w:t xml:space=\"preserve\">%s</w:t></w:r></w:p>" % p for p in paragraflar)
    yol = os.path.join(tempfile.mkdtemp(), "b.docx")
    with zipfile.ZipFile(yol, "w") as z:
        z.writestr("word/document.xml", "<w:document xmlns:w=\"%s\"><w:body>%s</w:body></w:document>" % (ns, govde))
    return yol


class DocxKurallari(unittest.TestCase):
    def dt(self, paragraflar):
        return f.alan_oku("dt", _docx(paragraflar))[0]

    def test_yeni_bicim_etiket_iki_nokta_deger_ayri_satirlarda(self):
        """Etiket, ':' ve değer ayrı paragraflarda: ':' değer sanılmamalı."""
        yol = _docx(["DOĞRUDAN TEMİN ONAY BELGESİ", "DOĞRUDAN TEMİN NUMARASI", ":", "26DT5011403", "İŞİN ADI", ":", "Vinç kiralama işi"])
        self.assertEqual(f.alan_oku("dt", yol)[0], "26DT5011403")
        self.assertEqual(f.alan_oku("is", yol)[0], "Vinç kiralama işi")

    def test_eski_bicim_etiketin_hemen_alt_satiri(self):
        yol = _docx(["BAŞLIK", "DOĞRUDAN TEMİN NUMARASI", "26DT141582", "İŞİN ADI", "Eski biçim işi"])
        self.assertEqual(f.alan_oku("dt", yol)[0], "26DT141582")
        self.assertEqual(f.alan_oku("is", yol)[0], "Eski biçim işi")

    def test_numara_etiketle_ayni_satirda(self):
        for satir in ("Doğrudan Temin Numarası: 26DT5011403", "Doğrudan Temin Numarası : 26DT5011403", "Doğrudan Temin Numarası   26DT5011403",
                      "DOĞRUDAN TEMİN NUMARASI - 26DT5011403 (EKAP)", "doğrudan temin numarası:26dt5011403"):
            self.assertEqual(self.dt(["x", satir, "y"]), "26DT5011403", satir)

    def test_etiket_satirinin_sirasi_onemsiz(self):
        for onceki in (0, 4, 11):
            self.assertEqual(self.dt(["satır %d" % i for i in range(onceki)] + ["Doğrudan Temin Numarası : 26DT7000001"]), "26DT7000001")

    def test_etiketle_baslamayan_satir_etiket_sayilmaz(self):
        # 'Doğrudan Temin Kapsamı' ve başlıktaki numara etiket değildir; ancak başlıktaki numara yedek olarak okunur
        self.assertEqual(self.dt(["26DT141582 Doğrudan Temin Alım Kaydı", "Doğrudan Temin Kapsamı"]), "26DT141582")
        self.assertIsNone(self.dt(["Doğrudan Temin Kapsamı", "Hizmet"]))

    def test_etiket_satiri_oncelikli(self):
        """Belgede başka bir DT numarası geçse bile 'Doğrudan Temin Numarası' satırındaki alınır."""
        self.assertEqual(self.dt(["26DT0000001 başlık", "Doğrudan Temin Numarası", ":", "26DT5011403"]), "26DT5011403")

    def test_is_adinda_ayrac_satirlari_atlanir(self):
        yol = _docx(["İŞİN ADI", ":", "–", "Gerçek değer"])
        self.assertEqual(f.alan_oku("is", yol)[0], "Gerçek değer")


class IsAdiTalepFormu(unittest.TestCase):
    """PDF'ten İşin Adı: Talep Formu'nda 'TALEBİN ADI, NİTELİĞİ VE MİKTARI' ile başlayan satır."""

    ETIKET = "TALEBİN ADI, NİTELİĞİ VE MİKTARI :"

    def test_etiketle_ayni_satirda_baslayan_iki_satirlik_yazi(self):
        metin = "\n".join(["TALEP FORMU", "TALEP SAHİBİ BİRİM:X", self.ETIKET + "Bölge Müdürlüğümüze bağlı Hilal TM'de bulunan su",
                           "deposunun (kampüs alanında bulunan bahçe sulama hattına) monte edilmesi işidir.", "13/01/2026", "Mehmet ÇELİK"])
        self.assertEqual(f.is_adi_talep_formundan(metin),
                         "Bölge Müdürlüğümüze bağlı Hilal TM'de bulunan su deposunun (kampüs alanında bulunan bahçe sulama hattına) monte edilmesi işidir.")

    def test_etiket_tek_basina_yazi_alt_satirlarda(self):
        metin = "\n".join([self.ETIKET, "     1- Sepetli vinç [ 1(bir)adet ] 20/01/2026 tarihinde 154 kV GÖLMARMARA TM 'de test ve bakım",
                           "     işinde kullanılmak üzere ihtiyaç duyulmuştur.", "     (Not:Sepetli vinçin Saat 09:00 da bulundurulması)",
                           "20/01/2026", "TALEP YETKİLİSİ"])
        self.assertEqual(f.is_adi_talep_formundan(metin),
                         "1- Sepetli vinç [ 1(bir)adet ] 20/01/2026 tarihinde 154 kV GÖLMARMARA TM 'de test ve bakım işinde kullanılmak üzere "
                         "ihtiyaç duyulmuştur. (Not:Sepetli vinçin Saat 09:00 da bulundurulması)")

    def test_tarih_yoksa_buyuk_harfli_baslikta_durur(self):
        metin = "\n".join([self.ETIKET + " Vinç kiralama işi", "devamı burada", "TALEP YETKİLİSİ", "Uğur ŞEKER"])
        self.assertEqual(f.is_adi_talep_formundan(metin), "Vinç kiralama işi devamı burada")

    def test_satir_sirasi_ve_yazim_farklari_onemsiz(self):
        for onceki in (0, 3, 15):
            m = "\n".join(["satır %d" % i for i in range(onceki)] + ["TALEBIN ADI,  NITELIGI VE  MIKTARI:Test işi", "13/01/2026"])
            self.assertEqual(f.is_adi_talep_formundan(m), "Test işi")
        self.assertEqual(f.is_adi_talep_formundan("talebin adı, niteliği ve miktarı : küçük harf\n01.02.2026"), "küçük harf")

    def test_etiket_yoksa_none(self):
        self.assertIsNone(f.is_adi_talep_formundan("TALEP FORMU\nYAKLAŞIK MALİYET: 5\n13/01/2026"))
        self.assertIsNone(f.is_adi_talep_formundan(self.ETIKET + "\n13/01/2026"))          # etiket var ama yazı boş

    @unittest.skipIf(f.fitz is None, "PyMuPDF kurulu değil")
    def test_pdf_ile_uctan_uca_ve_belge_turu_sarti(self):
        d = tempfile.mkdtemp()

        def pdf(ad, satirlar):
            yol = os.path.join(d, ad)
            belge = f.fitz.open()
            sayfa = belge.new_page()
            for i, s in enumerate(satirlar):
                sayfa.insert_text((50, 60 + 18 * i), s)
            belge.save(yol)
            belge.close()
            return yol
        talep = pdf("talep.pdf", ["TALEP FORMU", "TALEBIN ADI, NITELIGI VE MIKTARI : Vinc kiralama isi (1 adet) (Manisa)", "13/01/2026"])
        self.assertEqual(f.alan_oku("is", talep)[0], "Vinc kiralama isi")                 # sondaki parantezler atılır
        baska = pdf("ihtiyac.pdf", ["HIZMET IHTIYAC FORMU", "TALEBIN ADI, NITELIGI VE MIKTARI : Vinc kiralama isi", "13/01/2026"])
        deger, uyari = f.alan_oku("is", baska)
        self.assertIsNone(deger)                                                           # 'Talep Formu' yazmayan belgeden okunmaz
        self.assertIn("Talep Formu", uyari)


class GozatBirlikteDoldurma(unittest.TestCase):
    """Tek Gözat birden çok alanı doldurur: Talep Formu PDF'inde talep no/tarih + birim + işin adı; DOCX'te işin adı + DT."""

    def test_grup_secimi(self):
        for alan in ("talep", "birim", "is"):
            self.assertEqual(f.okunacak_alanlar(alan, "a.pdf"), ["talep", "birim", "is"], alan)
            self.assertEqual(f.okunacak_alanlar(alan, "A.PDF"), ["talep", "birim", "is"], alan)
        for alan in ("is", "dt"):
            self.assertEqual(f.okunacak_alanlar(alan, "a.docx"), ["is", "dt"], alan)       # DOCX grubu değişmedi
        self.assertEqual(f.okunacak_alanlar("dt", "a.pdf"), ["dt"])
        self.assertEqual(f.okunacak_alanlar("talep", "a.docx"), ["talep"])
        self.assertEqual(f.okunacak_alanlar("birim", "a.docx"), ["birim"])
        self.assertEqual(f.okunacak_alanlar("haric", "a.pdf"), ["haric"])

    @unittest.skipIf(f.fitz is None, "PyMuPDF kurulu değil")
    def test_talep_formu_pdf_uc_alani_doldurur_ve_bir_kez_okunur(self):
        d = tempfile.mkdtemp()

        def pdf(ad, satirlar):
            yol = os.path.join(d, ad)
            belge = f.fitz.open()
            sayfa = belge.new_page()
            for i, s in enumerate(satirlar):
                sayfa.insert_text((40, 50 + 16 * i), s)
            belge.save(yol)
            belge.close()
            return yol
        ust = ["T.C.", "GENEL MUDURLUGU", "3.Bolge Mudurlugu (Izmir)", "Manisa Birimi", "Sayi : E-60994212-934.01.99-3570050  23.01.2026"]
        talep = pdf("talep.pdf", ust + ["TALEP FORMU", "TALEBIN ADI, NITELIGI VE MIKTARI : Vinc kiralama isi (1 adet)", "20/01/2026"])
        sayac = {"n": 0}
        asil = f.pdf_sayfa_metinleri_mupdf
        f.pdf_sayfa_metinleri_mupdf = lambda *a, **k: (sayac.__setitem__("n", sayac["n"] + 1), asil(*a, **k))[1]
        try:
            f._ILK_SAYFA_ONBELLEK.clear()
            sonuc = f.alanlari_oku("birim", talep)
        finally:
            f.pdf_sayfa_metinleri_mupdf = asil
        self.assertEqual(sonuc["talep"][0], "3570050 - 23.01.2026")
        self.assertEqual(sonuc["birim"][0], "Manisa Birimi")
        self.assertEqual(sonuc["is"][0], "Vinc kiralama isi")
        self.assertEqual(sayac["n"], 1, "üç alan için PDF yalnızca bir kez okunmalı")
        # İhtiyaç Formu (gerçek Türkçe metinle): yalnızca birim dolar, talep no ve işin adı boş kalır (Talep Formu şartı)
        ihtiyac_metni = "\n".join(["T.C.", "GENEL MÜDÜRLÜĞÜ", "3.Bölge Müdürlüğü (İzmir)", "Manisa Birimi", "Sayı : E-1-3570050  23.01.2026",
                                   "Hizmet İhtiyaç Formu", "TALEBİN ADI, NİTELİĞİ VE MİKTARI : Vinç işi", "20/01/2026"])
        asil_metin = f._ilk_sayfa_metni
        f._ilk_sayfa_metni = lambda yol: ihtiyac_metni
        try:
            s2 = f.alanlari_oku("talep", talep)
        finally:
            f._ilk_sayfa_metni = asil_metin
        self.assertEqual(s2["birim"][0], "Manisa Birimi")
        self.assertIsNone(s2["talep"][0])
        self.assertIsNone(s2["is"][0])

    @unittest.skipIf(f.fitz is None, "PyMuPDF kurulu değil")
    def test_onbellek_dosya_degisince_yenilenir(self):
        d = tempfile.mkdtemp()
        yol = os.path.join(d, "a.pdf")
        for metin in ("BIRINCI belge metni burada yeterince uzun bir cumle olsun", "IKINCI belge metni burada yeterince uzun bir cumle"):
            belge = f.fitz.open()
            belge.new_page().insert_text((40, 60), metin)
            belge.save(yol)
            belge.close()
            os.utime(yol, (os.path.getmtime(yol) + 5, os.path.getmtime(yol) + 5))          # zaman damgası kesin değişsin
            self.assertIn(metin.split()[0], f._ilk_sayfa_metni(yol))


class Guvenlik(unittest.TestCase):
    def test_guvenli_uzantilar(self):
        for ok in ("a.pdf", "A.PDF", "x.docx", "r.jpeg", "C:\\a b\\c.xlsx"):
            self.assertTrue(f.guvenli_acilir(ok), ok)
        for kotu in ("a.exe", "a.bat", "a.cmd", "a.lnk", "a.js", "a.vbs", "a.ps1", "a.msi", "a.scr", "a.xlsm", "a.docm",
                     "a.pdf.exe", "uzantisiz", "", None):
            self.assertFalse(f.guvenli_acilir(kotu), kotu)

    def test_pdf_ac_calistirilabilir_dosyayi_acmaz(self):
        d = tempfile.mkdtemp()
        yol = os.path.join(d, "zararli.exe")
        open(yol, "wb").write(b"MZ")
        with self.assertRaises(f.GuvensizDosyaTuru):
            f.pdf_ac(yol)
        self.assertTrue(issubclass(f.GuvensizDosyaTuru, OSError))       # arayüzdeki `except OSError` yakalar

    def test_siskinlik_bombasi_sinirlanir(self):
        eski = f._AKIS_EN_COK
        f._AKIS_EN_COK = 1024 * 1024
        try:
            bomba = zlib.compress(b"\0" * (8 * 1024 * 1024))
            self.assertIsNone(f._akis_coz(b"<< /Filter /FlateDecode >>", bomba))
            normal = zlib.compress(b"merhaba")
            self.assertEqual(f._akis_coz(b"<< /Filter /FlateDecode >>", normal), b"merhaba")
        finally:
            f._AKIS_EN_COK = eski

    def test_bozuk_config_ezilmez(self):
        d = tempfile.mkdtemp()
        yol = os.path.join(d, "config.ini")
        icerik = "[ayarlar\nveritabani = benim\\ozel.db\n"          # bozuk bölüm başlığı
        open(yol, "w", encoding="utf-8").write(icerik)
        ayar = f.ayarlari_oku(yol)
        self.assertTrue(ayar["veritabani"].endswith("fatura.db"))     # varsayılanla devam eder, çökmez
        self.assertEqual(open(yol, encoding="utf-8").read(), icerik)   # kullanıcının dosyası ezilmedi


@unittest.skipIf(f.openpyxl is None, "openpyxl kurulu değil")
class ExcelCiktisi(unittest.TestCase):
    def test_formul_enjeksiyonu_ve_baglanti(self):
        import openpyxl
        d = tempfile.mkdtemp()
        pdf = os.path.join(d, "f.pdf"); open(pdf, "wb").write(b"%PDF")
        exe = os.path.join(d, "x.exe"); open(exe, "wb").write(b"MZ")
        x = os.path.join(d, "a.xlsx")
        f.xlsx_yaz(x, [fatura(1, firma_adi='=HYPERLINK("http://x")', pdf_yolu=pdf),
                       fatura(2, fatura_no="F-2", pdf_yolu=exe)], "Test")
        ws = openpyxl.load_workbook(x).active
        self.assertEqual(ws["C4"].data_type, "s")                      # '=...' formül değil metin
        self.assertTrue(ws["B4"].hyperlink)                            # güvenli türde fatura no bağlantılı
        self.assertIsNone(ws["B5"].hyperlink)                          # .exe'ye bağlantı verilmez

    def test_gizli_karakterle_formul_atlatilamaz(self):
        """Başa eklenen görünmez karakter temizlenince '=' başa gelirse de hücre formül olmamalı."""
        import openpyxl
        d = tempfile.mkdtemp()
        kotu = chr(11) + "=HYPERLINK(" + chr(34) + "http://x" + chr(34) + ")"
        for ad, yaz in (("a.xlsx", lambda y: f.xlsx_yaz(y, [fatura(1, firma_adi=kotu, fatura_no=chr(0) + "=1+1")], "T")),
                        ("s.xlsx", lambda y: f.satinalma_xlsx_yaz(y, [fatura(1, firma_adi=kotu, fatura_no=chr(0) + "=1+1")], "T", "2026"))):
            y = os.path.join(d, ad)
            yaz(y)
            ws = openpyxl.load_workbook(y).active
            for satir in ws.iter_rows(max_row=400):                   # şablonun boyutu 1 milyon satır görünür; yalnızca veri satırlarına bak
                for h in satir:
                    if h.data_type == "f":                                # yalnızca uygulamanın kendi toplam formülleri olabilir
                        self.assertTrue(str(h.value).startswith(("=SUBTOTAL(", "=SUM(")), "%s %s formül oldu: %r" % (ad, h.coordinate, h.value))

    def test_kirli_metin_disa_aktarmayi_bozmaz(self):
        import openpyxl
        d = tempfile.mkdtemp()
        kirli = fatura(1, yapilan_is="a" + chr(11) + "b" + chr(12) + "c" + chr(31), firma_adi="F" + chr(0) + "irma", fatura_no="x" * 40000)
        y = os.path.join(d, "k.xlsx")
        f.satinalma_xlsx_yaz(y, [kirli], "T", "2026")
        ws = openpyxl.load_workbook(y).active
        self.assertEqual(ws["G3"].value, "abc")
        self.assertEqual(ws["H3"].value, "Firma")
        self.assertLessEqual(len(ws["J3"].value), f.XLSX_HUCRE_EN_UZUN)             # Excel hücre sınırını aşmaz
        f.csv_yaz(os.path.join(d, "k.csv"), [kirli])                                  # CSV de çökmez

    def test_veritabani_denetim_karakterlerini_temizler(self):
        db, _ = yeni_db()
        i = db.ekle(dict(fatura(1), fatura_no="N1", firma_adi="Ali" + chr(11) + "Veli" + chr(10) + "Can", yapilan_is="x" + chr(0) + "y"), "u")
        r = db.getir(i)
        self.assertEqual((r["firma_adi"], r["yapilan_is"]), ("Ali Veli Can", "x y"))

    def test_satinalma_sablonu_eslesmesi(self):
        import openpyxl
        x = os.path.join(tempfile.mkdtemp(), "s.xlsx")
        f.satinalma_xlsx_yaz(x, [fatura(2), fatura(1, firma_adi="=1+1")], "2026 EKİM HARCAMA LİSTESİ", "2026")
        ws = openpyxl.load_workbook(x).active
        self.assertEqual(ws["A1"].value, "2026 EKİM HARCAMA LİSTESİ")
        satir = [ws.cell(3, k).value for k in range(1, 12)]
        self.assertEqual(satir[0], 1)                                  # SIRA NO
        self.assertEqual(ws["D3"].value, ws["F3"].value)               # talep no = belgenet no
        self.assertEqual(ws["H3"].data_type, "s")                      # '=...' firma adı formül olmaz
        self.assertEqual(ws["K4"].value, 1000.02)                      # KDV hariç tutar (TL)


def _tk_var():
    try:
        import tkinter
        tkinter.Tk().destroy()
        return True
    except Exception:
        return False


@unittest.skipUnless(_tk_var(), "ekran/Tk yok")
class KendiniSina(unittest.TestCase):
    def test_kendini_test_hepsi_tamam(self):
        """Paketlenmiş exe de aynı modla sınanır: FaturaTakip.exe --kendini-test rapor.txt"""
        rapor = os.path.join(tempfile.mkdtemp(), "r.txt")
        with self.assertRaises(SystemExit) as c:
            f.kendini_test(rapor)
        self.assertEqual(c.exception.code, 0, open(rapor, encoding="utf-8").read())


if __name__ == "__main__":
    unittest.main()
