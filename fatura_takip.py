# -*- coding: utf-8 -*-
r"""
Alış Faturası Kayıt ve Takip  -  Sürüm 2

Sadece Python standart kütüphanesi kullanır (tkinter, sqlite3, csv, ...).
Ek paket kurulumu GEREKMEZ.

Çalıştırma:   python fatura_takip.py      (konsolsuz açmak için dosya adını .pyw yapabilirsiniz)

Ayarlar: Programla aynı klasördeki config.ini dosyası.
  [ayarlar]
  veritabani = fatura.db                     ; ağ için: \\sunucu\paylasim\FaturaTakip\fatura.db
  tarama_klasoru =                           ; PDF "Gözat" penceresinin ilk açılacağı klasör (isteğe bağlı)
  tema = Açık                                ; Açık / Koyu / Mavi / Yeşil / Mor / Turuncu / Kırmızı (menüden de değiştirilir)
  kurum_anahtar =                            ; kendi kurumunuzun adından kelime(ler), virgülle (PDF'te alıcıyı firma sanmasın)
  kurum_vkn =                                ; kendi kurumunuzun vergi numarası (isteğe bağlı)

Tutarlar veritabanında KURUŞ (tam sayı) olarak tutulur; böylece yuvarlama hatası olmaz.
"""

import configparser
import csv
import getpass
import hashlib
import hmac
import json
import os
import re
import secrets
import sqlite3
import base64
import shutil
import subprocess
import sys
import tempfile
import threading
import time
import traceback
import zipfile
import zlib
import xml.etree.ElementTree as ET
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

# İsteğe bağlı: PyMuPDF kuruluysa PDF metni onunla okunur (daha güvenilir). Yoksa program
# yerleşik okuyucuyla aynen çalışır. Ağ klasöründeki "kutuphaneler" dizini de aranır.
_KUTUPHANE_KLASORU = os.path.join(os.path.dirname(os.path.abspath(sys.argv[0] if sys.argv and sys.argv[0] else __file__)),
                                  "kutuphaneler")
if os.path.isdir(_KUTUPHANE_KLASORU) and _KUTUPHANE_KLASORU not in sys.path:
    sys.path.append(_KUTUPHANE_KLASORU)
try:
    import openpyxl  # isteğe bağlı: biçimlendirilmiş .xlsx dışa aktarma; yoksa CSV kullanılır
    import openpyxl.styles
    import openpyxl.worksheet.properties
except Exception:
    openpyxl = None
try:
    import sv_ttk  # modern (Windows 11 tarzı) arayüz teması; yoksa klasik tema kullanılır
except Exception:
    sv_ttk = None
try:
    import pymupdf as fitz  # PyMuPDF (yeni ad)
except Exception:
    try:
        import fitz  # PyMuPDF (eski ad)
    except Exception:
        fitz = None

# ----------------------------------------------------------------------------
# 1) AYARLAR (config.ini)
# ----------------------------------------------------------------------------

UYGULAMA_ADI = "Alış Faturası Kayıt ve Takip"
GELISTIRICI = "Yunus Dağ"
IKON_YOLU = os.path.join(getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__))), "ikon.ico")
UYGULAMA_KLASORU = os.path.dirname(os.path.abspath(sys.argv[0] if sys.argv and sys.argv[0] else __file__))
CONFIG_YOLU = os.path.join(UYGULAMA_KLASORU, "config.ini")


def ayarlari_oku(config_yolu=CONFIG_YOLU):
    """config.ini'yi okur; yoksa varsayılan değerlerle oluşturur."""
    cfg = configparser.ConfigParser(interpolation=None)
    varsayilan = {"veritabani": "fatura.db", "tarama_klasoru": "", "tema": "Açık",
                  "kurum_anahtar": "", "kurum_vkn": ""}
    okunabilir = True
    if os.path.exists(config_yolu):
        try:
            cfg.read(config_yolu, encoding="utf-8")
        except (configparser.Error, UnicodeDecodeError, OSError):
            cfg = configparser.ConfigParser(interpolation=None)       # bozuk dosya: varsayılanlarla devam et
            okunabilir = False
    if not cfg.has_section("ayarlar"):
        cfg.add_section("ayarlar")
    degisti = False
    for anahtar, deger in varsayilan.items():
        if not cfg.has_option("ayarlar", anahtar):
            cfg.set("ayarlar", anahtar, deger)
            degisti = True
    if okunabilir and (degisti or not os.path.exists(config_yolu)):     # bozuk dosyanın üzerine yazma
        try:
            with open(config_yolu, "w", encoding="utf-8") as f:
                cfg.write(f)
        except OSError:
            pass  # yazılamıyorsa varsayılanlarla devam
    veritabani = cfg.get("ayarlar", "veritabani").strip() or "fatura.db"
    if not os.path.isabs(veritabani) and not veritabani.startswith("\\\\"):
        veritabani = os.path.join(os.path.dirname(config_yolu), veritabani)
    return {
        "veritabani": veritabani,
        "tarama_klasoru": cfg.get("ayarlar", "tarama_klasoru").strip(),
        "tema": cfg.get("ayarlar", "tema").strip() or "Açık",
        "kurum_anahtar": cfg.get("ayarlar", "kurum_anahtar").strip(),
        "kurum_vkn": re.sub(r"\D", "", cfg.get("ayarlar", "kurum_vkn")),
    }


def ayar_yaz(config_yolu, anahtar, deger):
    """config.ini içindeki tek bir ayarı günceller (yazılamazsa sessizce geçer)."""
    try:
        cfg = configparser.ConfigParser(interpolation=None)
        cfg.read(config_yolu, encoding="utf-8")
        if not cfg.has_section("ayarlar"):
            cfg.add_section("ayarlar")
        cfg.set("ayarlar", anahtar, deger)
        with open(config_yolu, "w", encoding="utf-8") as f:
            cfg.write(f)
    except (OSError, configparser.Error):
        pass


# ----------------------------------------------------------------------------
# 2) YARDIMCI FONKSİYONLAR (tutar / tarih biçimleri)
# ----------------------------------------------------------------------------

def tutar_ayristir(metin):
    """'1.234,56', '1234,56', '1234.56', '1.234' , '1.234,56 TL' -> kuruş (int). Boşsa None."""
    if metin is None:
        return None
    s = str(metin).upper().replace("TL", "").replace("₺", "").replace(" ", "").strip()
    if s == "":
        return None
    if "," in s and "." in s:
        # Hangisi sondaysa ondalık ayracıdır
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif "." in s:
        # Yalnız nokta: 1.234 veya 1.234.567 ise binlik, aksi halde ondalık
        if re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
            s = s.replace(".", "")
    try:
        d = Decimal(s)
    except InvalidOperation:
        raise ValueError("Geçersiz tutar: %s" % metin)
    if d < 0:
        raise ValueError("Tutar negatif olamaz.")
    return int((d * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def tutar_yaz(kurus, tl=False):
    """kuruş (int) -> '1.234,56' (tl=True ise sonuna ' TL')."""
    if kurus is None:
        return ""
    negatif = kurus < 0
    kurus = abs(int(kurus))
    lira, kr = divmod(kurus, 100)
    s = "{:,}".format(lira).replace(",", ".") + "," + "%02d" % kr
    if negatif:
        s = "-" + s
    return s + " TL" if tl else s


def tarih_ayristir(metin):
    """'GG.AA.YYYY' (veya GG/AA/YYYY, GG-AA-YYYY) -> 'YYYY-MM-DD'."""
    s = (metin or "").strip().replace("/", ".").replace("-", ".")
    try:
        return datetime.strptime(s, "%d.%m.%Y").strftime("%Y-%m-%d")
    except ValueError:
        raise ValueError("Tarih GG.AA.YYYY biçiminde olmalı (örnek: 05.03.2026).")


def tarih_yaz(iso):
    """'YYYY-MM-DD' -> 'GG.AA.YYYY'."""
    if not iso:
        return ""
    try:
        return datetime.strptime(iso[:10], "%Y-%m-%d").strftime("%d.%m.%Y")
    except ValueError:
        return iso


_DENETIM_KAR = re.compile(r"[\x00-\x1f\x7f-\x9f]")
_XLSX_YASAK = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")      # Excel'in hücreye yazılmasını reddettiği karakterler
XLSX_HUCRE_EN_UZUN = 32000                                          # Excel hücre sınırı 32.767 karakter


def temiz_metin(s):
    """Tek satırlık alan metninden denetim karakterlerini (satır sonu, sekme, dikey sekme, form besleme…) atar ve
    boşlukları sadeleştirir. PDF/DOCX'ten ya da Word'den yapıştırılan metinde bunlar sık görülür."""
    if s is None:
        return None
    return " ".join(_DENETIM_KAR.sub(" ", str(s)).split())


def xlsx_metin(d):
    """Excel hücresine güvenle yazılabilir değer: yasak karakterleri atar, aşırı uzun metni kısaltır."""
    if isinstance(d, str):
        d = _XLSX_YASAK.sub("", d)
        if len(d) > XLSX_HUCRE_EN_UZUN:
            d = d[:XLSX_HUCRE_EN_UZUN - 1] + "…"
    return d


def tr_kucuk(s):
    """Türkçe harfleri de hesaba katan küçük harfe çevirme (arama/mükerrer kontrolü için)."""
    return (s or "").replace("İ", "i").replace("I", "ı").lower().strip()


def yol_norm(yol):
    return os.path.normcase(os.path.normpath(yol.strip())) if yol and yol.strip() else ""


# ----------------------------------------------------------------------------
# 3) VERİTABANI KATMANI
# ----------------------------------------------------------------------------

class VeritabaniHatasi(Exception):
    pass


class MukerrerKayit(VeritabaniHatasi):
    pass


# ---- kullanıcı hesapları: şifre hash'leme ve "beni hatırla" oturum dosyası ---------------------
SIFRE_EN_AZ = 8
KULLANICI_ADI_EN_UZUN = 40
OTURUM_GUN = 90                       # "beni hatırla" anahtarının geçerlilik süresi
_PBKDF2_TUR = 600_000                 # OWASP önerisi (PBKDF2-HMAC-SHA256); eski hash'ler kendi tur sayısıyla doğrulanır
_SIFRE_HARFLERI = "abcdefghjkmnpqrstuvwxyz23456789"       # 0/O, 1/l/I gibi karışanlar yok


def sifre_hashle(sifre):
    """'pbkdf2$tur$tuz$hash' biçiminde tuzlu hash üretir (düz şifre hiçbir yerde saklanmaz)."""
    tuz = secrets.token_bytes(16)
    h = hashlib.pbkdf2_hmac("sha256", sifre.encode("utf-8"), tuz, _PBKDF2_TUR)
    return "pbkdf2$%d$%s$%s" % (_PBKDF2_TUR, tuz.hex(), h.hex())


def sifre_dogru_mu(sifre, kayit):
    try:
        _, tur, tuz, h = (kayit or "").split("$")
        beklenen = hashlib.pbkdf2_hmac("sha256", sifre.encode("utf-8"), bytes.fromhex(tuz), int(tur))
    except ValueError:
        return False
    return hmac.compare_digest(beklenen.hex(), h)


_sahte = {}


def _sahte_hash():
    """Olmayan kullanıcı adında da aynı süreyi harcamak için (yanıt süresinden hesap varlığı sezilmesin)."""
    if "h" not in _sahte:
        _sahte["h"] = sifre_hashle("sahte-sifre")
    return _sahte["h"]


def gecici_sifre(uzunluk=8):
    return "".join(secrets.choice(_SIFRE_HARFLERI) for _ in range(uzunluk))


def _oturum_dosyasi():
    klasor = os.path.join(os.environ.get("APPDATA") or os.path.expanduser("~"), "FaturaTakip")
    return os.path.join(klasor, "oturum.json")


def oturum_oku(db_yolu):
    """Bu bilgisayarda bu veritabanı için saklanan 'beni hatırla' anahtarı ('' = yok)."""
    try:
        with open(_oturum_dosyasi(), encoding="utf-8") as f:
            return str(json.load(f).get(yol_norm(db_yolu), ""))
    except (OSError, ValueError, AttributeError):
        return ""


def oturum_yaz(db_yolu, token):
    """Anahtarı bu bilgisayarın kullanıcı profiline yazar; token boşsa kaydı siler. Hata sessizce yutulur."""
    dosya = _oturum_dosyasi()
    try:
        try:
            with open(dosya, encoding="utf-8") as f:
                veri = json.load(f)
        except (OSError, ValueError):
            veri = {}
        if not isinstance(veri, dict):
            veri = {}
        if token:
            veri[yol_norm(db_yolu)] = token
        else:
            veri.pop(yol_norm(db_yolu), None)
        os.makedirs(os.path.dirname(dosya), exist_ok=True)
        with open(dosya, "w", encoding="utf-8") as f:
            json.dump(veri, f)
    except OSError:
        pass


ALANLAR = ["fatura_tarihi", "fatura_no", "firma_adi", "yapilan_is",
           "tutar_haric", "tutar_dahil", "pdf_yolu", "kategori", "talep_no", "talep_tarihi", "dt_no", "talep_eden_birim"]

# Yeni Kayıt formundaki kategori listesi ve Excel çıktısındaki sayfa adları (Faturalar 2026.xlsx ile aynı)
KATEGORILER = ["Minibüs-Otobüs Kiralama", "Kamyon Kiralama", "İş Makinesi Kiralama", "Asansör Bakım Onarım",
               "Doğalgaz Faturaları", "Araç Kiralama", "Personel ve Okul Servisi", "Klima-Havalandırma",
               "Isıtma-Tesisat Bakım Onarım", "Akaryakıt Alımı", "Elektrik İşleri", "İş Makinesi Bakım Onarım",
               "Yangın Sistemleri", "HGS", "Genel Alım"]
KATEGORISIZ = "Kategorisiz"     # kategorisi seçilmemiş (eski) kayıtların Excel sayfası


class FaturaDB:
    def __init__(self, yol):
        self.yol = yol
        klasor = os.path.dirname(yol)
        if klasor and not os.path.isdir(klasor):
            raise VeritabaniHatasi("Veritabanı klasörüne ulaşılamıyor:\n%s\n\n"
                                   "Ağ bağlantısını ve config.ini içindeki yolu kontrol edin." % klasor)
        self._olustur()

    def _baglan(self):
        try:
            con = sqlite3.connect(self.yol, timeout=15)
            con.row_factory = sqlite3.Row
            return con
        except sqlite3.Error as e:
            raise VeritabaniHatasi("Veritabanı açılamadı:\n%s\n\n%s" % (self.yol, e))

    def _calistir(self, fonk):
        """Kilitlenme gibi hataları anlaşılır Türkçe mesaja çevirir."""
        con = self._baglan()
        try:
            sonuc = fonk(con)
            con.commit()
            return sonuc
        except sqlite3.OperationalError as e:
            if "locked" in str(e).lower():
                raise VeritabaniHatasi("Veritabanı şu an başka bir kullanıcı tarafından kullanılıyor.\n"
                                       "Birkaç saniye sonra tekrar deneyin.")
            raise VeritabaniHatasi("Veritabanı hatası: %s" % e)
        except sqlite3.Error as e:
            raise VeritabaniHatasi("Veritabanı hatası: %s" % e)
        finally:
            con.close()

    def _olustur(self):
        def f(con):
            con.execute("""
                CREATE TABLE IF NOT EXISTS faturalar (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    fatura_tarihi TEXT NOT NULL,
                    fatura_no TEXT NOT NULL,
                    firma_adi TEXT NOT NULL,
                    yapilan_is TEXT,
                    tutar_haric INTEGER,
                    tutar_dahil INTEGER,
                    pdf_yolu TEXT,
                    kaydeden TEXT,
                    kayit_tarihi TEXT,
                    son_duzenleyen TEXT,
                    duzenleme_tarihi TEXT,
                    firma_anahtar TEXT,
                    pdf_anahtar TEXT
                )""")
            con.execute("CREATE INDEX IF NOT EXISTS ix_tarih ON faturalar(fatura_tarihi)")
            # Sürüm 1'den gelen veritabanını bozmadan yeni sütun/tablo ekle
            sutunlar = [r[1] for r in con.execute("PRAGMA table_info(faturalar)")]
            if "vkn" not in sutunlar:
                con.execute("ALTER TABLE faturalar ADD COLUMN vkn TEXT")
            for yeni in ("kategori", "talep_no", "talep_tarihi", "dt_no", "talep_eden_birim"):
                if yeni not in sutunlar:
                    con.execute("ALTER TABLE faturalar ADD COLUMN %s TEXT" % yeni)
            con.execute("CREATE TABLE IF NOT EXISTS firma_vkn (vkn TEXT PRIMARY KEY, firma_adi TEXT NOT NULL)")
            con.execute("""
                CREATE TABLE IF NOT EXISTS ekler (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    fatura_id INTEGER NOT NULL,
                    tur TEXT,
                    dosya_yolu TEXT NOT NULL,
                    aciklama TEXT,
                    ekleyen TEXT,
                    ekleme_tarihi TEXT
                )""")
            con.execute("CREATE INDEX IF NOT EXISTS ix_ek_fatura ON ekler(fatura_id)")
            con.execute("""
                CREATE TABLE IF NOT EXISTS kullanicilar (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    kullanici_adi TEXT NOT NULL,
                    kullanici_anahtar TEXT NOT NULL UNIQUE,
                    sifre TEXT NOT NULL,
                    yonetici INTEGER NOT NULL DEFAULT 0,
                    aktif INTEGER NOT NULL DEFAULT 1,
                    sifre_degismeli INTEGER NOT NULL DEFAULT 0,
                    olusturma_tarihi TEXT
                )""")
            con.execute("""
                CREATE TABLE IF NOT EXISTS oturumlar (
                    anahtar_hash TEXT PRIMARY KEY,
                    kullanici_id INTEGER NOT NULL,
                    olusturma_tarihi TEXT
                )""")
        self._calistir(f)

    # -- kullanıcı hesapları -------------------------------------------------
    @staticmethod
    def _sifre_kontrol(sifre):
        if len(sifre or "") < SIFRE_EN_AZ:
            raise VeritabaniHatasi("Şifre en az %d karakter olmalı." % SIFRE_EN_AZ)

    def kullanici_sayisi(self):
        return self._calistir(lambda con: con.execute("SELECT COUNT(*) FROM kullanicilar").fetchone()[0])

    def kullanicilar(self):
        def f(con):
            return [dict(r) for r in con.execute(
                "SELECT id, kullanici_adi, yonetici, aktif, sifre_degismeli FROM kullanicilar "
                "ORDER BY kullanici_anahtar")]
        return self._calistir(f)

    def _kullanici_ekle(self, con, ad, sifre, yonetici, sifre_degismeli):
        ad = (ad or "").strip()
        if not ad:
            raise VeritabaniHatasi("Kullanıcı adı boş olamaz.")
        if len(ad) > KULLANICI_ADI_EN_UZUN or any(ord(c) < 32 for c in ad):
            raise VeritabaniHatasi("Kullanıcı adı en fazla %d karakter olmalı ve denetim karakteri içermemeli."
                                   % KULLANICI_ADI_EN_UZUN)
        self._sifre_kontrol(sifre)
        if con.execute("SELECT 1 FROM kullanicilar WHERE kullanici_anahtar=?", (tr_kucuk(ad),)).fetchone():
            raise MukerrerKayit("'%s' adlı bir kullanıcı zaten var." % ad)
        con.execute("INSERT INTO kullanicilar (kullanici_adi, kullanici_anahtar, sifre, yonetici, aktif, "
                    "sifre_degismeli, olusturma_tarihi) VALUES (?,?,?,?,1,?,?)",
                    (ad, tr_kucuk(ad), sifre_hashle(sifre), 1 if yonetici else 0, 1 if sifre_degismeli else 0,
                     datetime.now().strftime("%Y-%m-%d %H:%M:%S")))

    def kullanici_ekle(self, ad, sifre, yonetici=False, sifre_degismeli=True):
        self._calistir(lambda con: self._kullanici_ekle(con, ad, sifre, yonetici, sifre_degismeli))

    def ilk_yonetici_olustur(self, ad, sifre):
        """Yalnızca hiç hesap yokken çalışır (aynı anda iki kişi açarsa ikincisi reddedilir)."""
        def f(con):
            if con.execute("SELECT COUNT(*) FROM kullanicilar").fetchone()[0]:
                raise VeritabaniHatasi("Yönetici hesabı başka biri tarafından az önce oluşturuldu.\n"
                                       "Uygulamayı yeniden açıp kendi hesabınızla giriş yapın.")
            self._kullanici_ekle(con, ad, sifre, True, False)
        self._calistir(f)

    def giris_dene(self, ad, sifre):
        """-> (kayıt | None, hata metni). Pasif hesabın durumu yalnızca şifre doğruysa söylenir."""
        def f(con):
            r = con.execute("SELECT * FROM kullanicilar WHERE kullanici_anahtar=?", (tr_kucuk(ad),)).fetchone()
            if not r:
                sifre_dogru_mu(sifre, _sahte_hash())
                return None, "Kullanıcı adı veya şifre hatalı."
            if not sifre_dogru_mu(sifre, r["sifre"]):
                return None, "Kullanıcı adı veya şifre hatalı."
            if not r["aktif"]:
                return None, "Bu hesap pasife alınmış. Yöneticinize başvurun."
            try:
                eski_tur = int(r["sifre"].split("$")[1])
            except (ValueError, IndexError):
                eski_tur = 0
            if eski_tur < _PBKDF2_TUR:              # eski (daha az turlu) hash'i girişte yenisiyle değiştir
                con.execute("UPDATE kullanicilar SET sifre=? WHERE id=?", (sifre_hashle(sifre), r["id"]))
            return dict(r), ""
        return self._calistir(f)

    def sifre_degistir(self, kid, yeni, eski=None):
        """eski verilirse mevcut şifre doğrulanır. Kullanıcı kendi şifresini seçtiği için 'değiştirmeli' bayrağı kalkar.
        Kullanıcının tüm 'beni hatırla' oturumları düşürülür (çağıran, bu bilgisayar için gerekirse yenisini açar)."""
        self._sifre_kontrol(yeni)

        def f(con):
            r = con.execute("SELECT sifre FROM kullanicilar WHERE id=?", (kid,)).fetchone()
            if not r:
                raise VeritabaniHatasi("Kullanıcı bulunamadı.")
            if eski is not None and not sifre_dogru_mu(eski, r["sifre"]):
                raise VeritabaniHatasi("Mevcut şifre hatalı.")
            con.execute("UPDATE kullanicilar SET sifre=?, sifre_degismeli=0 WHERE id=?", (sifre_hashle(yeni), kid))
            con.execute("DELETE FROM oturumlar WHERE kullanici_id=?", (kid,))   # eski şifreyle açılmış oturumlar düşer
        self._calistir(f)

    def sifre_sifirla(self, kid, gecici):
        """Yönetici işlemi: geçici şifre verir, ilk girişte değiştirtir, kullanıcının 'hatırla' oturumlarını düşürür."""
        self._sifre_kontrol(gecici)

        def f(con):
            con.execute("UPDATE kullanicilar SET sifre=?, sifre_degismeli=1 WHERE id=?", (sifre_hashle(gecici), kid))
            con.execute("DELETE FROM oturumlar WHERE kullanici_id=?", (kid,))
        self._calistir(f)

    def kullanici_durumu(self, kid, aktif=None, yonetici=None):
        """Aktif/yönetici bayrağını değiştirir; en az bir aktif yönetici her zaman kalır."""
        def f(con):
            if aktif is not None:
                con.execute("UPDATE kullanicilar SET aktif=? WHERE id=?", (1 if aktif else 0, kid))
                if not aktif:
                    con.execute("DELETE FROM oturumlar WHERE kullanici_id=?", (kid,))
            if yonetici is not None:
                con.execute("UPDATE kullanicilar SET yonetici=? WHERE id=?", (1 if yonetici else 0, kid))
            if not con.execute("SELECT 1 FROM kullanicilar WHERE yonetici=1 AND aktif=1").fetchone():
                raise VeritabaniHatasi("En az bir aktif yönetici kalmalı.")
        self._calistir(f)

    def yonetici_kurtar(self, ad, sifre):
        """Komut satırı kurtarması: hesabı yoksa açar; varsa yönetici + aktif yapıp şifresini yeniler."""
        def f(con):
            if con.execute("SELECT 1 FROM kullanicilar WHERE kullanici_anahtar=?", (tr_kucuk(ad),)).fetchone():
                self._sifre_kontrol(sifre)
                con.execute("UPDATE kullanicilar SET sifre=?, yonetici=1, aktif=1, sifre_degismeli=0 "
                            "WHERE kullanici_anahtar=?", (sifre_hashle(sifre), tr_kucuk(ad)))
                con.execute("DELETE FROM oturumlar WHERE kullanici_id = "
                            "(SELECT id FROM kullanicilar WHERE kullanici_anahtar=?)", (tr_kucuk(ad),))
            else:
                self._kullanici_ekle(con, ad, sifre, True, False)
        self._calistir(f)

    # "Beni hatırla": bilgisayarda yalnızca rastgele bir anahtar durur; veritabanında yalnızca onun hash'i
    def oturum_olustur(self, kid):
        token = secrets.token_urlsafe(32)
        self._calistir(lambda con: con.execute(
            "INSERT INTO oturumlar (anahtar_hash, kullanici_id, olusturma_tarihi) VALUES (?,?,?)",
            (hashlib.sha256(token.encode()).hexdigest(), kid, datetime.now().strftime("%Y-%m-%d %H:%M:%S"))))
        return token

    def oturum_dogrula(self, token):
        """'Beni hatırla' anahtarı OTURUM_GUN günden eskiyse geçersizdir (süresi dolanlar silinir)."""
        sinir = (datetime.now() - timedelta(days=OTURUM_GUN)).strftime("%Y-%m-%d %H:%M:%S")

        def f(con):
            con.execute("DELETE FROM oturumlar WHERE olusturma_tarihi < ?", (sinir,))
            r = con.execute("SELECT k.* FROM oturumlar o JOIN kullanicilar k ON k.id = o.kullanici_id "
                            "WHERE o.anahtar_hash=? AND k.aktif=1",
                            (hashlib.sha256((token or "").encode()).hexdigest(),)).fetchone()
            return dict(r) if r else None
        return self._calistir(f)

    def oturum_kapat(self, token):
        self._calistir(lambda con: con.execute("DELETE FROM oturumlar WHERE anahtar_hash=?",
                                               (hashlib.sha256((token or "").encode()).hexdigest(),)))

    # -- kontroller ----------------------------------------------------------
    @staticmethod
    def _anahtar_firma(veri):
        return tr_kucuk(veri["firma_adi"]) + "|" + tr_kucuk(veri["fatura_no"])

    def _mukerrer_kontrol(self, con, veri, kendi_id=None):
        cur = con.execute("SELECT id FROM faturalar WHERE firma_anahtar=? AND id IS NOT ?",
                          (self._anahtar_firma(veri), kendi_id))
        if cur.fetchone():
            raise MukerrerKayit("Bu firmaya ait aynı fatura numarası zaten kayıtlı:\n%s  /  %s"
                                % (veri["firma_adi"], veri["fatura_no"]))
        pa = yol_norm(veri.get("pdf_yolu"))
        if pa:
            cur = con.execute("SELECT id, fatura_no, firma_adi FROM faturalar WHERE pdf_anahtar=? AND id IS NOT ?",
                              (pa, kendi_id))
            r = cur.fetchone()
            if r:
                raise MukerrerKayit("Bu PDF dosyası zaten başka bir faturaya bağlı:\n%s  /  %s"
                                    % (r["firma_adi"], r["fatura_no"]))

    # -- işlemler ------------------------------------------------------------
    def ekle(self, veri, kullanici):
        simdi = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        veri = {k: (temiz_metin(v) if isinstance(v, str) else v) for k, v in veri.items()}

        def f(con):
            self._mukerrer_kontrol(con, veri)
            cur = con.execute("""
                INSERT INTO faturalar (fatura_tarihi, fatura_no, firma_adi, yapilan_is, tutar_haric,
                    tutar_dahil, pdf_yolu, kaydeden, kayit_tarihi, firma_anahtar, pdf_anahtar, vkn, kategori,
                    talep_no, talep_tarihi, dt_no, talep_eden_birim)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (veri["fatura_tarihi"], veri["fatura_no"], veri["firma_adi"], veri["yapilan_is"],
                 veri["tutar_haric"], veri.get("tutar_dahil"), veri["pdf_yolu"], kullanici, simdi,
                 self._anahtar_firma(veri), yol_norm(veri.get("pdf_yolu")), veri.get("vkn") or None,
                 veri.get("kategori") or None, veri.get("talep_no"), veri.get("talep_tarihi"), veri.get("dt_no"),
                 veri.get("talep_eden_birim")))
            self._vkn_isle(con, veri)
            return cur.lastrowid
        return self._calistir(f)

    def guncelle(self, fid, veri, kullanici):
        simdi = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        veri = {k: (temiz_metin(v) if isinstance(v, str) else v) for k, v in veri.items()}

        def f(con):
            self._mukerrer_kontrol(con, veri, kendi_id=fid)
            con.execute("""
                UPDATE faturalar SET fatura_tarihi=?, fatura_no=?, firma_adi=?, yapilan_is=?,
                    tutar_haric=?, pdf_yolu=?, son_duzenleyen=?, duzenleme_tarihi=?,
                    firma_anahtar=?, pdf_anahtar=?, vkn=COALESCE(NULLIF(?, ''), vkn), kategori=?,
                    talep_no=?, talep_tarihi=?, dt_no=?, talep_eden_birim=? WHERE id=?""",
                (veri["fatura_tarihi"], veri["fatura_no"], veri["firma_adi"], veri["yapilan_is"],
                 veri["tutar_haric"], veri["pdf_yolu"], kullanici, simdi,
                 self._anahtar_firma(veri), yol_norm(veri.get("pdf_yolu")), veri.get("vkn") or "", veri.get("kategori") or None,
                 veri.get("talep_no"), veri.get("talep_tarihi"), veri.get("dt_no"),
                 veri.get("talep_eden_birim"), fid))
            self._vkn_isle(con, veri)
        self._calistir(f)

    def sil(self, fid):
        def f(con):
            con.execute("DELETE FROM ekler WHERE fatura_id=?", (fid,))
            con.execute("DELETE FROM faturalar WHERE id=?", (fid,))
        self._calistir(f)

    # -- ek evraklar (talep, yazı, irsaliye ...): dosya kopyalanmaz, yalnızca konumu tutulur --
    def ekler(self, fid):
        def f(con):
            return [dict(r) for r in con.execute("SELECT * FROM ekler WHERE fatura_id=? ORDER BY id", (fid,))]
        return self._calistir(f)

    def ek_ekle(self, fid, tur, yol, aciklama, kullanici):
        simdi = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        tur, yol, aciklama = temiz_metin(tur), temiz_metin(yol), temiz_metin(aciklama)

        def f(con):
            if con.execute("SELECT 1 FROM ekler WHERE fatura_id=? AND dosya_yolu=?", (fid, yol)).fetchone():
                raise MukerrerKayit("Bu dosya bu faturaya zaten eklenmiş:\n%s" % yol)
            con.execute("INSERT INTO ekler (fatura_id, tur, dosya_yolu, aciklama, ekleyen, ekleme_tarihi) "
                        "VALUES (?,?,?,?,?,?)", (fid, tur, yol, aciklama, kullanici, simdi))
        self._calistir(f)

    def ek_sil(self, ek_id):
        self._calistir(lambda con: con.execute("DELETE FROM ekler WHERE id=?", (ek_id,)))

    def getir(self, fid):
        def f(con):
            r = con.execute("SELECT * FROM faturalar WHERE id=?", (fid,)).fetchone()
            return dict(r) if r else None
        return self._calistir(f)

    def firmalar(self):
        def f(con):
            return [r[0] for r in con.execute(
                "SELECT DISTINCT firma_adi FROM faturalar ORDER BY firma_adi COLLATE NOCASE")]
        return sorted(self._calistir(f), key=tr_kucuk)

    @staticmethod
    def _vkn_isle(con, veri):
        """VKN okunmuşsa 'vkn -> firma adı' eşlemesini kaydeder (sonraki faturalarda ad otomatik gelsin)."""
        vkn = (veri.get("vkn") or "").strip()
        if vkn and (veri.get("firma_adi") or "").strip():
            con.execute("INSERT OR REPLACE INTO firma_vkn (vkn, firma_adi) VALUES (?, ?)",
                        (vkn, veri["firma_adi"].strip()))

    def vkn_firma(self, vkn):
        """Daha önce kaydedilmiş VKN için firma adını döndürür (yoksa None)."""
        if not vkn:
            return None

        def f(con):
            r = con.execute("SELECT firma_adi FROM firma_vkn WHERE vkn=?", (vkn,)).fetchone()
            return r[0] if r else None
        return self._calistir(f)

    def hepsi(self):
        def f(con):
            return [dict(r) for r in con.execute(
                "SELECT f.*, (SELECT COUNT(*) FROM ekler e WHERE e.fatura_id = f.id) AS ek_sayisi "
                "FROM faturalar f ORDER BY f.fatura_tarihi DESC, f.id DESC")]
        return self._calistir(f)

    def liste(self, arama="", firma="", bas=None, bit=None, yil=None, ay=None):
        """Filtrelenmiş kayıtları döndürür. bas/bit: 'YYYY-MM-DD' veya None."""
        return filtrele(self.hepsi(), arama, firma, bas, bit, yil, ay)


def toplamlar(satirlar):
    adet = len(satirlar)
    haric = sum(r["tutar_haric"] or 0 for r in satirlar)
    return adet, haric


AY_ADLARI = ["Ocak", "Şubat", "Mart", "Nisan", "Mayıs", "Haziran",
             "Temmuz", "Ağustos", "Eylül", "Ekim", "Kasım", "Aralık"]


def filtrele(satirlar, arama="", firma="", bas=None, bit=None, yil=None, ay=None):
    """Kayıt listesini arama, firma, tarih aralığı, yıl ve aya göre süzer (yil/ay: int veya None)."""
    sonuc = []
    a = tr_kucuk(arama)
    for r in satirlar:
        t = r["fatura_tarihi"] or ""
        if firma and r["firma_adi"] != firma:
            continue
        if bas and t < bas:
            continue
        if bit and t > bit:
            continue
        if yil and t[:4] != "%04d" % yil:
            continue
        if ay and t[5:7] != "%02d" % ay:
            continue
        if a:
            havuz = tr_kucuk(" ".join([r["fatura_no"] or "", r["firma_adi"] or "", r["yapilan_is"] or "",
                                    r.get("kategori") or "", r.get("talep_no") or "", r.get("dt_no") or "",
                                    r.get("talep_eden_birim") or ""]))
            if a not in havuz:
                continue
        sonuc.append(r)
    return sonuc


def fatura_ara(satirlar, firma="", fatura_no="", en_az=None, en_cok=None):
    """Fatura Bul: verilen tüm ölçütleri birlikte sağlayan kayıtlar (VE). firma / fatura_no: içinde geçen metin
    (büyük-küçük harf ve Türkçe İ/ı farkı önemsiz; fatura no'da boşluklar yok sayılır). en_az / en_cok: KDV hariç
    tutar sınırları, kuruş cinsinden ve sınırlar dahil. Boş ölçüt yok sayılır; tutar sınırı varken tutarsız kayıt çıkmaz."""
    f = tr_kucuk(firma)
    n = tr_kucuk(fatura_no).replace(" ", "")
    sonuc = []
    for r in satirlar:
        if f and f not in tr_kucuk(r["firma_adi"]):
            continue
        if n and n not in tr_kucuk(r["fatura_no"]).replace(" ", ""):
            continue
        t = r["tutar_haric"]
        if (en_az is not None or en_cok is not None) and t is None:
            continue
        if en_az is not None and t < en_az:
            continue
        if en_cok is not None and t > en_cok:
            continue
        sonuc.append(r)
    return sonuc


def aylik_ozet(satirlar, yil):
    """Seçilen yılın 12 ayı için adet / KDV hariç toplamları + son satırda TOPLAM."""
    sonuc = []
    for k in range(1, 13):
        adet, haric = toplamlar(filtrele(satirlar, yil=yil, ay=k))
        sonuc.append({"ay": AY_ADLARI[k - 1], "adet": adet, "haric": haric})
    sonuc.append({"ay": "TOPLAM", "adet": sum(x["adet"] for x in sonuc),
                  "haric": sum(x["haric"] for x in sonuc)})
    return sonuc


def _xlsx_stilleri():
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    ince = Side(style="thin", color="BFC7D1")
    return {
        "kenar": Border(left=ince, right=ince, top=ince, bottom=ince),
        "baslik_dolgu": PatternFill("solid", fgColor="1F4E79"),
        "baslik_yazi": Font(name="Calibri", bold=True, color="FFFFFF", size=11),
        "zebra": PatternFill("solid", fgColor="EEF3F9"),
        "toplam_dolgu": PatternFill("solid", fgColor="D9E2EF"),
        "toplam_yazi": Font(name="Calibri", bold=True, size=11),
        "normal": Font(name="Calibri", size=11),
        "baslik_yazisi": Font(name="Calibri", bold=True, size=14, color="1F4E79"),
        "orta": Alignment(horizontal="center", vertical="center"),
        "sol": Alignment(horizontal="left", vertical="center", wrap_text=True),
        "sag": Alignment(horizontal="right", vertical="center"),
        "baslik_orta": Alignment(horizontal="center", vertical="center", wrap_text=True),
    }


def _xlsx_sayfa_adi(metin):
    return re.sub(r"[\\/*?:\[\]]", "-", metin)[:31] or "Faturalar"


def xlsx_yaz(yol, satirlar, baslik="Alış Faturaları"):
    """Faturaları biçimlendirilmiş bir Excel (.xlsx) çalışma kitabı olarak yazar (openpyxl gerekir).
    Her kategori ayrı bir sayfadır; kategorisi seçilmemiş eski kayıtlar 'Kategorisiz' sayfasına gider."""
    if openpyxl is None:
        raise RuntimeError("XLSX için 'openpyxl' kütüphanesi kurulu değil.")
    st = _xlsx_stilleri()
    gruplar = {}
    for r in satirlar:
        ad = r.get("kategori") if r.get("kategori") in KATEGORILER else KATEGORISIZ
        gruplar.setdefault(ad, []).append(r)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for ad in KATEGORILER + [KATEGORISIZ]:
        if ad in gruplar:
            _kategori_sayfasi_yaz(wb.create_sheet(_xlsx_sayfa_adi(ad)), gruplar[ad], "%s - %s" % (baslik, ad), st)
    if not wb.sheetnames:
        _kategori_sayfasi_yaz(wb.create_sheet(_xlsx_sayfa_adi(baslik)), [], baslik, st)
    wb.save(yol)


def _kategori_sayfasi_yaz(ws, satirlar, baslik, st):
    from openpyxl.utils import get_column_letter
    sutunlar = ["Fatura Tarihi", "Fatura No", "Firma Adı", "Yapılan İş", "Tutar (KDV Hariç)", "Talep No",
                "Talep Eden Birim", "Ek Evrak", "PDF Konumu"]
    genislik_en_az = [14, 20, 34, 40, 18, 12, 40, 10, 40]
    son = len(sutunlar)

    ws["A1"] = "%s   (%d fatura, düzenlenme: %s)" % (baslik, len(satirlar), date.today().strftime("%d.%m.%Y"))
    ws["A1"].font = st["baslik_yazisi"]
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=son)
    ws.row_dimensions[1].height = 26

    bas = 3
    for k, ad in enumerate(sutunlar, 1):
        h = ws.cell(row=bas, column=k, value=ad)
        h.font, h.fill, h.border, h.alignment = st["baslik_yazi"], st["baslik_dolgu"], st["kenar"], st["baslik_orta"]
    ws.row_dimensions[bas].height = 30

    sirali = sorted(satirlar, key=lambda r: (r["fatura_tarihi"] or "", r["id"]))
    for i, r in enumerate(sirali):
        sat = bas + 1 + i
        try:
            tarih = datetime.strptime((r["fatura_tarihi"] or "")[:10], "%Y-%m-%d").date()
        except ValueError:
            tarih = r["fatura_tarihi"]
        haric = r["tutar_haric"] / 100 if r["tutar_haric"] is not None else None
        degerler = [tarih, r["fatura_no"], r["firma_adi"], r["yapilan_is"] or "", haric,
                    r.get("talep_no") or "", r.get("talep_eden_birim") or "", r.get("ek_sayisi") or None,
                    r["pdf_yolu"] or ""]
        for k, d in enumerate(degerler, 1):
            d = xlsx_metin(d)                      # ÖNCE temizle, kontrolü temizlenmiş değerde yap (gizli karakterle atlatılamasın)
            c = ws.cell(row=sat, column=k, value=d)
            if isinstance(d, str) and d.startswith("="):
                c.data_type = "s"                  # PDF/DOCX'ten gelen '=...' metni formül olarak çalışmasın
            c.font, c.border = st["normal"], st["kenar"]
            if i % 2:
                c.fill = st["zebra"]
            if k == 1:
                c.number_format, c.alignment = "DD.MM.YYYY", st["orta"]
            elif k == 5:
                c.number_format, c.alignment = "#,##0.00", st["sag"]
            elif k in (6, 8):
                c.alignment = st["orta"]
            else:
                c.alignment = st["sol"]
        if r["pdf_yolu"] and guvenli_acilir(r["pdf_yolu"]):
            for sutun in (2, 9):                # Fatura No ve PDF Konumu hücreleri faturanın PDF'ini açar
                try:
                    hucre = ws.cell(row=sat, column=sutun)
                    hucre.hyperlink = r["pdf_yolu"]
                    hucre.font = openpyxl.styles.Font(name="Calibri", size=11, color="0563C1", underline="single")
                except Exception:
                    pass

    ilk, bit = bas + 1, bas + len(sirali)
    tr = bit + 1
    for k in range(1, son + 1):
        c = ws.cell(row=tr, column=k)
        c.font, c.fill, c.border = st["toplam_yazi"], st["toplam_dolgu"], st["kenar"]
    ws.cell(row=tr, column=3, value="TOPLAM (%d fatura)" % len(sirali)).alignment = st["sag"]
    c = ws.cell(row=tr, column=5, value="=SUBTOTAL(109,E%d:E%d)" % (ilk, bit))
    c.number_format, c.alignment = "#,##0.00", st["sag"]
    ws.cell(row=tr, column=8, value="=SUBTOTAL(109,H%d:H%d)" % (ilk, bit)).alignment = st["orta"]

    for k, ad in enumerate(sutunlar, 1):
        en = max([len(str(ws.cell(row=s, column=k).value or "")) for s in range(ilk, bit + 1)] + [len(ad)])
        sinir = 60 if k in (3, 4, 7, 9) else 24
        ws.column_dimensions[get_column_letter(k)].width = min(max(genislik_en_az[k - 1], en + 3), sinir)

    ws.auto_filter.ref = "A%d:%s%d" % (bas, get_column_letter(son), bit)
    ws.freeze_panes = ws.cell(row=bas + 1, column=1)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr = openpyxl.worksheet.properties.PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "%d:%d" % (bas, bas)


SATINALMA_SABLONU = os.path.join(getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__))),
                                 "sablon", "satinalma_sablonu.xlsx")


def tr_buyuk(metin):
    """Türkçe büyük harfe çevirme ('Ekim' -> 'EKİM')."""
    return (metin or "").replace("i", "İ").replace("ı", "I").upper()


def _excel_sayisi(metin):
    """Yalnızca rakamlardan oluşan numaraları (talep no gibi) Excel'de sayı olarak yazmak için int yapar."""
    m = (metin or "").strip()
    return int(m) if m.isdigit() and len(m) <= 15 and not m.startswith("0") else m


def _iso_tarih(iso):
    try:
        return datetime.strptime((iso or "")[:10], "%Y-%m-%d").date()
    except ValueError:
        return None


# Satınalma 'Makine İkmal Harcama Listesi' sütunları  <-  kayıt alanları (başka sütunlar elle doldurulur)
SATINALMA_SUTUNLARI = (
    (2, "dt_no"),                                    # B  EKAP doğrudan temin kayıt no
    (3, "talep_eden_birim"),                         # C  Talep eden birim
    (4, "talep_no"),                                 # D  Talep no
    (5, "talep_tarihi"),                             # E  Belgenet tarihi (= talep tarihi)
    (6, "talep_no"),                                 # F  Belgenet no (= talep no)
    (7, "yapilan_is"),                               # G  İşin konusu
    (8, "firma_adi"),                                # H  Kalan firma
    (9, "fatura_tarihi"),                            # I  Fatura tarihi
    (10, "fatura_no"),                               # J  Fatura no
    (11, "tutar_haric"),                             # K  Ödeme tutarı (KDV hariç)
)


def satinalma_xlsx_yaz(yol, satirlar, baslik, sayfa_adi):
    """Kayıtları satınalma 'Makine İkmal Harcama Listesi' biçiminde yazar: sablon/satinalma_sablonu.xlsx
    açılır (biçim, sütun genişlikleri, filtre, dondurma aynen korunur), kayıtlar 3. satırdan itibaren doldurulur."""
    if openpyxl is None:
        raise RuntimeError("XLSX için 'openpyxl' kütüphanesi kurulu değil.")
    if not os.path.exists(SATINALMA_SABLONU):
        raise FileNotFoundError("Satınalma şablonu bulunamadı:\n%s" % SATINALMA_SABLONU)
    from copy import copy
    wb = openpyxl.load_workbook(SATINALMA_SABLONU)
    ws = wb.active
    ws.title = _xlsx_sayfa_adi(sayfa_adi)
    ws["A1"] = baslik
    ilk, son_sutun = 3, 28
    son = ilk
    while ws.cell(son + 1, 1).value is not None or ws.cell(son + 2, 1).value is not None:
        son += 1                                        # şablondaki numaralı/biçimli satırların sonu (1 satır boşluğa dayanıklı)
    sirali = sorted(satirlar, key=lambda x: (x["fatura_tarihi"] or "", x["id"]))
    son_yazilan = max(son, ilk + len(sirali) - 1)
    for sat in range(ilk, son_yazilan + 1):
        if sat > son:                                   # kayıt sayısı şablondaki satırdan fazlaysa son satırın biçimini çoğalt
            for k in range(1, son_sutun + 1):
                ws.cell(sat, k)._style = copy(ws.cell(son, k)._style)
            ws.row_dimensions[sat].height = ws.row_dimensions[son].height
        ws.cell(sat, 1).value = sat - ilk + 1           # SIRA NO baştan sıralı
        for k in range(2, son_sutun + 1):               # şablonda eski veri kalmış olsa bile temizle
            ws.cell(sat, k).value = None
    for i, rec in enumerate(sirali):
        sat = ilk + i
        for k, alan in SATINALMA_SUTUNLARI:
            d = rec.get(alan)
            if alan in ("talep_tarihi", "fatura_tarihi"):
                d = _iso_tarih(d)
            elif alan == "tutar_haric":
                d = d / 100 if d is not None else None
            elif alan == "talep_no":
                d = _excel_sayisi(d) if d else None
            else:
                d = (d or "").strip() or None
            d = xlsx_metin(d)                           # ÖNCE temizle, kontrolü temizlenmiş değerde yap
            c = ws.cell(sat, k, value=d)
            if isinstance(d, str) and d.startswith("="):
                c.data_type = "s"                       # PDF/DOCX'ten gelen '=...' metni formül olarak çalışmasın
    ws.auto_filter.ref = "A2:%s%d" % (openpyxl.utils.get_column_letter(son_sutun), son_yazilan)
    wb.save(yol)


def ozet_xlsx_yaz(yol, ozet, yil):
    """Aylık özeti biçimlendirilmiş Excel (.xlsx) olarak yazar."""
    if openpyxl is None:
        raise RuntimeError("XLSX için 'openpyxl' kütüphanesi kurulu değil.")
    st = _xlsx_stilleri()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Aylık Özet %s" % yil
    ws["A1"] = "Aylık Fatura Özeti - %s" % yil
    ws["A1"].font = st["baslik_yazisi"]
    ws.merge_cells("A1:C1")
    ws.row_dimensions[1].height = 26
    for k, ad in enumerate(["Ay", "Fatura Adedi", "Tutar (KDV Hariç)"], 1):
        h = ws.cell(row=3, column=k, value=ad)
        h.font, h.fill, h.border, h.alignment = st["baslik_yazi"], st["baslik_dolgu"], st["kenar"], st["orta"]
    ws.row_dimensions[3].height = 24
    for i, x in enumerate(ozet):
        sat = 4 + i
        toplam = x["ay"] == "TOPLAM"
        degerler = [x["ay"], x["adet"], x["haric"] / 100]
        for k, d in enumerate(degerler, 1):
            c = ws.cell(row=sat, column=k, value=d)
            c.border = st["kenar"]
            c.font = st["toplam_yazi"] if toplam else st["normal"]
            if toplam:
                c.fill = st["toplam_dolgu"]
            elif i % 2:
                c.fill = st["zebra"]
            if k == 3:
                c.number_format, c.alignment = "#,##0.00", st["sag"]
            elif k == 2:
                c.alignment = st["orta"]
    for harf, gen in zip("ABC", (16, 16, 22)):
        ws.column_dimensions[harf].width = gen
    ws.freeze_panes = "A4"
    ws.sheet_view.showGridLines = False
    wb.save(yol)


def ozet_csv_yaz(yol, ozet):
    with open(yol, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Ay", "Fatura Adedi", "KDV Hariç Toplam"])
        for x in ozet:
            w.writerow([x["ay"], x["adet"], tutar_yaz(x["haric"])])


def _csv_guvenli(deger):
    """Excel'in formül sanmaması için =, +, -, @ ile başlayan metnin önüne ' koyar (PDF/DOCX'ten gelen metin olabilir)."""
    if isinstance(deger, str):
        deger = _XLSX_YASAK.sub("", deger)
        if deger[:1] in ("=", "+", "-", "@", "\t", "\r"):
            return "'" + deger
    return deger


def csv_yaz(yol, satirlar):
    """Excel'de Türkçe karakter bozulmadan açılır: utf-8-sig + noktalı virgül."""
    with open(yol, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(["Fatura Tarihi", "Fatura No", "Firma Adı", "Yapılan İş", "Kategori",
                    "Fatura Tutarı (KDV Hariç)", "Talep No", "Talep Eden Birim", "PDF Konumu"])
        for r in satirlar:
            w.writerow([tarih_yaz(r["fatura_tarihi"])] + [_csv_guvenli(x) for x in (
                r["fatura_no"], r["firma_adi"], r["yapilan_is"] or "", r.get("kategori") or "")]
                + [tutar_yaz(r["tutar_haric"])] + [_csv_guvenli(x) for x in (
                    r.get("talep_no") or "", r.get("talep_eden_birim") or "", r["pdf_yolu"] or "")])


# ----------------------------------------------------------------------------
# 4) WINDOWS YARDIMCILARI (PDF aç / klasörde göster)
# ----------------------------------------------------------------------------

# Uygulamadan "Aç" ile açılabilen belge türleri. Yollar ortak veritabanından geldiği için (başka biri .exe/.bat/.lnk gibi
# çalıştırılabilir bir dosyayı kayda bağlamış olabilir) yalnızca bu güvenli belge/resim türleri açılır.
GUVENLI_UZANTILAR = frozenset((".pdf", ".doc", ".docx", ".rtf", ".odt", ".txt", ".xls", ".xlsx", ".ods", ".csv",
                               ".ppt", ".pptx", ".jpg", ".jpeg", ".png", ".gif", ".bmp", ".tif", ".tiff", ".eml", ".msg"))


class GuvensizDosyaTuru(OSError):
    pass


def guvenli_acilir(yol):
    """Dosya, açılması güvenli bir belge/resim türünde mi? (Yalnızca uzantıya bakar; var olup olmadığına bakmaz.)"""
    return os.path.splitext(str(yol or "").strip())[1].lower() in GUVENLI_UZANTILAR


def pdf_ac(yol):
    if not yol or not os.path.exists(yol):
        raise FileNotFoundError("Dosya bulunamadı:\n%s\n\nTaşınmış ya da silinmiş olabilir; "
                                "ağ bağlantısını da kontrol edin." % (yol or "(yol girilmemiş)"))
    if not guvenli_acilir(yol):
        raise GuvensizDosyaTuru("Güvenlik nedeniyle bu dosya türü uygulamadan açılmaz:\n%s\n\n"
                                "Dosyayı 'Klasörde Göster' ile bulup kendiniz açabilirsiniz." % yol)
    if hasattr(os, "startfile"):
        os.startfile(yol)  # Windows
    else:
        subprocess.Popen(["xdg-open", yol])


def klasorde_goster(yol):
    if not yol or not os.path.exists(yol):
        raise FileNotFoundError("Dosya bulunamadı:\n%s" % (yol or "(yol girilmemiş)"))
    if os.name == "nt":
        subprocess.Popen(["explorer", "/select,", os.path.normpath(yol)])
    else:
        subprocess.Popen(["xdg-open", os.path.dirname(yol)])


# ----------------------------------------------------------------------------
# 5) PDF'DEN FATURA BİLGİSİ OKUMA  (yalnızca standart kütüphane)
# ----------------------------------------------------------------------------
# Yöntem 1: PDF'in içindeki yazı katmanı (e-fatura PDF'leri ve tarayıcıda "aranabilir PDF"
#           seçeneğiyle taranmış PDF'ler). Hızlıdır ve en güvenilir yoldur.
# Yöntem 2: Windows'un yerleşik OCR'ı (PowerShell üzerinden). Yazı katmanı olmayan,
#           düz resim olarak taranmış PDF'ler için yedek yöntemdir.
# Okunan metinden fatura alanları düzenli ifadelerle (regex) ayıklanır. Sonuç her zaman
# kullanıcı tarafından kontrol edilir; hiçbir şey onaysız kaydedilmez.


class _Ad(bytes):
    """İçerik akışındaki /Ad (name) işlenenlerini string işlenenlerden ayırmak için."""


_ACIK_DIZI = object()
_KIMLIK = (1.0, 0.0, 0.0, 1.0, 0.0, 0.0)


def _carp(a, b):
    return (a[0] * b[0] + a[1] * b[2], a[0] * b[1] + a[1] * b[3],
            a[2] * b[0] + a[3] * b[2], a[2] * b[1] + a[3] * b[3],
            a[4] * b[0] + a[5] * b[2] + b[4], a[4] * b[1] + a[5] * b[3] + b[5])


_AKIS_EN_COK = 64 * 1024 * 1024            # tek bir PDF akışının açılmış en büyük boyutu (sıkıştırma bombası önlemi)
_PDF_EN_COK = 300 * 1024 * 1024            # yerleşik okuyucunun belleğe alacağı en büyük PDF
_DOCX_XML_EN_COK = 30 * 1024 * 1024        # DOCX içindeki document.xml'in açılmış en büyük boyutu


def _akis_coz(sozluk, ham):
    """Akışın filtrelerini (Flate, ASCII85, ASCIIHex) çözer. Desteklenmeyen filtrede None döner."""
    m = re.search(rb"/Filter\s*(\[[^\]]*\]|/[A-Za-z0-9]+)", sozluk)
    if not m:
        return ham
    veri = ham
    for ad in re.findall(rb"/([A-Za-z0-9]+)", m.group(1)):
        try:
            if ad in (b"FlateDecode", b"Fl"):
                dz = zlib.decompressobj()
                veri = dz.decompress(veri, _AKIS_EN_COK)
                if dz.unconsumed_tail:                  # sınır aşıldı: bu akış okunmaz
                    return None
            elif ad in (b"ASCII85Decode", b"A85"):
                v = veri.strip()
                if v.startswith(b"<~"):
                    v = v[2:]
                if v.endswith(b"~>"):
                    v = v[:-2]
                veri = base64.a85decode(v)
            elif ad in (b"ASCIIHexDecode", b"AHx"):
                hx = re.sub(rb"[^0-9A-Fa-f]", b"", veri.split(b">")[0])
                if len(hx) % 2:
                    hx += b"0"
                veri = bytes.fromhex(hx.decode())
            else:
                return None
        except Exception:
            return None
    return veri


def _govde_ayir(govde):
    """'<< sözlük >> stream ... endstream' gövdesini (sözlük, akış) olarak ayırır."""
    m = re.search(rb"(?<!end)stream(?:\r\n|\n|\r)", govde)
    if not m:
        return govde, None
    sozluk = govde[:m.start()]
    ham = govde[m.end():]
    uz = re.search(rb"/Length\s+(\d+)\b(?!\s+\d+\s+R)", sozluk)
    if uz and int(uz.group(1)) <= len(ham):
        ham = ham[:int(uz.group(1))]
    else:
        k = ham.rfind(b"endstream")
        if k >= 0:
            ham = ham[:k]
    return sozluk, ham


def _pdf_nesneleri(veri):
    """PDF'teki nesneleri {no: (sözlük_baytları, akış_baytları|None)} olarak çıkarır (ObjStm dahil)."""
    nesneler = {}
    for m in re.finditer(rb"(?<![0-9])(\d{1,7})\s+\d{1,5}\s+obj\b", veri):
        bit = veri.find(b"endobj", m.end())
        govde = veri[m.end(): bit if bit >= 0 else len(veri)]
        nesneler[int(m.group(1))] = _govde_ayir(govde)
    for no, (sozluk, akis) in list(nesneler.items()):
        if akis is None or not re.search(rb"/Type\s*/ObjStm", sozluk):
            continue
        try:
            d = _akis_coz(sozluk, akis)
            n = int(re.search(rb"/N\s+(\d+)", sozluk).group(1))
            first = int(re.search(rb"/First\s+(\d+)", sozluk).group(1))
            sayilar = re.findall(rb"\d+", d[:first])
            cift = [(int(sayilar[2 * i]), int(sayilar[2 * i + 1])) for i in range(min(n, len(sayilar) // 2))]
            for i, (onum, off) in enumerate(cift):
                son = cift[i + 1][1] if i + 1 < len(cift) else len(d) - first
                nesneler[onum] = (d[first + off: first + son], None)
        except Exception:
            continue
    return nesneler


def _deger_al(s, anahtar):
    """Sözlük baytlarında /Anahtar değerini döndürür (sözlük, dizi, referans veya düz değer)."""
    m = re.search(rb"/" + anahtar + rb"(?![A-Za-z0-9])\s*", s)
    if not m:
        return None
    i = m.end()
    if s[i:i + 2] == b"<<":
        derin, j = 0, i
        while j < len(s):
            if s[j:j + 2] == b"<<":
                derin += 1
                j += 2
            elif s[j:j + 2] == b">>":
                derin -= 1
                j += 2
                if derin == 0:
                    return s[i:j]
            else:
                j += 1
        return s[i:]
    if s[i:i + 1] == b"[":
        k = s.find(b"]", i)
        return s[i:k + 1] if k >= 0 else s[i:]
    r = re.match(rb"\d+\s+\d+\s+R", s[i:])
    if r:
        return r.group(0)
    t = re.match(rb"[^\s/<>\[\]()]+", s[i:])
    return t.group(0) if t else None


def _ref_no(deger):
    m = re.match(rb"\s*(\d+)\s+\d+\s+R", deger or b"")
    return int(m.group(1)) if m else None


def _cozumle(nesneler, deger):
    """Değer bir referanssa o nesnenin sözlüğünü, değilse değerin kendisini verir."""
    no = _ref_no(deger)
    if no is not None:
        return nesneler.get(no, (b"", None))[0]
    return deger or b""


def _hex_utf16(hx):
    hx = hx.decode() if isinstance(hx, bytes) else hx
    if len(hx) % 2:
        hx += "0"
    try:
        return bytes.fromhex(hx).decode("utf-16-be", "ignore")
    except ValueError:
        return ""


def _cmap_oku(cm):
    """ToUnicode CMap'ini okur -> (kod_genişliği_bayt, {kod: metin})."""
    harita = {}
    genislik = 0
    cs = re.search(rb"begincodespacerange(.*?)endcodespacerange", cm, re.S)
    if cs:
        h = re.findall(rb"<([0-9A-Fa-f]+)>", cs.group(1))
        if h:
            genislik = len(h[0]) // 2
    for blok in re.findall(rb"beginbfchar(.*?)endbfchar", cm, re.S):
        for a, b in re.findall(rb"<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]*)>", blok):
            harita[int(a, 16)] = _hex_utf16(b)
    for blok in re.findall(rb"beginbfrange(.*?)endbfrange", cm, re.S):
        for m in re.finditer(rb"<([0-9A-Fa-f]+)>\s*<([0-9A-Fa-f]+)>\s*(\[[^\]]*\]|<[0-9A-Fa-f]*>)", blok):
            lo, hi = int(m.group(1), 16), int(m.group(2), 16)
            if hi < lo or hi - lo > 65535:
                continue
            hedef = m.group(3)
            if hedef.startswith(b"["):
                for k, hx in enumerate(re.findall(rb"<([0-9A-Fa-f]*)>", hedef)):
                    if lo + k <= hi:
                        harita[lo + k] = _hex_utf16(hx)
            else:
                hx = hedef[1:-1].decode()
                if not hx:
                    continue
                taban, uz = int(hx, 16), len(hx)
                for k in range(hi - lo + 1):
                    harita[lo + k] = _hex_utf16(format(taban + k, "0%dx" % uz))
    if not genislik:
        genislik = 2 if any(k > 255 for k in harita) else 1
    return genislik, harita


_VARSAYILAN_FONT = {"genislik": 1, "harita": None, "kodlama": "cp1252", "okunamaz": False}


def _font_bilgisi(nesneler, no, onbellek):
    if no is None:
        return _VARSAYILAN_FONT
    if no in onbellek:
        return onbellek[no]
    bilgi = dict(_VARSAYILAN_FONT)
    try:
        s = nesneler[no][0]
        tu = _deger_al(s, b"ToUnicode")
        tno = _ref_no(tu) if tu else None
        if tno in nesneler and nesneler[tno][1] is not None:
            cm = _akis_coz(nesneler[tno][0], nesneler[tno][1])
            if cm:
                bilgi["genislik"], bilgi["harita"] = _cmap_oku(cm)
        if bilgi["harita"] is None:
            if re.search(rb"/Subtype\s*/Type0", s):
                bilgi["okunamaz"] = True
            enc = _deger_al(s, b"Encoding") or b""
            if b"MacRoman" in enc:
                bilgi["kodlama"] = "mac_roman"
    except Exception:
        pass
    onbellek[no] = bilgi
    return bilgi


def _metni_coz(ham, bilgi):
    if bilgi["harita"] is not None:
        g = bilgi["genislik"] or 1
        cikti = []
        for i in range(0, len(ham) - g + 1, g):
            kod = int.from_bytes(ham[i:i + g], "big")
            cikti.append(bilgi["harita"].get(kod, " " if (kod == 32 and g == 1) else ""))
        return "".join(cikti)
    if bilgi["okunamaz"]:
        return ""
    return ham.decode(bilgi["kodlama"], "ignore")


_ADI_RE = re.compile(rb"/[^\s/<>\[\]()%]*")
_TOKEN_RE = re.compile(rb"[^\s/<>\[\]()%]+")
_SAYI_RE = re.compile(rb"[+-]?(?:\d+\.?\d*|\.\d+)")
_KACIS = {ord("n"): 10, ord("r"): 13, ord("t"): 9, ord("b"): 8, ord("f"): 12}


def _icerik_parcalari(icerik, font_bul):
    """İçerik akışındaki metin parçalarını (y, x, bitiş_x, metin, boyut) olarak döndürür."""
    parcalar = []
    n = len(icerik)
    i = 0
    ope = []
    ctm, ctm_yigin = _KIMLIK, []
    tm = tlm = _KIMLIK
    lead, boyut, bilgi = 0.0, 10.0, _VARSAYILAN_FONT

    def parca_ekle(metin):
        nonlocal tm
        if not metin:
            return
        efb = (boyut * (tm[0] ** 2 + tm[1] ** 2) ** 0.5 * (ctm[0] ** 2 + ctm[1] ** 2) ** 0.5) or boyut
        x = tm[4] * ctm[0] + tm[5] * ctm[2] + ctm[4]
        y = tm[4] * ctm[1] + tm[5] * ctm[3] + ctm[5]
        parcalar.append((y, x, x + len(metin) * 0.5 * efb, metin, efb))
        tm = (tm[0], tm[1], tm[2], tm[3],
              tm[4] + len(metin) * 0.5 * boyut * tm[0], tm[5] + len(metin) * 0.5 * boyut * tm[1])

    while i < n:
        c = icerik[i:i + 1]
        if c in (b" ", b"\t", b"\r", b"\n", b"\x00", b"\x0c"):
            i += 1
            continue
        if c == b"%":
            j = i
            while j < n and icerik[j] not in (10, 13):
                j += 1
            i = j
            continue
        if c == b"(":
            derin, j, buf = 1, i + 1, bytearray()
            while j < n and derin:
                ch = icerik[j]
                if ch == 0x5C:
                    j += 1
                    if j >= n:
                        break
                    e = icerik[j]
                    if e in _KACIS:
                        buf.append(_KACIS[e])
                        j += 1
                    elif 48 <= e <= 55:
                        okt = bytes([e])
                        j += 1
                        while len(okt) < 3 and j < n and 48 <= icerik[j] <= 55:
                            okt += bytes([icerik[j]])
                            j += 1
                        buf.append(int(okt, 8) & 255)
                    elif e in (10, 13):
                        j += 1
                        if e == 13 and j < n and icerik[j] == 10:
                            j += 1
                    else:
                        buf.append(e)
                        j += 1
                elif ch == 0x28:
                    derin += 1
                    buf.append(ch)
                    j += 1
                elif ch == 0x29:
                    derin -= 1
                    if derin:
                        buf.append(ch)
                    j += 1
                else:
                    buf.append(ch)
                    j += 1
            ope.append(bytes(buf))
            i = j
            continue
        if c == b"<":
            if icerik[i:i + 2] == b"<<":
                ope.append(None)
                i += 2
                continue
            j = icerik.find(b">", i)
            if j < 0:
                break
            hx = re.sub(rb"[^0-9A-Fa-f]", b"", icerik[i + 1:j])
            if len(hx) % 2:
                hx += b"0"
            ope.append(bytes.fromhex(hx.decode()))
            i = j + 1
            continue
        if c == b">":
            i += 2 if icerik[i:i + 2] == b">>" else 1
            continue
        if c == b"[":
            ope.append(_ACIK_DIZI)
            i += 1
            continue
        if c == b"]":
            k = len(ope) - 1
            while k >= 0 and ope[k] is not _ACIK_DIZI:
                k -= 1
            dizi = ope[k + 1:] if k >= 0 else []
            del ope[max(k, 0):]
            ope.append(dizi)
            i += 1
            continue
        if c == b"/":
            m = _ADI_RE.match(icerik, i)
            ope.append(_Ad(m.group(0)[1:]))
            i = m.end()
            continue
        m = _TOKEN_RE.match(icerik, i)
        if not m:
            i += 1
            continue
        tok = m.group(0)
        i = m.end()
        if _SAYI_RE.fullmatch(tok):
            ope.append(float(tok))
            continue
        # ---- operatör ----
        if tok == b"BI":                       # satır içi resim: ikili veriyi atla
            k = re.compile(rb"\sEI(?=\s|$)").search(icerik, i)
            i = k.end() if k else n
            ope = []
            continue
        sayilar = [x for x in ope if isinstance(x, float)]
        if tok == b"q":
            ctm_yigin.append(ctm)
        elif tok == b"Q":
            if ctm_yigin:
                ctm = ctm_yigin.pop()
        elif tok == b"cm" and len(sayilar) >= 6:
            ctm = _carp(tuple(sayilar[-6:]), ctm)
        elif tok == b"BT":
            tm = tlm = _KIMLIK
        elif tok == b"Tf" and len(ope) >= 2 and isinstance(ope[-2], _Ad):
            bilgi = font_bul(bytes(ope[-2]))
            if isinstance(ope[-1], float) and ope[-1]:
                boyut = abs(ope[-1])
        elif tok in (b"Td", b"TD") and len(sayilar) >= 2:
            tx, ty = sayilar[-2], sayilar[-1]
            if tok == b"TD":
                lead = -ty
            tlm = _carp((1.0, 0.0, 0.0, 1.0, tx, ty), tlm)
            tm = tlm
        elif tok == b"Tm" and len(sayilar) >= 6:
            tlm = tm = tuple(sayilar[-6:])
        elif tok == b"TL" and sayilar:
            lead = sayilar[-1]
        elif tok in (b"T*", b"'", b'"'):
            tlm = _carp((1.0, 0.0, 0.0, 1.0, 0.0, -lead), tlm)
            tm = tlm
            if tok != b"T*" and ope and isinstance(ope[-1], bytes):
                parca_ekle(_metni_coz(ope[-1], bilgi))
        elif tok == b"Tj" and ope and isinstance(ope[-1], bytes):
            parca_ekle(_metni_coz(ope[-1], bilgi))
        elif tok == b"TJ" and ope and isinstance(ope[-1], list):
            metin = []
            for oge in ope[-1]:
                if isinstance(oge, bytes):
                    metin.append(_metni_coz(oge, bilgi))
                elif isinstance(oge, float) and oge < -250:
                    metin.append(" ")
            parca_ekle("".join(metin))
        ope = []
    return parcalar


def _satirlara_cevir(parcalar):
    """Konumlarına göre parçaları satırlara dizer (yukarıdan aşağıya, soldan sağa)."""
    satirlar = []
    for p in sorted(parcalar, key=lambda p: -p[0]):
        if satirlar and abs(satirlar[-1][0] - p[0]) <= max(2.0, 0.45 * p[4]):
            satirlar[-1][1].append(p)
        else:
            satirlar.append([p[0], [p]])
    cikti = []
    for _, ogeler in satirlar:
        ogeler.sort(key=lambda p: p[1])
        s = ""
        onceki = None
        for p in ogeler:
            if onceki is None:
                s = p[3]
            elif len(onceki[3].strip()) == 1 and len(p[3].strip()) == 1 and p[1] - onceki[2] < 0.3 * p[4]:
                s += p[3]
            else:
                s += " " + p[3]
            onceki = p
        s = re.sub(r"[ \t]+", " ", s).strip()
        if s:
            cikti.append(s)
    return "\n".join(cikti)


def _sayfa_sirasi(nesneler):
    sonuc = []
    kok = None
    for no, (s, _) in nesneler.items():
        if re.search(rb"/Type\s*/Catalog", s):
            p = _deger_al(s, b"Pages")
            kok = _ref_no(p) if p else None
            break

    def gez(no, d=0):
        if no not in nesneler or d > 20:
            return
        s = nesneler[no][0]
        if re.search(rb"/Type\s*/Pages", s):
            for k in re.findall(rb"(\d+)\s+\d+\s+R", _deger_al(s, b"Kids") or b""):
                gez(int(k), d + 1)
        elif re.search(rb"/Type\s*/Page(?![A-Za-z])", s):
            sonuc.append(no)
    if kok is not None:
        gez(kok)
    if not sonuc:
        sonuc = sorted(no for no, (s, _) in nesneler.items() if re.search(rb"/Type\s*/Page(?![A-Za-z])", s))
    return sonuc


def _sayfa_fontlari(nesneler, sozluk, derinlik=0):
    kay = _deger_al(sozluk, b"Resources")
    if kay is None:
        par = _deger_al(sozluk, b"Parent")
        no = _ref_no(par) if par else None
        if no in nesneler and derinlik < 10:
            return _sayfa_fontlari(nesneler, nesneler[no][0], derinlik + 1)
        return {}
    font = _deger_al(_cozumle(nesneler, kay), b"Font")
    if font is None:
        return {}
    font_s = _cozumle(nesneler, font)
    return {ad: int(no) for ad, no in re.findall(rb"/([^\s/<>\[\]()]+)\s+(\d+)\s+\d+\s+R", font_s)}


def pdf_sayfa_metinleri(yol, en_cok_sayfa=60):
    """Yerleşik okuyucuyla her sayfanın yazı katmanını ayrı ayrı okur -> (sayfa_metinleri, toplam_sayfa).
    Yazı katmanı olmayan sayfaların metni boş gelir."""
    if os.path.getsize(yol) > _PDF_EN_COK:
        return [], 0                            # aşırı büyük dosya belleğe alınmaz (PyMuPDF/OCR yoluna kalır)
    with open(yol, "rb") as f:
        veri = f.read()
    if re.search(rb"/Encrypt\b", veri):
        return [], 0
    nesneler = _pdf_nesneleri(veri)
    onbellek = {}
    sayfa_metinleri = []
    sirali = _sayfa_sirasi(nesneler)
    for sayfa_no in sirali[:en_cok_sayfa]:
        s = nesneler[sayfa_no][0]
        icerik_deger = _deger_al(s, b"Contents") or b""
        if icerik_deger.startswith(b"["):
            refler = [int(x) for x in re.findall(rb"(\d+)\s+\d+\s+R", icerik_deger)]
        else:
            r = _ref_no(icerik_deger)
            refler = [r] if r is not None else []
        icerik = b"\n".join(
            (_akis_coz(nesneler[r][0], nesneler[r][1]) or b"")
            for r in refler if r in nesneler and nesneler[r][1] is not None)
        fontlar = _sayfa_fontlari(nesneler, s)
        parcalar = _icerik_parcalari(
            icerik, lambda ad: _font_bilgisi(nesneler, fontlar.get(ad), onbellek))
        sayfa_metinleri.append(_satirlara_cevir(parcalar))
    return sayfa_metinleri, len(sirali)


def pdf_metin_cikar(yol, en_cok_sayfa=5):
    """PDF'in ilk sayfalarındaki yazı katmanı metni (birleşik). Yazı katmanı yoksa boş metin gelir."""
    metinler, _ = pdf_sayfa_metinleri(yol, en_cok_sayfa)
    return "\n\n".join(m for m in metinler if m)


def _metin_yeterli(metin):
    """Çıkan metin gerçekten bir faturaya benziyor mu? (Resim PDF'lerde metin boş ya da anlamsız gelir.)"""
    k = tr_kucuk(metin)
    anahtarlar = ("fatura", "toplam", "tutar", "tarih", "kdv", "vergi", "vkn", "ödenecek", "mal hizmet")
    return len(k) > 80 and sum(1 for a in anahtarlar if a in k) >= 2


# ---- Windows OCR (PowerShell + Windows.Media.Ocr) ----------------------------

_OCR_PS = r'''
$ErrorActionPreference = 'Stop'
try {
Add-Type -AssemblyName System.Runtime.WindowsRuntime
[void][Windows.Storage.StorageFile,Windows.Storage,ContentType=WindowsRuntime]
[void][Windows.Data.Pdf.PdfDocument,Windows.Data.Pdf,ContentType=WindowsRuntime]
[void][Windows.Data.Pdf.PdfPageRenderOptions,Windows.Data.Pdf,ContentType=WindowsRuntime]
[void][Windows.Media.Ocr.OcrEngine,Windows.Foundation,ContentType=WindowsRuntime]
[void][Windows.Media.Ocr.OcrResult,Windows.Foundation,ContentType=WindowsRuntime]
[void][Windows.Graphics.Imaging.BitmapDecoder,Windows.Graphics.Imaging,ContentType=WindowsRuntime]
[void][Windows.Graphics.Imaging.SoftwareBitmap,Windows.Graphics.Imaging,ContentType=WindowsRuntime]
[void][Windows.Storage.Streams.InMemoryRandomAccessStream,Windows.Storage.Streams,ContentType=WindowsRuntime]
[void][Windows.Globalization.Language,Windows.Globalization,ContentType=WindowsRuntime]
$m = [System.WindowsRuntimeSystemExtensions].GetMethods() | Where-Object { $_.Name -eq 'AsTask' -and $_.GetParameters().Count -eq 1 }
$asOp = $m | Where-Object { $_.IsGenericMethod -and $_.GetParameters()[0].ParameterType.Name -eq 'IAsyncOperation`1' } | Select-Object -First 1
$asAct = $m | Where-Object { -not $_.IsGenericMethod -and $_.GetParameters()[0].ParameterType.Name -eq 'IAsyncAction' } | Select-Object -First 1
function Op($o, $t) { $k = $asOp.MakeGenericMethod($t).Invoke($null, @($o)); $k.Wait(-1) | Out-Null; $k.Result }
function Act($o) { $k = $asAct.Invoke($null, @($o)); $k.Wait(-1) | Out-Null }
$dosya = Op ([Windows.Storage.StorageFile]::GetFileFromPathAsync($env:FATURA_PDF)) ([Windows.Storage.StorageFile])
$belge = Op ([Windows.Data.Pdf.PdfDocument]::LoadFromFileAsync($dosya)) ([Windows.Data.Pdf.PdfDocument])
$motor = $null
try { $motor = [Windows.Media.Ocr.OcrEngine]::TryCreateFromLanguage((New-Object Windows.Globalization.Language 'tr')) } catch {}
if (-not $motor) { $motor = [Windows.Media.Ocr.OcrEngine]::TryCreateFromUserProfileLanguages() }
if (-not $motor) { throw 'Windows OCR dil paketi bulunamadi' }
$sb = New-Object System.Text.StringBuilder
$liste = @($env:FATURA_SAYFALAR.Split(',') | ForEach-Object { [int]$_ })
foreach ($i in $liste) {
  if ($i -ge [int]$belge.PageCount) { [void]$sb.Append([char]12); continue }
  $sayfa = $belge.GetPage($i)
  $akis = New-Object Windows.Storage.Streams.InMemoryRandomAccessStream
  $sec = New-Object Windows.Data.Pdf.PdfPageRenderOptions
  $sec.DestinationWidth = [uint32]2480
  Act ($sayfa.RenderToStreamAsync($akis, $sec))
  $akis.Seek([uint64]0)
  $kod = Op ([Windows.Graphics.Imaging.BitmapDecoder]::CreateAsync($akis)) ([Windows.Graphics.Imaging.BitmapDecoder])
  $bmp = Op ($kod.GetSoftwareBitmapAsync()) ([Windows.Graphics.Imaging.SoftwareBitmap])
  $sonuc = Op ($motor.RecognizeAsync($bmp)) ([Windows.Media.Ocr.OcrResult])
  foreach ($s in $sonuc.Lines) { [void]$sb.AppendLine($s.Text) }
  [void]$sb.Append([char]12)
  $sayfa.Dispose()
}
[IO.File]::WriteAllText($env:FATURA_OUT, $sb.ToString(), (New-Object System.Text.UTF8Encoding $false))
} catch { [Console]::Error.WriteLine($_.Exception.Message); exit 2 }
'''


def ocr_ile_oku(yol, sayfalar=(0, 1, 2), zaman_asimi=None):
    """Windows'un yerleşik OCR'ı ile PDF'in verilen sayfalarını (0 tabanlı dizinler) okur.
    -> sayfalarla aynı sırada metin listesi. Windows dışında hata verir."""
    if os.name != "nt":
        raise RuntimeError("Windows OCR yalnızca Windows'ta çalışır.")
    sayfalar = [int(i) for i in sayfalar]
    if not sayfalar:
        return []
    zaman_asimi = zaman_asimi or max(150, 30 * len(sayfalar))
    klasor = tempfile.mkdtemp(prefix="fatura_ocr_")
    try:
        yerel = os.path.join(klasor, "belge.pdf")      # ağ yolundan değil, yerel kopyadan okunur
        shutil.copyfile(yol, yerel)
        cikti = os.path.join(klasor, "cikti.txt")
        env = dict(os.environ, FATURA_PDF=yerel, FATURA_OUT=cikti,
                   FATURA_SAYFALAR=",".join(str(i) for i in sayfalar))
        betik = os.path.join(klasor, "ocr.ps1")        # -EncodedCommand yerine dosyadan çalıştırılır (daha güvenilir)
        with open(betik, "w", encoding="utf-8-sig") as bf:
            bf.write(_OCR_PS)
        p = subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-ExecutionPolicy", "Bypass", "-File", betik],
            env=env, capture_output=True, timeout=zaman_asimi,
            creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        if p.returncode != 0 or not os.path.exists(cikti):
            hata = (p.stderr or b"").decode("utf-8", "ignore").strip() or "bilinmeyen hata"
            raise RuntimeError(hata[:300])
        with open(cikti, encoding="utf-8") as f:
            parcalar = [x.strip() for x in f.read().split("\f")]
        return (parcalar + [""] * len(sayfalar))[:len(sayfalar)]
    except subprocess.TimeoutExpired:
        raise RuntimeError("OCR zaman aşımına uğradı.")
    except FileNotFoundError:
        raise RuntimeError("PowerShell bulunamadı.")
    finally:
        shutil.rmtree(klasor, ignore_errors=True)


# ---- Metinden fatura alanlarını ayıklama -----------------------------------------

_TUTAR_RE = r"(\d{1,3}(?:\.\d{3})+,\d{2}|\d+,\d{2})"
_TARIH_RE = r"(\d{1,2}\s*[./-]\s*\d{1,2}\s*[./-]\s*\d{4})"
_ORANLAR = (0, 1, 8, 10, 18, 20)
_I = "[İIiı]"
_FIRMA_SONEK = re.compile(
    r"(?<![A-Za-zÇĞİÖŞÜçğıöşü])A\.?\s?Ş\b|ANON" + _I + r"M|L" + _I + r"M" + _I + r"TED|\bLTD\b|\bŞT" + _I +
    r"\b|SAN\.|T" + _I + r"C\.|KOLL\.|ORTAKLI[ĞG]I|HOLD" + _I + r"NG|\bŞ" + _I + r"RKET" + _I, re.I)
_KARA_LISTE = ("sayın", "fatura", "tarih", "vkn", "tckn", "vergi", "adres", "tel", "mersis", "ticaret sicil",
               "web", "posta", "senaryo", "ettn", "irsaliye", "sipariş", "ödeme", "banka", "iban", "e-arşiv",
               "e-fatura", "no:", "no ", "mah.", "cad.", "sok.")
_BIRIMLER = {"adet", "ad", "kg", "kgm", "mt", "m", "mtr", "lt", "ltr", "paket", "set", "kutu", "saat", "gün",
             "gun", "ay", "c62", "niu", "top", "çift", "rulo", "metre"}


def _tarih_bul(metin):
    for et in (r"Fatura\s*Tarihi", r"D.zenlenme\s*Tarihi", r"D.zenleme\s*Tarihi"):
        for m in re.finditer(et + r"\s*:?\s*" + _TARIH_RE, metin, re.I):
            try:
                iso = tarih_ayristir(re.sub(r"\s+", "", m.group(1)))
                if 2000 <= int(iso[:4]) <= 2100:
                    return iso
            except ValueError:
                pass
    for m in re.finditer(_TARIH_RE, metin):
        try:
            iso = tarih_ayristir(re.sub(r"\s+", "", m.group(1)))
            if 2000 <= int(iso[:4]) <= 2100:
                return iso
        except ValueError:
            pass
    return None


def _fatura_no_bul(metin):
    m = re.search(r"\b([A-Z0-9]{3}20\d{2}\d{9})\b", metin.upper())
    if m:
        return m.group(1)
    m = re.search(r"Fatura\s*(?:No|Numaras.)\s*[:.]?\s*([A-Za-z0-9][A-Za-z0-9\-/_.]{2,24})", metin, re.I)
    if m:
        return m.group(1).strip(".")
    return None


def _tutar_etiket(metin, kaliplar):
    for k in kaliplar:
        for m in re.finditer(k, metin, re.I):
            a = re.search(_TUTAR_RE, metin[m.end(): m.end() + 90])
            if a:
                try:
                    return tutar_ayristir(a.group(1))
                except ValueError:
                    pass
    return None


def _oran_uyumlu(haric, dahil):
    return any(abs(haric * (100 + o) - dahil * 100) <= 100 for o in _ORANLAR)


def _tutar_cifti_bul(metin):
    """Etiket bulunamazsa: metindeki tutarlar arasında 'KDV hariç x (1+oran) = KDV dahil' olan çifti ara."""
    tutarlar = set()
    for m in re.finditer(_TUTAR_RE, metin):
        try:
            tutarlar.add(tutar_ayristir(m.group(1)))
        except ValueError:
            pass
    tutarlar = sorted((t for t in tutarlar if t > 0), reverse=True)
    for d in tutarlar:
        for h in tutarlar:
            if h < d and _oran_uyumlu(h, d):
                return h, d
    return None


def _tutarlari_bul(metin):
    mal = _tutar_etiket(metin, [r"Mal\s*Hizmet\s*Toplam\s*Tutar"])
    isk = _tutar_etiket(metin, [r"Toplam\s*.skonto"])
    kdv = _tutar_etiket(metin, [r"Hesaplanan\s*K\.?D\.?V"])
    dahil = _tutar_etiket(metin, [r"Vergiler\s*Dahil\s*Toplam\s*Tutar", r"Vergiler\s*Dahil",
                                  r"KDV\s*Dahil\s*Toplam", r".denecek\s*Tutar", r"Genel\s*Toplam"])
    haric = None
    if mal is not None:
        haric = mal - (isk or 0)
    else:
        haric = _tutar_etiket(metin, [r"Vergiler\s*Hari.", r"KDV\s*Hari.\s*Toplam", r"Matrah", r"Ara\s*Toplam"])
    if haric is None and dahil is not None and kdv is not None:
        haric = dahil - kdv
    if dahil is None and haric is not None and kdv is not None:
        dahil = haric + kdv
    tutarli = haric is not None and dahil is not None and 0 <= haric <= dahil and _oran_uyumlu(haric, dahil)
    if not tutarli:
        cift = _tutar_cifti_bul(metin)
        if cift:
            haric, dahil = cift
        elif haric is not None and dahil is not None and haric > dahil:
            haric = dahil = None
    return haric, dahil


_VKN_ETIKET = re.compile(
    r"(?:VKN|TCKN|Vergi\s*Kimlik\s*(?:No|Numaras.)|T\.?\s?C\.?\s*Kimlik\s*(?:No|Numaras.)|"
    r"Vergi\s*(?:No|Numaras.)|V\.?\s?K\.?\s?N\.?)\s*[:/.]?\s*(?<!\d)(\d{10}|\d{11})(?!\d)", re.I)


def _vkn_bul(metin, kurum_vkn=""):
    """Satıcının VKN (10 hane) ya da TCKN (11 hane, şahıs işletmesi) numarasını bulur.
    Önce 'SAYIN' (alıcı) bölümünden önceye bakar; iki sütunlu faturalarda satıcı bilgisi
    'Sayın'dan sonra gelebildiği için bulunamazsa tüm metinde ilk uygun numara alınır."""
    sayin = re.search(r"SAYIN", metin.replace("ı", "I").upper())
    bolgeler = [metin[:sayin.start()]] if sayin else []
    bolgeler.append(metin)
    for bolge in bolgeler:
        for m in _VKN_ETIKET.finditer(bolge):
            if m.group(1) != kurum_vkn:
                return m.group(1)
    return None


_SADE = str.maketrans("ıİIşŞğĞüÜöÖçÇ", "iiissgguuoocc")


def _sade(s):
    """Türkçe harfleri ASCII'ye indirger (OCR 'TEİAŞ' yerine 'TEIAS' okusa da eşleşsin)."""
    return (s or "").translate(_SADE).lower().strip()


_BELGE_ONEKI = re.compile(r"^\s*e\s*-\s*(?:fatura|ar[şs]iv(?:\s*fatura)?)\s*", re.I)
_DEVAM_RE = re.compile(r"^(?:T[İI]C|T[İI]CARET|LTD|ŞT[İI]|SAN|SANAY[İI]|A\.?\s?Ş|ANON[İI]M|L[İI]M[İI]TED|"
                       r"[İI]TH|[İI]HR|PAZ|MAK[İI]NE)", re.I)


def _devam_satiri_ekle(ad, bolge, i):
    """Firma adı satır sonunda bağlaçla bitiyorsa ya da sonraki satır unvan devamı gibi duruyorsa (TİC.LTD.ŞTİ) ekler."""
    if i + 1 >= len(bolge):
        return ad
    sonraki = bolge[i + 1].strip()
    bagla_biter = re.search(r"(VE|&|T[İI]C\.?|SAN\.?)\s*$", ad, re.I)
    kisa_metin = not re.search(r"\d", sonraki) and len(sonraki) <= 40
    if kisa_metin and (bagla_biter or not _FIRMA_SONEK.search(ad)) and _DEVAM_RE.match(sonraki):
        return ad + " " + sonraki
    if kisa_metin and bagla_biter and not any(b in tr_kucuk(sonraki) for b in _KARA_LISTE):
        return ad + " " + sonraki
    return ad


def _firma_bul(satirlar, kurum_anahtar=""):
    anahtarlar = [_sade(a) for a in kurum_anahtar.split(",") if a.strip()]
    bolge = []
    for s in satirlar[:45]:
        if "sayın" in tr_kucuk(s):
            break
        bolge.append(s)

    def kurum_mu(s):
        k = _sade(s)
        return any(a in k for a in anahtarlar)

    adaylar = [i for i, s in enumerate(bolge)
               if _FIRMA_SONEK.search(s) and not kurum_mu(s) and len(re.findall(r"[A-Za-zÇĞİÖŞÜçğıöşü]{2,}", s)) >= 2]
    if adaylar:
        i = adaylar[0]
        ad = _BELGE_ONEKI.sub("", bolge[i])
        ad = _devam_satiri_ekle(ad, bolge, i)
        baglac = re.match(r"\s*(VE|T[İI]CARET|SANAY[İI]|T[İI]C\.|SAN\.)\b", ad, re.I)
        if baglac and i > 0 and not kurum_mu(bolge[i - 1]) and not re.search(r"\d", bolge[i - 1]):
            ad = bolge[i - 1] + " " + ad
        return re.sub(r"\s+", " ", ad).strip(" -:,")
    for s in bolge:
        harfler = re.findall(r"[A-Za-zÇĞİÖŞÜçğıöşü]", s)
        if len(harfler) < 6 or kurum_mu(s):
            continue
        if any(b in tr_kucuk(s) for b in _KARA_LISTE) or re.search(r"\d{4,}", s):
            continue
        if sum(1 for h in harfler if h == h.upper()) / len(harfler) > 0.7:
            return re.sub(r"\s+", " ", s).strip(" -:,")
    return None


def _is_bul(satirlar):
    bas = None
    for i, s in enumerate(satirlar):
        k = tr_kucuk(s)
        if re.search(r"mal\s*hizmet|a[çc][ıi]klama", k) and re.search(r"miktar|birim|fiyat", k):
            bas = i
            break
    if bas is None:
        return None
    kalemler = []
    for s in satirlar[bas + 1: bas + 9]:
        k = tr_kucuk(s)
        if re.search(r"toplam|iskonto|hesaplanan|kdv|yaln[ıi]z|not\s*:", k):
            break
        parcalar = re.sub(r"^\s*\d{1,3}\s+", "", s).split()
        cikti = []
        for j, t in enumerate(parcalar):
            tl = tr_kucuk(t)
            if cikti and (re.fullmatch(r"%?\d+[.,]\d+%?|%\d+", t) or tl in _BIRIMLER):
                break
            if cikti and re.fullmatch(r"\d+", t) and j + 1 < len(parcalar) and tr_kucuk(parcalar[j + 1]) in _BIRIMLER:
                break
            cikti.append(t)
        ad = " ".join(cikti).strip()
        if len(ad) >= 3 and re.search(r"[A-Za-zÇĞİÖŞÜçğıöşü]{3}", ad):
            kalemler.append(ad)
        if len(kalemler) == 2:
            break
    return ("; ".join(kalemler))[:120] if kalemler else None


def fatura_metnini_coz(metin, ayar=None):
    """Okunan metinden alanları ayıklar. Bulunamayanlar sonuca girmez. Tutarlar kuruş, tarih ISO biçimindedir."""
    ayar = ayar or {}
    metin = (metin or "").replace("\xa0", " ")
    satirlar = [s.strip() for s in metin.splitlines() if s.strip()]
    sonuc = {}
    tarih = _tarih_bul(metin)
    if tarih:
        sonuc["fatura_tarihi"] = tarih
    no = _fatura_no_bul(metin)
    if no:
        sonuc["fatura_no"] = no
    vkn = _vkn_bul(metin, ayar.get("kurum_vkn", ""))
    if vkn:
        sonuc["vkn"] = vkn
    firma = _firma_bul(satirlar, ayar.get("kurum_anahtar", ""))
    if firma:
        sonuc["firma_adi"] = firma
    haric, dahil = _tutarlari_bul(metin)
    if haric is not None:
        sonuc["tutar_haric"] = haric
    if dahil is not None:
        sonuc["tutar_dahil"] = dahil
    return sonuc


def pdf_sayfa_metinleri_mupdf(yol, en_cok_sayfa=60):
    """PyMuPDF kuruluysa sayfa sayfa metin çıkarır -> (sayfa_metinleri, toplam_sayfa); yoksa ([], 0)."""
    if fitz is None:
        return [], 0
    try:
        with fitz.open(yol) as belge:
            if belge.needs_pass:
                return [], 0
            toplam = len(belge)
            return [belge[i].get_text("text", sort=True).strip() for i in range(min(toplam, en_cok_sayfa))], toplam
    except Exception:
        return [], 0


_FATURA_IPUCLARI = ("fatura no", "fatura numaras", "fatura tarihi", "ettn", "mal hizmet toplam", "vergiler dahil",
                    "ödenecek tutar", "hesaplanan kdv", "e-arşiv", "e-fatura", "senaryo", "özelleştirme",
                    "düzenlenme tarihi", "sayın")
_OCR_SAYFA_SINIRI = 20


def _fatura_puani(metin):
    """Sayfada kaç farklı 'fatura ipucu' var? (Fatura sayfasını talep/yazı sayfalarından ayırmak için.)"""
    k = tr_kucuk(metin)
    return sum(1 for a in _FATURA_IPUCLARI if a in k)


def sayfa_ayristir(metin, toplam):
    """'7', '7-8' veya '2,5' -> 0 tabanlı sayfa dizinleri (boşsa None)."""
    s = (metin or "").strip().replace(" ", "")
    if not s:
        return None
    biçim_hata = ValueError("Fatura sayfası '7', '7-8' ya da '2,5' biçiminde yazılmalı.")
    numaralar = []
    try:
        for parca in s.split(","):
            if "-" in parca:
                a, b = (int(x) for x in parca.split("-", 1))
                numaralar += list(range(min(a, b), max(a, b) + 1))
            else:
                numaralar.append(int(parca))
    except ValueError:
        raise biçim_hata
    if not numaralar or len(numaralar) > 20:
        raise biçim_hata
    if any(n < 1 or (toplam and n > toplam) for n in numaralar):
        raise ValueError("Fatura sayfası 1 ile %d arasında olmalı." % toplam if toplam else "Sayfa numarası geçersiz.")
    return sorted({n - 1 for n in numaralar})


def pdf_bilgi_cikar(yol, ayar=None, sayfa=None):
    """PDF'i okur ve fatura sayfasındaki bilgileri ayıklar.
    sayfa: None/boş = fatura sayfasını otomatik bul; '7', '7-8', '2,5' = yalnızca bu sayfalar okunur.
    Okuma sırası: PyMuPDF (varsa) -> yerleşik yazı katmanı okuyucu -> Windows OCR (yazısı olmayan sayfalar).
    -> {yontem, metin, alanlar, hata, sayfa, sayfa_sayisi, sayfa_bilgisi}"""
    sonuc = {"yontem": None, "metin": "", "alanlar": {}, "hata": None,
             "sayfa": None, "sayfa_sayisi": None, "sayfa_bilgisi": ""}
    sayfalar, toplam = pdf_sayfa_metinleri_mupdf(yol)
    if not any(_metin_yeterli(m) for m in sayfalar):
        try:
            s2, t2 = pdf_sayfa_metinleri(yol)
        except Exception:
            s2, t2 = [], 0
        if any(_metin_yeterli(m) for m in s2) or not sayfalar:
            sayfalar = s2
        toplam = max(toplam, t2)
    sayfalar = list(sayfalar)
    n = max(toplam, len(sayfalar))
    sonuc["sayfa_sayisi"] = n or None

    ocr = {"hata": None, "yapilan": set()}

    def ocr_calistir(dizinler):
        if not dizinler:
            return
        if os.name != "nt":
            ocr["hata"] = "PDF'te okunabilir yazı katmanı yok (resim olarak taranmış). OCR yalnızca Windows'ta çalışır."
            return
        try:
            metinler = ocr_ile_oku(yol, dizinler)
        except Exception as ex:
            ocr["hata"] = "Windows OCR çalışmadı: %s" % ex
            return
        for i, m in zip(dizinler, metinler):
            if len(m.strip()) > 40:
                while len(sayfalar) <= i:
                    sayfalar.append("")
                sayfalar[i] = m
                ocr["yapilan"].add(i)

    def en_iyi_sayfa():
        puanlar = [(_fatura_puani(m) if _metin_yeterli(m) else 0) for m in sayfalar]
        if not puanlar or max(puanlar) < 2:
            return None
        return puanlar.index(max(puanlar))

    istenen = sayfa_ayristir(sayfa, n)
    secilen = None
    if istenen:
        ocr_calistir([i for i in istenen if i >= len(sayfalar) or not _metin_yeterli(sayfalar[i])])
        secilen = istenen
        sonuc["sayfa_bilgisi"] = "Yalnızca %s. sayfa okundu." % ", ".join(str(i + 1) for i in istenen)
    else:
        i = en_iyi_sayfa()
        if i is None:                # yazı katmanı olmayan (taranmış) sayfaları OCR'dan geçirip tekrar dene
            hedef = [k for k in range(min(n or 3, _OCR_SAYFA_SINIRI))
                     if k >= len(sayfalar) or len(sayfalar[k].strip()) < 40]
            ocr_calistir(hedef)
            i = en_iyi_sayfa()
        if i is not None:
            secilen = [i]
            if n > 1:
                sonuc["sayfa_bilgisi"] = ("%d sayfalık belgede fatura %d. sayfada bulundu. Yanlışsa 'Fatura sayfası' "
                                          "kutusuna doğru sayfayı yazıp 'PDF'den Oku'ya basın." % (n, i + 1))

    if secilen:
        metin = "\n\n".join(sayfalar[i] for i in secilen if i < len(sayfalar))
        sonuc["sayfa"] = ", ".join(str(i + 1) for i in secilen)
        kaynak_ocr = any(i in ocr["yapilan"] for i in secilen)
    else:                            # fatura sayfası ayırt edilemedi: eski davranış, ilk sayfaların metni
        ilk = range(min(5, len(sayfalar)))
        metin = "\n\n".join(sayfalar[i] for i in ilk if sayfalar[i].strip())
        kaynak_ocr = any(i in ocr["yapilan"] for i in ilk)
        if n > 1 and metin.strip():
            sonuc["sayfa_bilgisi"] = ("Fatura sayfası otomatik belirlenemedi, ilk sayfalar okundu. Doğru sayfayı "
                                      "'Fatura sayfası' kutusuna yazıp 'PDF'den Oku'ya basın.")
    if metin.strip() and (secilen or _metin_yeterli(metin)):
        sonuc["yontem"] = "ocr" if kaynak_ocr else "metin"
    if sonuc["yontem"] is None or not _metin_yeterli(metin):
        sonuc["hata"] = ocr["hata"] or (None if sonuc["yontem"] else
                                        "OCR çalıştı ama sayfada okunabilir yazı bulamadı (tarama çok silik/eğik olabilir).")
    sonuc["metin"] = metin
    if metin.strip():
        sonuc["alanlar"] = fatura_metnini_coz(metin, ayar)
    return sonuc


# ----------------------------------------------------------------------------
# 5b) ALAN OKUMA KURALLARI  (Yeni Kayıt formundaki her satırın "Gözat" butonu)
# ----------------------------------------------------------------------------
# Her kural: kural(yol, secenek) -> (deger_metni | None, not_metni)
# Kuralların kaydı OKUMA_KURALLARI sözlüğündedir: (alan, dosya_uzantisi) -> kural.
# Fatura tarihi / no / firma / tutar bu sözlükte değildir: fatura PDF'i tek seferde okunur (pdf_bilgi_cikar).
# Yeni bir alan/dosya türü eklemek için kuralı yazıp sözlüğe bir satır eklemek yeterlidir.

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def docx_satirlari(yol):
    """DOCX'in metnini belge sırasıyla satır satır verir (paragraflar ve tablo hücreleri dahil).
    Ek kütüphane gerekmez: DOCX bir zip, içindeki word/document.xml okunur."""
    try:
        with zipfile.ZipFile(yol) as z:
            if z.getinfo("word/document.xml").file_size > _DOCX_XML_EN_COK:
                raise ValueError("DOCX dosyası beklenenden çok büyük, okunmadı.")
            ham = z.read("word/document.xml")
        if b"<!DOCTYPE" in ham or b"<!ENTITY" in ham:   # gerçek Word belgelerinde olmaz; XXE / "billion laughs" önlemi
            raise ValueError("DOCX dosyası güvenli görünmüyor, okunmadı.")
        kok = ET.fromstring(ham)
    except (zipfile.BadZipFile, KeyError, ET.ParseError):
        raise ValueError("DOCX dosyası açılamadı (bozuk ya da eski .doc biçimi olabilir).")
    satirlar = []
    for p in kok.iter(_W + "p"):
        parca = []
        for e in p.iter():
            if e.tag == _W + "t":
                parca.append(e.text or "")
            elif e.tag in (_W + "tab", _W + "br"):
                parca.append(" ")
        t = " ".join("".join(parca).split())
        if t:
            satirlar.append(t)
    return satirlar


def _etiket_sonrasi(satirlar, etiket):
    """Etiket satırının hemen altındaki (ya da 'Etiket: değer' biçiminde aynı satırdaki) değeri döndürür."""
    e = tr_kucuk(etiket)
    for i, satir in enumerate(satirlar):
        k = tr_kucuk(satir).rstrip(" :")
        if k == e:
            return satirlar[i + 1] if i + 1 < len(satirlar) else None
        if k.startswith(e) and ":" in satir[len(etiket):len(etiket) + 3]:
            return satir.split(":", 1)[1].strip() or None
    return None


def kural_docx_etiket(etiket, desen=None):
    """DOCX'te 'etiket' satırının karşılığını okur. Bulunamazsa (varsa) 'desen' tüm metinde aranır."""
    def kural(yol, secenek):
        satirlar = docx_satirlari(yol)
        deger = _etiket_sonrasi(satirlar, etiket)
        if not deger and desen:
            m = re.search(desen, "\n".join(satirlar))
            deger = m.group(0) if m else None
        return deger, ("" if deger else "Belgede '%s' satırı bulunamadı." % etiket)
    return kural


def _ilk_sayfa_metni(yol):
    sayfalar, _ = pdf_sayfa_metinleri_mupdf(yol, 1)
    if not sayfalar or len(sayfalar[0].strip()) < 40:
        try:
            sayfalar, _ = pdf_sayfa_metinleri(yol, 1)
        except Exception:
            sayfalar = []
    if (not sayfalar or len(sayfalar[0].strip()) < 40) and os.name == "nt":
        try:
            sayfalar = ocr_ile_oku(yol, [0])      # taranmış talep yazısı
        except Exception:
            sayfalar = []
    return sayfalar[0] if sayfalar else ""


_TARIH_DESENI = r"\b\d{1,2}[./]\d{1,2}[./]\d{4}\b"


def talep_coz(metin):
    """Talep yazısından (no, tarih_iso, uyarı) çıkarır.
    Belgede 'Talep Formu' ('Mal Alımı Talep Formu' dahil) yazmıyorsa hiçbir şey okunmaz.
    No: 'Sayı' satırındaki numaranın son 7 rakamı. Tarih: 'Sayı' satırından önceki (sayfanın sağ üstündeki) tarih."""
    metin = (metin or "").replace("\xa0", " ")
    if "talep formu" not in tr_kucuk(metin):
        return None, None, "Bu belge 'Talep Formu' ya da 'Mal Alımı Talep Formu' değil; talep no ve tarihi okunmadı."
    no = tarih = None
    m = re.search(r"Say[ıiİI]\s*:\s*(\S+)", metin)
    if m:
        rakamlar = re.sub(r"\D", "", m.group(1))
        no = rakamlar[-7:] if len(rakamlar) >= 7 else None
    t = re.search(_TARIH_DESENI, metin[:m.start()] if m else metin) or re.search(_TARIH_DESENI, metin)
    if t:
        try:
            tarih = tarih_ayristir(t.group(0))
        except ValueError:
            pass
    uyari = ""
    if not no:
        uyari = "'Sayı' satırı bulunamadı."
    return no, tarih, uyari


def talep_yaz(no, tarih_iso):
    """Formdaki tek hücrenin metni: '3570347 - 23.01.2026'."""
    return " - ".join(x for x in (no or "", tarih_yaz(tarih_iso)) if x)


def talep_ayristir(metin):
    """Hücre metni -> (no, tarih_iso). Tarih yoksa None; biçim bozuksa ValueError."""
    metin = (metin or "").strip()
    t = re.search(_TARIH_DESENI, metin)
    tarih = tarih_ayristir(t.group(0)) if t else None
    no = (metin[:t.start()] + metin[t.end():] if t else metin).strip(" -–\t")
    return no or None, tarih


def kural_talep_pdf(yol, secenek):
    no, tarih, uyari = talep_coz(_ilk_sayfa_metni(yol))
    if not no and not tarih:
        return None, uyari or "Talep bilgisi okunamadı."
    return talep_yaz(no, tarih), uyari


# Talep Eden Birim okunabilmesi için belgede (ilk sayfada) bunlardan biri geçmeli; boşluk sayısı önemsiz.
BIRIM_ANAHTARLARI = ("Talep Formu", "Malzeme Talep Formu", "Hizmet İhtiyaç Formu", "Malzeme İhtiyaç Formu")


def birim_coz(metin):
    """Belgenin ilk sayfasından talep eden birimi çıkarır -> (birim | None, uyarı).
    Belgede BIRIM_ANAHTARLARI'ndan biri geçmeli; yoksa hiçbir şey okunmaz (yanlış PDF seçilmiş olabilir). Birim her zaman sayfanın en üstündeki 4. satırdır:
        T.C.  /  TÜRKİYE ELEKTRİK İLETİM A.Ş. GENEL MÜDÜRLÜĞÜ  /  3.Bölge Müdürlüğü (İzmir)  /  <BİRİM>"""
    ham = (metin or "").replace("\xa0", " ")
    k = tr_kucuk(ham)
    if not any(a in " ".join(k.split()) for a in map(tr_kucuk, BIRIM_ANAHTARLARI)):
        return None, "Belgede %s ifadelerinden biri geçmiyor; birim okunmadı." % ", ".join(
            "'%s'" % a for a in BIRIM_ANAHTARLARI)
    satirlar = [s.strip() for s in ham.splitlines() if s.strip()]
    bas = 0                                            # 'T.C.' satırı 1. satır; bulunamazsa sayfanın ilk satırı
    for i, s in enumerate(satirlar[:6]):
        if re.fullmatch(r"t\.?\s?c\.?", tr_kucuk(s)):
            bas = i
            break
    if bas + 3 >= len(satirlar):
        return None, "Sayfanın üstündeki 4. satır (talep eden birim) bulunamadı."
    return " ".join(satirlar[bas + 3].split()), ""


def kural_birim_pdf(yol, secenek):
    return birim_coz(_ilk_sayfa_metni(yol))


_IS_BASLIK_ONEKI = re.compile(r"^(hizmet|malzeme)", re.I)


def _kenar_cizgileri(sayfa):
    """Sayfadaki tablo çizgileri -> ([(x, y0, y1) dikey çizgiler], [(y, x0, x1) yatay çizgiler])."""
    dikey, yatay = [], []
    try:
        cizimler = sayfa.get_drawings()
    except Exception:
        return dikey, yatay
    for d in cizimler:
        r = d["rect"]
        if r.width < 2 and r.height > 8:
            dikey.append(((r.x0 + r.x1) / 2, r.y0, r.y1))
        elif r.height < 2 and r.width > 8:
            yatay.append(((r.y0 + r.y1) / 2, r.x0, r.x1))
    return dikey, yatay


def is_adi_pdf_tablodan(yol):
    """Ihtiyaç Formu'ndaki 'HİZMETİN / MALZEMENİN ADI VE KARAKTERİSTİĞİ' başlığının hemen altındaki hücrenin metni.
    Tablo çizgilerinden başlığın sütunu ve altındaki hücre bulunur (PyMuPDF gerekir); bulunamazsa None."""
    if fitz is None:
        return None
    try:
        with fitz.open(yol) as belge:
            if belge.needs_pass or not len(belge):
                return None
            sayfa = belge[0]
            kelimeler = sayfa.get_text("words")
            dikey, yatay = _kenar_cizgileri(sayfa)
    except Exception:
        return None
    if not dikey:
        return None
    for kar in kelimeler:
        if not tr_kucuk(kar[4]).startswith("karakteristi"):
            continue
        # başlığın ilk satırı: hemen üstteki '<Hizmetin|Malzemenin> ADI VE' satırı
        ust = [w for w in kelimeler if 0 < kar[1] - w[1] < 20 and abs(w[0] - kar[0]) < 150]
        if not any(tr_kucuk(w[4]) == "adı" for w in ust) or not any(_IS_BASLIK_ONEKI.match(w[4]) for w in ust):
            continue
        cx = (kar[0] + kar[2]) / 2
        sol = max([x for x, y0, y1 in dikey if x < cx - 1 and y0 <= kar[1] + 1 and y1 >= kar[3] - 1], default=None)
        sag = min([x for x, y0, y1 in dikey if x > cx + 1 and y0 <= kar[1] + 1 and y1 >= kar[3] - 1], default=None)
        if sol is None or sag is None:
            continue
        yler = sorted(y for y, x0, x1 in yatay if x0 <= cx <= x1 and y >= kar[3] - 1)
        if not yler:
            continue
        bas = yler[0]                                    # başlık hücresinin alt kenarı
        bit = next((y for y in yler if y > bas + 3), bas + 150)
        satirlar = {}
        for w in kelimeler:
            mx, my = (w[0] + w[2]) / 2, (w[1] + w[3]) / 2
            if sol < mx < sag and bas < my < bit:
                satirlar.setdefault(round(my / 4), []).append(w)
        metin = " ".join(" ".join(x[4] for x in sorted(s, key=lambda x: x[0]))
                         for _, s in sorted(satirlar.items()))
        return " ".join(metin.split()) or None
    return None


_IS_DURDUR = re.compile(r"^\d+\s*(adet|ad\.?)?$|\d{1,2}[/.]\d{1,2}[/.]\d{4}|tarih|imza|ihtiyaç sahibi|alt birim|stok|"
                        r"adı soyadı|ünvan", re.I)
_IS_MIKTAR = re.compile(r"\s+\d+\s*(?:adet|ad\.?)\b(\s+\d+)*", re.I)
_IS_BASLIK_ARTIGI = re.compile(r"^(kodu|birimi|miktar[ıi]?|miktar|\(varsa\)|proje kodu|karakteristi.*|"
                               r"(hizmet|malzeme)\w*\s+ad[ıi]\s+ve)$", re.I)


def is_adi_pdf_metinden(metin):
    """Tablo çizgisi okunamadığında (taranmış PDF / PyMuPDF yok): 'KARAKTERİSTİĞİ' başlığından sonraki satırlar,
    başlığın yan sütun başlıkları atlanıp miktar/tarih/imza satırına gelene kadar alınır."""
    satirlar = [s.strip() for s in (metin or "").splitlines() if s.strip()]
    for i, s in enumerate(satirlar):
        if not tr_kucuk(s).startswith("karakteristi"):
            continue
        sonuc = []
        for t in satirlar[i + 1:i + 12]:
            if not sonuc and _IS_BASLIK_ARTIGI.match(t.strip()):
                continue
            if _IS_DURDUR.search(t):
                break
            sonuc.append(_IS_MIKTAR.sub("", t))        # satır içindeki '1 Adet' miktar sütunu karışmasın
        if sonuc:
            return " ".join(" ".join(sonuc).split())
    return None


_SONDAKI_PARANTEZ = re.compile(r"\s*\([^()]*\)\s*$")


def is_adi_parantez_at(ad):
    """İşin adının sonundaki parantez içi kısımları atar: 'Sepetli Vinç Kiralama (1 Adet) (Gölmarmara)' ->
    'Sepetli Vinç Kiralama'. Ortadaki parantezlere dokunmaz; ad yalnızca parantezden oluşuyorsa olduğu gibi kalır."""
    sade = ad or ""
    while _SONDAKI_PARANTEZ.search(sade):
        sade = _SONDAKI_PARANTEZ.sub("", sade)
    return sade or ad


def kural_is_adi_pdf(yol, secenek):
    metin = _ilk_sayfa_metni(yol)
    if "talep formu" not in tr_kucuk(metin):          # 'Talep Formu' ve 'Mal Alımı Talep Formu' ikisini de kapsar
        return None, "Bu belge 'Talep Formu' ya da 'Mal Alımı Talep Formu' değil; işin adı okunmadı."
    deger = is_adi_pdf_tablodan(yol) or is_adi_pdf_metinden(metin)
    if not deger:
        return None, "PDF'te 'Hizmetin / Malzemenin adı ve karakteristiği' başlığı bulunamadı."
    return is_adi_parantez_at(deger), ""


OKUMA_KURALLARI = {
    ("is", ".pdf"): kural_is_adi_pdf,
    ("is", ".docx"): kural_docx_etiket("İşin Adı"),
    ("dt", ".docx"): kural_docx_etiket("Doğrudan Temin Numarası", r"\b\d{2}DT\d{5,}\b"),
    ("talep", ".pdf"): kural_talep_pdf,
    ("birim", ".pdf"): kural_birim_pdf,
}
OKUNABILIR_TURLER = (".pdf", ".docx")


FATURA_ALANLARI = ("tarih", "no", "firma", "haric")
DOCX_BIRLIKTE = ("is", "dt")     # Doğrudan Temin Onay Belgesi (DOCX) seçilince bu alanlar birlikte doldurulur


def okunacak_alanlar(alan, yol):
    """Gözat'a basılan alanla birlikte aynı dosyadan doldurulacak alanlar."""
    if alan in DOCX_BIRLIKTE and yol.lower().endswith(".docx"):
        return list(DOCX_BIRLIKTE)
    return [alan]


def alanlari_oku(alan, yol, secenek=None):
    """okunacak_alanlar() için her alanı okur -> {alan: (deger | None, not_metni)}."""
    secenek = secenek if secenek is not None else {}
    return {a: alan_oku(a, yol, secenek) for a in okunacak_alanlar(alan, yol)}


def alan_oku(alan, yol, secenek=None):
    """Seçilen dosyadan 'alan' için değeri okur -> (deger | None, not_metni).
    Dosya türü (PDF / DOCX) desteklenmiyorsa ya da o alan o tür için kural tanımlı değilse deger None gelir."""
    secenek = secenek if secenek is not None else {}
    uzanti = os.path.splitext(yol)[1].lower()
    if uzanti not in OKUNABILIR_TURLER:
        return None, "Desteklenmeyen dosya türü (%s). PDF veya DOCX seçin." % (uzanti or "uzantısız")
    kural = OKUMA_KURALLARI.get((alan, uzanti))
    if kural is None:
        return None, "Bu alan için %s okuma kuralı henüz tanımlı değil." % uzanti.lstrip(".").upper()
    return kural(yol, secenek)


# ----------------------------------------------------------------------------
# 6) ARAYÜZ
# ----------------------------------------------------------------------------

TEMALAR = {
    "Açık": dict(bg="#fafafa", fg="#1c1c1c", alan="#ffffff", sec="#005fb8", sec_fg="#ffffff",
                 baslik="#e9e9e9", oto="#fff3c4", cizgi="#d0d0d0", dugme="#f3f3f3", zebra="#f4f7fb",
                 **{"not": "#6b6b6b"}),
    "Koyu": dict(bg="#1c1c1c", fg="#ffffff", alan="#2b2b2b", sec="#57c8ff", sec_fg="#000000",
                 baslik="#323232", oto="#5a4a12", cizgi="#454545", dugme="#2f2f2f", zebra="#242424",
                 **{"not": "#a0a0a0"}),
    "Mavi": dict(bg="#e8f0fa", fg="#12263f", alan="#ffffff", sec="#1f5fae", sec_fg="#ffffff",
                 baslik="#c9dcf2", oto="#fff0b3", cizgi="#a9c3e3", dugme="#d3e3f6", zebra="#f0f6fd",
                 **{"not": "#4a6a8f"}),
    "Yeşil": dict(bg="#e9f4ec", fg="#14301e", alan="#ffffff", sec="#2e7d4f", sec_fg="#ffffff",
                  baslik="#cde6d5", oto="#fff0b3", cizgi="#a8cdb4", dugme="#d5ebdc", zebra="#f1f9f4",
                  **{"not": "#4c7058"}),
    "Mor": dict(bg="#f0eaf8", fg="#26183d", alan="#ffffff", sec="#6b3fa0", sec_fg="#ffffff",
                baslik="#dccdef", oto="#fff0b3", cizgi="#c3add8", dugme="#e4d9f2", zebra="#f7f3fc",
                **{"not": "#6d5a8a"}),
    "Turuncu": dict(bg="#fdf1e6", fg="#3a2210", alan="#ffffff", sec="#d9660b", sec_fg="#ffffff",
                    baslik="#f6dcc2", oto="#fff0b3", cizgi="#e8c4a0", dugme="#f9e3cd", zebra="#fef8f2",
                    **{"not": "#8a6240"}),
    "Kırmızı": dict(bg="#faebeb", fg="#3a1414", alan="#ffffff", sec="#b3261e", sec_fg="#ffffff",
                    baslik="#f1cccc", oto="#fff0b3", cizgi="#e0aaaa", dugme="#f6d7d7", zebra="#fdf4f4",
                    **{"not": "#8a4f4f"}),
}


UYGULAMA_KIMLIGI = "YunusDag.FaturaTakip"


def hata_gunlugu_yolu():
    return os.path.join(os.path.dirname(_oturum_dosyasi()), "hata.log")


def hata_kaydet(metin):
    """Beklenmeyen hatanın ayrıntısını kullanıcı profilindeki hata.log dosyasına ekler (1 MB'ı aşınca sıfırlar)."""
    yol = hata_gunlugu_yolu()
    try:
        os.makedirs(os.path.dirname(yol), exist_ok=True)
        if os.path.exists(yol) and os.path.getsize(yol) > 1024 * 1024:
            os.remove(yol)
        with open(yol, "a", encoding="utf-8") as f:
            f.write("=== %s ===\n%s\n" % (datetime.now().strftime("%Y-%m-%d %H:%M:%S"), metin))
    except OSError:
        pass


def windows_kimligi_ayarla():
    """Windows'a bu sürecin kendi uygulaması olduğunu bildirir (AppUserModelID). Bu olmadan 'python fatura_takip.py'
    ile açıldığında görev çubuğu pencereyi python.exe'ye bağlar: Python ikonu görünür, alt-tab ve gruplama bozulur."""
    if os.name != "nt":
        return
    try:
        import ctypes
        ctypes.windll.shell32.SetCurrentProcessExplicitAppUserModelID(UYGULAMA_KIMLIGI)
    except Exception:
        pass


def arayuzu_calistir():
    import tkinter as tk
    from tkinter import ttk, messagebox, filedialog

    ayar = ayarlari_oku()
    kullanici = getpass.getuser()

    windows_kimligi_ayarla()                       # pencere açılmadan ÖNCE: görev çubuğu Python'u değil uygulamayı görsün
    kok = tk.Tk()
    beklenmeyen = {"acik": False}

    def beklenmeyen_hata(tur, deger, iz):
        """Düğme/menü işleyicilerinde yakalanmamış hata olursa sessiz kalmasın: günlüğe yaz, kullanıcıya göster."""
        hata_kaydet("".join(traceback.format_exception(tur, deger, iz)))
        if beklenmeyen["acik"]:
            return
        beklenmeyen["acik"] = True
        try:
            messagebox.showerror(UYGULAMA_ADI, "Beklenmeyen bir hata oluştu:\n\n%s: %s\n\nAyrıntı şuraya kaydedildi:\n%s\n\n"
                                 "İşlem tamamlanmamış olabilir; tekrar deneyin." % (tur.__name__, str(deger)[:300],
                                                                                    hata_gunlugu_yolu()))
        except Exception:
            pass
        finally:
            beklenmeyen["acik"] = False
    kok.report_callback_exception = beklenmeyen_hata
    kok.title(UYGULAMA_ADI)
    kok.geometry("1240x680")
    kok.minsize(980, 520)
    kok.withdraw()                                 # giriş yapılana kadar ana pencere görünmez
    try:
        kok.iconbitmap(IKON_YOLU)                  # ana pencerenin görev çubuğu / alt-tab ikonu
        kok.iconbitmap(default=IKON_YOLU)          # sonradan açılan tüm pencerelerde de aynı ikon
    except tk.TclError:
        pass

    try:
        db = FaturaDB(ayar["veritabani"])
    except VeritabaniHatasi as e:
        messagebox.showerror(UYGULAMA_ADI, str(e))
        kok.destroy()
        return

    # ---- tema ----------------------------------------------------------------
    stil = ttk.Style(kok)
    tema = {"ad": "Açık", "r": TEMALAR["Açık"]}
    v_tema = tk.StringVar(value=ayar["tema"] if ayar["tema"] in TEMALAR else "Açık")

    def tema_uygula(ad, kaydet=False):
        if ad not in TEMALAR:
            ad = "Açık"
        r = TEMALAR[ad]
        tema["ad"], tema["r"] = ad, r
        if sv_ttk is not None and ad in ("Açık", "Koyu"):
            sv_ttk.set_theme("light" if ad == "Açık" else "dark")
            stil.configure("Treeview", rowheight=30, font=("Segoe UI", 10))
            stil.configure("Treeview.Heading", font=("Segoe UI", 10, "bold"))
            stil.configure("Oto.TEntry", fieldbackground=r["oto"], foreground=r["fg"])
            stil.configure("Oto.TCombobox", fieldbackground=r["oto"], foreground=r["fg"])
            stil.configure("Toplam.TLabel", font=("Segoe UI", 10, "bold"))
            stil.configure("Not.TLabel", foreground=r["not"])
            stil.configure("Imza.TLabel", foreground=r["not"], font=("Segoe UI", 9))
            stil.configure("ImzaAd.TLabel", foreground=r["sec"], font=("Segoe UI Semibold", 9))
            stil.map("Toolbutton", foreground=[("selected", r["sec"])])
            kok.configure(bg=r["bg"])
            for w in kok.winfo_children():
                if isinstance(w, tk.Toplevel):
                    w.configure(bg=r["bg"])
            if kaydet:
                ayar_yaz(CONFIG_YOLU, "tema", ad)
            return
        try:
            stil.theme_use("clam")
        except tk.TclError:
            pass
        stil.configure(".", background=r["bg"], foreground=r["fg"], fieldbackground=r["alan"],
                       bordercolor=r["cizgi"], troughcolor=r["bg"], lightcolor=r["bg"], darkcolor=r["bg"])
        stil.configure("TButton", background=r["dugme"], foreground=r["fg"], padding=(8, 4))
        stil.map("TButton", background=[("active", r["baslik"]), ("pressed", r["cizgi"])])
        stil.configure("TEntry", fieldbackground=r["alan"], foreground=r["fg"], insertcolor=r["fg"])
        stil.configure("Oto.TEntry", fieldbackground=r["oto"], foreground=r["fg"], insertcolor=r["fg"])
        stil.configure("TCombobox", fieldbackground=r["alan"], background=r["dugme"], foreground=r["fg"],
                       arrowcolor=r["fg"], insertcolor=r["fg"])
        stil.configure("Oto.TCombobox", fieldbackground=r["oto"], background=r["dugme"], foreground=r["fg"],
                       arrowcolor=r["fg"], insertcolor=r["fg"])
        stil.map("TCombobox", fieldbackground=[("readonly", r["alan"])], foreground=[("readonly", r["fg"])],
                 selectbackground=[("readonly", r["alan"])], selectforeground=[("readonly", r["fg"])])
        stil.configure("Treeview", background=r["alan"], fieldbackground=r["alan"], foreground=r["fg"],
                       rowheight=24)
        stil.map("Treeview", background=[("selected", r["sec"])], foreground=[("selected", r["sec_fg"])])
        stil.configure("Treeview.Heading", background=r["baslik"], foreground=r["fg"], relief="flat",
                       padding=(4, 4))
        stil.map("Treeview.Heading", background=[("active", r["cizgi"])])
        stil.configure("Toolbutton", background=r["dugme"], foreground=r["fg"], padding=(7, 3))
        stil.map("Toolbutton", background=[("selected", r["sec"]), ("active", r["baslik"])],
                 foreground=[("selected", r["sec_fg"])])
        stil.configure("TScrollbar", background=r["dugme"], troughcolor=r["bg"], arrowcolor=r["fg"])
        stil.configure("Toplam.TLabel", font=("Segoe UI", 10, "bold"))
        stil.configure("Not.TLabel", foreground=r["not"])
        stil.configure("Imza.TLabel", foreground=r["not"], font=("Segoe UI", 9))
        stil.configure("ImzaAd.TLabel", foreground=r["sec"], font=("Segoe UI Semibold", 9))
        kok.configure(bg=r["bg"])
        kok.option_add("*TCombobox*Listbox.background", r["alan"])
        kok.option_add("*TCombobox*Listbox.foreground", r["fg"])
        kok.option_add("*TCombobox*Listbox.selectBackground", r["sec"])
        kok.option_add("*TCombobox*Listbox.selectForeground", r["sec_fg"])
        for w in kok.winfo_children():
            if isinstance(w, tk.Toplevel):
                w.configure(bg=r["bg"])
                for t in w.winfo_children():
                    if isinstance(t, tk.Text):
                        t.configure(bg=r["alan"], fg=r["fg"], insertbackground=r["fg"])
        if kaydet:
            ayar_yaz(CONFIG_YOLU, "tema", ad)

    tema_uygula(v_tema.get())

    def yeni_pencere(baslik):
        p = tk.Toplevel(kok)
        p.title(baslik)
        p.configure(bg=tema["r"]["bg"])
        p.transient(kok)
        return p

    def arka_planda(is_fonk, bitince):
        """Uzun işi (PDF okuma) ayrı iş parçacığında yapar; arayüz donmaz. bitince(değer, hata) ana iş parçacığında çağrılır."""
        kutu = {}

        def calis():
            try:
                kutu["deger"] = is_fonk()
            except Exception as ex:  # noqa
                kutu["hata"] = ex
        t = threading.Thread(target=calis, daemon=True)
        t.start()

        def bekle():
            if t.is_alive():
                kok.after(200, bekle)
                return
            bitince(kutu.get("deger"), kutu.get("hata"))
        kok.after(200, bekle)

    def hata_goster(e):
        messagebox.showerror(UYGULAMA_ADI, str(e), parent=kok)

    # ---- hesaplar: giriş, şifre, kullanıcı yönetimi ------------------------------
    oturum = {"id": None, "ad": "", "yonetici": False, "token": "", "yeniden": False}

    def diyalog(baslik, aciklama, alanlar, tamam="Tamam", onay=None, kontrol=None, ana=None):
        """Küçük modal form. alanlar: [(etiket, gizli_mi, ilk_deger)]; onay: onay kutusunun metni.
        kontrol(degerler, onay_isaretli) -> hata metni | None  (hata varsa pencere açık kalır).
        Döner: (degerler, onay_isaretli) ya da vazgeçilirse None."""
        pen = tk.Toplevel(kok)
        pen.title(baslik)
        pen.configure(bg=tema["r"]["bg"])
        pen.resizable(False, False)
        if ana is not None:
            pen.transient(ana)
        cerceve = ttk.Frame(pen, padding=16)
        cerceve.pack()
        sat = 0
        if aciklama:
            ttk.Label(cerceve, text=aciklama, wraplength=360, justify="left").grid(
                row=sat, column=0, columnspan=2, sticky="w", pady=(0, 10))
            sat += 1
        degiskenler, girisler = [], []
        for etiket, gizli, ilk in alanlar:
            ttk.Label(cerceve, text=etiket).grid(row=sat, column=0, sticky="w", pady=4)
            v = tk.StringVar(value=ilk)
            g = ttk.Entry(cerceve, textvariable=v, width=30, show="\u2022" if gizli else "")
            g.grid(row=sat, column=1, sticky="w", padx=(12, 0), pady=4)
            degiskenler.append(v)
            girisler.append(g)
            sat += 1
        onay_v = tk.BooleanVar(value=False)
        if onay:
            ttk.Checkbutton(cerceve, text=onay, variable=onay_v).grid(
                row=sat, column=0, columnspan=2, sticky="w", pady=(6, 0))
            sat += 1
        hata_v = tk.StringVar()
        ttk.Label(cerceve, textvariable=hata_v, foreground="#d93025", wraplength=360, justify="left").grid(
            row=sat, column=0, columnspan=2, sticky="w", pady=(8, 0))
        sat += 1
        sonuc = {}

        def onayla(*_):
            degerler = [v.get() for v in degiskenler]
            try:
                hata = kontrol(degerler, onay_v.get()) if kontrol else None
            except VeritabaniHatasi as ex:
                hata = str(ex)
            if hata:
                hata_v.set(hata)
                return
            sonuc["deger"] = (degerler, onay_v.get())
            pen.destroy()
        dugme = ttk.Frame(cerceve)
        dugme.grid(row=sat, column=0, columnspan=2, sticky="e", pady=(12, 0))
        ttk.Button(dugme, text=tamam, command=onayla).pack(side="left", padx=4)
        ttk.Button(dugme, text="Vazgeç", command=pen.destroy).pack(side="left", padx=4)
        pen.bind("<Return>", onayla)
        pen.bind("<Escape>", lambda e: pen.destroy())
        pen.update_idletasks()
        pen.geometry("+%d+%d" % ((pen.winfo_screenwidth() - pen.winfo_reqwidth()) // 2,
                                 (pen.winfo_screenheight() - pen.winfo_reqheight()) // 3))
        pen.wait_visibility()
        pen.grab_set()
        girisler[0].focus_set()
        kok.wait_window(pen)
        return sonuc.get("deger")

    def yeni_sifre_iste(k, zorunlu, eski_iste=False, ana=None):
        """Yeni şifre ister ve kaydeder. Döner: kaydedildi mi."""
        alanlar = ([("Mevcut şifre", True, "")] if eski_iste else []) + \
                  [("Yeni şifre", True, ""), ("Yeni şifre (tekrar)", True, "")]

        def kontrol(d, _onay):
            if d[-2] != d[-1]:
                return "Yeni şifreler aynı değil."
            db.sifre_degistir(k["id"], d[-2], eski=d[0] if eski_iste else None)
            return None
        aciklama = ("Geçici şifreyle giriş yaptınız. Devam etmek için kendi şifrenizi belirleyin (en az %d karakter)."
                    if zorunlu else "Yeni şifreniz en az %d karakter olmalı.") % SIFRE_EN_AZ
        return diyalog("Şifre Değiştir", aciklama, alanlar, "Kaydet", kontrol=kontrol, ana=ana) is not None

    def ilk_yonetici_olustur():
        bulunan = {}

        def kontrol(d, _onay):
            if d[1] != d[2]:
                return "Şifreler aynı değil."
            db.ilk_yonetici_olustur(d[0], d[1])
            bulunan["k"] = db.giris_dene(d[0], d[1])[0]
            return None
        diyalog("İlk Kurulum", "Veritabanında henüz hesap yok. Yönetici hesabını şimdi siz oluşturuyorsunuz; "
                "diğer kullanıcıları sonra Dosya > Kullanıcılar bölümünden ekleyeceksiniz.",
                [("Kullanıcı adı", False, ""), ("Şifre", True, ""), ("Şifre (tekrar)", True, "")],
                "Hesabı Oluştur", kontrol=kontrol)
        return bulunan.get("k")

    def giris_diyalogu():
        bulunan = {"token": ""}
        kilit = {"deneme": 0, "bitis": 0.0}

        def kontrol(d, hatirla):
            kalan = kilit["bitis"] - time.time()
            if kalan > 0:
                return "Çok fazla hatalı deneme. %d saniye sonra tekrar deneyin." % (kalan + 1)
            k, hata = db.giris_dene(d[0].strip(), d[1])
            if not k:
                kilit["deneme"] += 1
                if kilit["deneme"] % 5 == 0:
                    kilit["bitis"] = time.time() + 15
                return hata
            bulunan["token"] = db.oturum_olustur(k["id"]) if hatirla else ""
            oturum_yaz(db.yol, bulunan["token"])
            bulunan["k"] = k
            return None
        diyalog("Giriş", "Kullanıcı adınız ve şifrenizle giriş yapın.",
                [("Kullanıcı adı", False, ""), ("Şifre", True, "")], "Giriş",
                onay="Beni bu bilgisayarda hatırla", kontrol=kontrol)
        return bulunan.get("k"), bulunan["token"]

    def hatirlamayi_yenile(kid):
        """Şifre değişince eski anahtarlar silinir; bu bilgisayarda 'beni hatırla' açıksa yeni anahtar verilir."""
        token = db.oturum_olustur(kid)
        oturum_yaz(db.yol, token)
        return token

    def giris_yap():
        """Açılışta oturum açtırır. False dönerse uygulama kapanır."""
        try:
            k, token = None, ""
            if db.kullanici_sayisi() == 0:
                k = ilk_yonetici_olustur()
            else:
                kayitli = oturum_oku(db.yol)
                if kayitli:
                    k = db.oturum_dogrula(kayitli)
                    token = kayitli if k else ""
                if not k:
                    k, token = giris_diyalogu()
            if not k:
                return False
            if k["sifre_degismeli"]:
                if not yeni_sifre_iste(k, True):
                    return False
                if token:
                    token = hatirlamayi_yenile(k["id"])
        except VeritabaniHatasi as e:
            messagebox.showerror(UYGULAMA_ADI, str(e))
            return False
        oturum.update(id=k["id"], ad=k["kullanici_adi"], yonetici=bool(k["yonetici"]), token=token)
        return True

    if not giris_yap():
        kok.destroy()
        return False
    kullanici = oturum["ad"]

    # ---- durum ---------------------------------------------------------------
    durum = {"satirlar": [], "gorunen": [], "sirala": ("fatura_tarihi", True)}
    v_arama = tk.StringVar()
    v_firma = tk.StringVar(value="Tümü")
    v_bas = tk.StringVar()
    v_bit = tk.StringVar()
    v_toplam = tk.StringVar()
    v_yil = tk.StringVar()
    v_ay = tk.IntVar(value=0)     # 0 = Tümü, 1..12 = ay

    # ---- üst bar 1: yıl + ay sekmeleri -----------------------------------------
    ust1 = ttk.Frame(kok, padding=(8, 8, 8, 0))
    ust1.pack(fill="x")
    ttk.Label(ust1, text="Yıl:").pack(side="left")
    cb_yil = ttk.Combobox(ust1, textvariable=v_yil, width=7, state="readonly")
    cb_yil.pack(side="left", padx=(4, 10))
    ay_dugmeleri = []
    rb = ttk.Radiobutton(ust1, text="Tümü", value=0, variable=v_ay, style="Toolbutton")
    rb.pack(side="left", padx=1)
    ay_dugmeleri.append(rb)
    for k in range(1, 13):
        rb = ttk.Radiobutton(ust1, text=AY_ADLARI[k - 1], value=k, variable=v_ay, style="Toolbutton")
        rb.pack(side="left", padx=1)
        ay_dugmeleri.append(rb)

    # ---- üst bar 2: arama + filtre -------------------------------------------
    ust = ttk.Frame(kok, padding=(8, 8, 8, 2))
    ust.pack(fill="x")
    ttk.Label(ust, text="Ara:").grid(row=0, column=0, sticky="w")
    e_ara = ttk.Entry(ust, textvariable=v_arama, width=34)
    e_ara.grid(row=0, column=1, padx=(4, 12))
    ttk.Label(ust, text="Firma:").grid(row=0, column=2, sticky="w")
    cb_firma = ttk.Combobox(ust, textvariable=v_firma, width=30, state="readonly")
    cb_firma.grid(row=0, column=3, padx=(4, 12))
    ttk.Label(ust, text="Tarih (GG.AA.YYYY):").grid(row=0, column=4, sticky="w")
    ttk.Entry(ust, textvariable=v_bas, width=11).grid(row=0, column=5, padx=(4, 2))
    ttk.Label(ust, text="-").grid(row=0, column=6)
    ttk.Entry(ust, textvariable=v_bit, width=11).grid(row=0, column=7, padx=(2, 8))
    btn_temizle = ttk.Button(ust, text="Filtreleri Temizle")
    btn_temizle.grid(row=0, column=8, padx=4)
    btn_bul = ttk.Button(ust, text="\U0001F50D Fatura Bul")
    btn_bul.grid(row=0, column=9, padx=4)

    # ---- tablo ---------------------------------------------------------------
    orta = ttk.Frame(kok, padding=(8, 4))
    orta.pack(fill="both", expand=True)
    sutunlar = [
        ("fatura_tarihi", "Fatura Tarihi", 95, "center"),
        ("fatura_no", "Fatura No", 150, "w"),
        ("firma_adi", "Firma Adı", 340, "w"),
        ("yapilan_is", "İşin Adı", 280, "w"),
        ("kategori", "Kategori", 170, "w"),
        ("talep", "Talep No - Tarih", 150, "w"),
        ("talep_eden_birim", "Talep Eden Birim", 260, "w"),
        ("dt_no", "DT No", 100, "w"),
        ("tutar_haric", "Tutar (KDV Hariç)", 130, "e"),
        ("pdf_var", "PDF", 45, "center"),
        ("ek_sayisi", "Ek", 40, "center"),
    ]
    tree = ttk.Treeview(orta, columns=[s[0] for s in sutunlar], show="headings", selectmode="browse")
    for ad, baslik, genislik, hiza in sutunlar:
        tree.heading(ad, text=baslik, command=lambda a=ad: sirala_tikla(a))
        tree.column(ad, width=genislik, minwidth=genislik, anchor=hiza,
                    stretch=(ad in ("firma_adi", "yapilan_is")))
    sb = ttk.Scrollbar(orta, orient="vertical", command=tree.yview)
    sb_x = ttk.Scrollbar(orta, orient="horizontal", command=tree.xview)
    tree.configure(yscrollcommand=sb.set, xscrollcommand=sb_x.set)
    tree.grid(row=0, column=0, sticky="nsew")
    sb.grid(row=0, column=1, sticky="ns")
    sb_x.grid(row=1, column=0, sticky="ew")
    orta.rowconfigure(0, weight=1)
    orta.columnconfigure(0, weight=1)

    # ---- alt bar -------------------------------------------------------------
    alt = ttk.Frame(kok, padding=(8, 4, 8, 8))
    alt.pack(fill="x")
    btn_ekle = ttk.Button(alt, text="Yeni Kayıt")
    btn_duzenle = ttk.Button(alt, text="Düzenle")
    btn_sil = ttk.Button(alt, text="Sil")
    btn_evrak = ttk.Button(alt, text="Evraklar...")
    btn_klasor = ttk.Button(alt, text="Klasörde Göster")
    btn_ozet = ttk.Button(alt, text="Aylık Özet")
    btn_csv = ttk.Button(alt, text="Excel'e Aktar")
    btn_satinalma = ttk.Button(alt, text="Satınalmaya Aktar")
    for b in (btn_ekle, btn_duzenle, btn_sil, btn_evrak, btn_klasor, btn_ozet, btn_csv, btn_satinalma):
        b.pack(side="left", padx=3)
    ttk.Label(alt, textvariable=v_toplam, style="Toplam.TLabel").pack(side="right")

    # ---- imza şeridi ---------------------------------------------------------
    imza = ttk.Frame(kok, padding=(12, 0, 14, 6))
    imza.pack(fill="x", side="bottom")
    ttk.Label(imza, text=GELISTIRICI, style="ImzaAd.TLabel", cursor="hand2").pack(side="right")
    ttk.Label(imza, text="Developed by ", style="Imza.TLabel").pack(side="right")
    ttk.Separator(kok, orient="horizontal").pack(fill="x", side="bottom", padx=10)

    # ---- liste yenileme ------------------------------------------------------
    def filtre_degerleri():
        bas = bit = None
        try:
            if v_bas.get().strip():
                bas = tarih_ayristir(v_bas.get())
            if v_bit.get().strip():
                bit = tarih_ayristir(v_bit.get())
        except ValueError:
            return None  # yazım tamamlanmamış; filtreyi uygulama
        return bas, bit

    def anahtar(r, sutun):
        if sutun == "tutar_haric":
            return r[sutun] if r[sutun] is not None else -1
        if sutun == "pdf_var":
            return 1 if r["pdf_yolu"] else 0
        if sutun == "ek_sayisi":
            return r.get("ek_sayisi") or 0
        if sutun == "fatura_tarihi":
            return r["fatura_tarihi"] or ""
        if sutun == "talep":
            return (r.get("talep_tarihi") or "", r.get("talep_no") or "")
        return tr_kucuk(r.get(sutun) or "")

    def goster(donem=""):
        tree.delete(*tree.get_children())
        sutun, azalan = durum["sirala"]
        satirlar = sorted(durum["satirlar"], key=lambda r: anahtar(r, sutun), reverse=azalan)
        tree.tag_configure("cift", background=tema["r"]["zebra"])
        for sira, r in enumerate(satirlar):
            tree.insert("", "end", iid=str(r["id"]), tags=("cift",) if sira % 2 else (), values=(
                tarih_yaz(r["fatura_tarihi"]), r["fatura_no"], r["firma_adi"], r["yapilan_is"] or "",
                r.get("kategori") or "", talep_yaz(r.get("talep_no"), r.get("talep_tarihi")),
                r.get("talep_eden_birim") or "", r.get("dt_no") or "",
                tutar_yaz(r["tutar_haric"]), "✔" if r["pdf_yolu"] else "",
                r.get("ek_sayisi") or ""))
        durum["gorunen"] = satirlar
        adet, haric = toplamlar(satirlar)
        v_toplam.set("%s%d fatura   |   KDV Hariç: %s"
                     % (donem + "  -  " if donem else "", adet, tutar_yaz(haric, True)))
        for ad, baslik, _, _ in sutunlar:
            ok = (" ▼" if azalan else " ▲") if ad == sutun else ""
            tree.heading(ad, text=baslik + ok)

    def yenile(*_):
        if durum.get("askida"):
            return
        try:
            fv = filtre_degerleri()
            if fv is None:
                return
            tum = db.hepsi()
            # yıl listesi
            yillar = sorted({int(r["fatura_tarihi"][:4]) for r in tum} | {date.today().year}, reverse=True)
            cb_yil["values"] = ["Tümü"] + [str(y) for y in yillar]
            if v_yil.get() not in ["Tümü"] + [str(y) for y in yillar]:
                v_yil.set(str(yillar[0]))
            yil = None if v_yil.get() == "Tümü" else int(v_yil.get())
            # firma listesi
            firmalar = sorted({r["firma_adi"] for r in tum}, key=tr_kucuk)
            cb_firma["values"] = ["Tümü"] + firmalar
            if v_firma.get() not in ["Tümü"] + firmalar:
                v_firma.set("Tümü")
            # ay düğmelerindeki adetler (seçili yıla göre)
            yil_kayitlari = filtrele(tum, yil=yil)
            sayac = {k: 0 for k in range(1, 13)}
            for r in yil_kayitlari:
                sayac[int(r["fatura_tarihi"][5:7])] += 1
            ay_dugmeleri[0].configure(text="Tümü (%d)" % len(yil_kayitlari))
            for k in range(1, 13):
                ay_dugmeleri[k].configure(text="%s (%d)" % (AY_ADLARI[k - 1], sayac[k]) if sayac[k] else AY_ADLARI[k - 1])
            ay = v_ay.get() or None
            firma = "" if v_firma.get() == "Tümü" else v_firma.get()
            durum["satirlar"] = filtrele(tum, v_arama.get(), firma, fv[0], fv[1], yil, ay)
            donem = ((AY_ADLARI[ay - 1] + " ") if ay else "") + (str(yil) if yil else "Tüm yıllar")
            goster(donem)
        except VeritabaniHatasi as e:
            hata_goster(e)

    def sirala_tikla(sutun):
        eski, azalan = durum["sirala"]
        durum["sirala"] = (sutun, (not azalan) if eski == sutun else False)
        yenile()

    def fatura_bul_penceresi(*_):
        """Mercek: tüm faturalar içinde firma adı / fatura no / tutara göre arar; seçileni ana listede gösterir."""
        eski = durum.get("bul_pen")
        if eski is not None:
            try:
                if eski.winfo_exists():
                    eski.deiconify()
                    eski.lift()
                    eski.focus_force()
                    return
            except tk.TclError:
                pass
        pen = yeni_pencere("Fatura Bul")
        durum["bul_pen"] = pen
        pen.geometry("940x520")
        pen.minsize(760, 380)
        pen.bind("<Destroy>", lambda e: durum.update(bul_pen=None) if e.widget is pen else None)
        v_f, v_n, v_a, v_b = tk.StringVar(), tk.StringVar(), tk.StringVar(), tk.StringVar()
        v_bilgi = tk.StringVar()
        kayit = {"tum": [], "bulunan": {}}

        ust_k = ttk.Frame(pen, padding=(12, 12, 12, 4))
        ust_k.pack(fill="x")
        ttk.Label(ust_k, text="Firma adı:").grid(row=0, column=0, sticky="w")
        e_f = ttk.Entry(ust_k, textvariable=v_f, width=36)
        e_f.grid(row=0, column=1, padx=(6, 18), pady=3, sticky="w")
        ttk.Label(ust_k, text="Fatura no:").grid(row=0, column=2, sticky="w")
        e_n = ttk.Entry(ust_k, textvariable=v_n, width=24)
        e_n.grid(row=0, column=3, padx=(6, 0), pady=3, sticky="w")
        ttk.Label(ust_k, text="Tutar (KDV hariç):").grid(row=1, column=0, sticky="w")
        tutar_k = ttk.Frame(ust_k)
        tutar_k.grid(row=1, column=1, columnspan=3, padx=(6, 0), pady=3, sticky="w")
        e_a = ttk.Entry(tutar_k, textvariable=v_a, width=14)
        e_a.pack(side="left")
        ttk.Label(tutar_k, text="  -  ").pack(side="left")
        e_b = ttk.Entry(tutar_k, textvariable=v_b, width=14)
        e_b.pack(side="left")
        ttk.Label(tutar_k, text="   Tek tutar yazarsanız tam eşleşme, iki tutar yazarsanız aralık aranır.",
                  style="Not.TLabel").pack(side="left")
        ttk.Label(pen, textvariable=v_bilgi, style="Not.TLabel", padding=(12, 0, 12, 4)).pack(anchor="w")

        orta_k = ttk.Frame(pen, padding=(12, 0))
        orta_k.pack(fill="both", expand=True)
        kolonlar = [("tarih", "Fatura Tarihi", 100, "center"), ("no", "Fatura No", 150, "w"),
                    ("firma", "Firma Adı", 300, "w"), ("tutar", "Tutar (KDV Hariç)", 120, "e"),
                    ("is", "İşin Adı", 260, "w")]
        agac = ttk.Treeview(orta_k, columns=[k[0] for k in kolonlar], show="headings", selectmode="browse")
        for ad, baslik, gen, hiza in kolonlar:
            agac.heading(ad, text=baslik)
            agac.column(ad, width=gen, minwidth=gen // 2, anchor=hiza, stretch=(ad in ("firma", "is")))
        sb_k = ttk.Scrollbar(orta_k, orient="vertical", command=agac.yview)
        agac.configure(yscrollcommand=sb_k.set)
        agac.pack(side="left", fill="both", expand=True)
        sb_k.pack(side="right", fill="y")
        agac.tag_configure("cift", background=tema["r"]["zebra"])
        ipucu_bagla(agac)

        def ara(*_):
            hata = ""
            sinirlar = []
            for v in (v_a, v_b):
                try:
                    sinirlar.append(tutar_ayristir(v.get()))
                except ValueError:
                    sinirlar.append(None)
                    hata = "Geçersiz tutar: %s" % v.get().strip()
            en_az, en_cok = sinirlar
            if en_az is not None and not v_b.get().strip():
                en_cok = en_az                      # tek tutar: tam eşleşme
            bulunan = fatura_ara(kayit["tum"], v_f.get(), v_n.get(), en_az, en_cok)
            kayit["bulunan"] = {str(x["id"]): x for x in bulunan}
            agac.delete(*agac.get_children())
            for sira, x in enumerate(bulunan):
                agac.insert("", "end", iid=str(x["id"]), tags=("cift",) if sira % 2 else (), values=(
                    tarih_yaz(x["fatura_tarihi"]), x["fatura_no"], x["firma_adi"], tutar_yaz(x["tutar_haric"]),
                    x["yapilan_is"] or ""))
            if bulunan:
                agac.selection_set(str(bulunan[0]["id"]))
            v_bilgi.set(hata or ("%d fatura bulundu (toplam %d kayıt)" % (len(bulunan), len(kayit["tum"]))
                                 if bulunan else "Eşleşen fatura yok (toplam %d kayıt)" % len(kayit["tum"])))

        def yukle(*_):
            try:
                kayit["tum"] = db.hepsi()
            except VeritabaniHatasi as ex:
                kayit["tum"] = []
                hata_goster(ex)
            ara()

        def listede_goster(*_):
            s = agac.selection()
            if not s:
                return
            x = kayit["bulunan"][s[0]]
            durum["askida"] = True                  # tek seferde tüm filtreleri ayarla, sonra bir kez yenile
            try:
                v_arama.set("")
                v_firma.set("Tümü")
                v_bas.set("")
                v_bit.set("")
                v_yil.set(x["fatura_tarihi"][:4])
                v_ay.set(int(x["fatura_tarihi"][5:7]))
            finally:
                durum["askida"] = False
            yenile()
            if tree.exists(s[0]):
                tree.selection_set(s[0])
                tree.focus(s[0])
                tree.see(s[0])
            pen.destroy()
            kok.lift()
            tree.focus_set()

        def temizle_alanlar():
            for v in (v_f, v_n, v_a, v_b):
                v.set("")
            e_f.focus_set()

        for v in (v_f, v_n, v_a, v_b):
            v.trace_add("write", ara)
        alt_k = ttk.Frame(pen, padding=(12, 8, 12, 12))
        alt_k.pack(fill="x")
        ttk.Button(alt_k, text="Listede Göster", command=listede_goster).pack(side="left", padx=3)
        ttk.Button(alt_k, text="Alanları Temizle", command=temizle_alanlar).pack(side="left", padx=3)
        ttk.Button(alt_k, text="Yenile", command=yukle).pack(side="left", padx=3)
        ttk.Button(alt_k, text="Kapat", command=pen.destroy).pack(side="right", padx=3)
        ttk.Label(alt_k, text="Çift tıklama / Enter: ana listede göster", style="Not.TLabel").pack(side="right", padx=12)
        agac.bind("<Double-1>", listede_goster)
        agac.bind("<Return>", listede_goster)
        for e in (e_f, e_n, e_a, e_b):
            e.bind("<Return>", listede_goster)
            e.bind("<Down>", lambda ev: (agac.focus_set(), "break")[1])
        pen.bind("<Escape>", lambda e: pen.destroy())
        yukle()
        e_f.focus_set()

    def temizle():
        v_arama.set("")
        v_firma.set("Tümü")
        v_bas.set("")
        v_bit.set("")
        v_ay.set(0)
        yenile()

    def ipucu_bagla(agac, gecikme=500):
        """Tabloda bir hücrenin üstünde bekleyince, sütuna sığmayan metnin tamamını balon olarak gösterir."""
        import tkinter.font as tkfont
        d = {"is": None, "pen": None, "hucre": None}

        def kapat(*_):
            if d["is"]:
                agac.after_cancel(d["is"])
                d["is"] = None
            if d["pen"]:
                d["pen"].destroy()
                d["pen"] = None

        def goster(satir, sutun, x, y):
            d["is"] = None
            try:
                metin = str(agac.set(satir, sutun))
                ad = stil.lookup("Treeview", "font") or "TkDefaultFont"
                sig = agac.column(sutun, "width") - 14          # hücre kenar boşluğu
                if not metin or tkfont.Font(font=ad).measure(metin) <= sig:
                    return                                      # sığıyor, balon gerekmez
                r = tema["r"]
                pen = tk.Toplevel(agac)
                pen.wm_overrideredirect(True)
                pen.wm_geometry("+%d+%d" % (x + 14, y + 18))
                tk.Label(pen, text=metin, bg=r["alan"], fg=r["fg"], justify="left", wraplength=520,
                         padx=8, pady=4, relief="solid", borderwidth=1, font=ad).pack()
                d["pen"] = pen
            except tk.TclError:
                pass

        def hareket(e):
            satir, sutun = agac.identify_row(e.y), agac.identify_column(e.x)
            if not satir or not sutun:
                kapat()
                d["hucre"] = None
                return
            hucre = (satir, sutun)
            if hucre == d["hucre"]:
                return                                          # aynı hücrede: balon açıksa kalsın
            kapat()
            d["hucre"] = hucre
            d["is"] = agac.after(gecikme, goster, satir, agac["columns"][int(sutun[1:]) - 1],
                                 e.x_root, e.y_root)

        def ayril(*_):
            kapat()
            d["hucre"] = None

        agac.bind("<Motion>", hareket, add="+")
        for olay in ("<Leave>", "<ButtonPress>", "<MouseWheel>", "<Key>"):
            agac.bind(olay, ayril, add="+")

    def secili_id():
        s = tree.selection()
        if not s:
            messagebox.showinfo(UYGULAMA_ADI, "Önce listeden bir fatura seçin.", parent=kok)
            return None
        return int(s[0])

    # ---- metin gösterici (okunan ham metin) ----------------------------------
    def metin_goster(baslik, metin):
        pen = yeni_pencere(baslik)
        pen.geometry("760x560")
        t = tk.Text(pen, wrap="word", bg=tema["r"]["alan"], fg=tema["r"]["fg"], insertbackground=tema["r"]["fg"],
                    font=("Consolas", 10))
        t.insert("1.0", metin or "(boş)")
        t.configure(state="normal")
        sb2 = ttk.Scrollbar(pen, command=t.yview)
        t.configure(yscrollcommand=sb2.set)
        sb2.pack(side="right", fill="y")
        t.pack(fill="both", expand=True, padx=6, pady=6)
        pen.bind("<Escape>", lambda e: pen.destroy())

    def durum_metni(sonuc, bulunan):
        if sonuc["yontem"] == "metin":
            kaynak = "PDF'in yazı katmanından okundu"
        elif sonuc["yontem"] == "ocr":
            kaynak = "Tarama Windows OCR ile okundu (hata payı daha yüksek olabilir)"
        else:
            return "PDF'den okunamadı. %s Bilgileri elle girebilirsiniz." % (sonuc["hata"] or "")
        if bulunan == 0:
            return "%s ama fatura bilgisi ayıklanamadı. Bilgileri elle girin." % kaynak
        return "%s. %d/4 alan bulundu; sarı alanları mutlaka kontrol edin. (Yapılan İş elle girilir.)" % (kaynak, bulunan)

    # ---- fatura formu --------------------------------------------------------
    def form_ac(fid=None):
        mevcut = db.getir(fid) if fid else None
        if fid and not mevcut:
            messagebox.showwarning(UYGULAMA_ADI, "Kayıt bulunamadı (başkası silmiş olabilir).", parent=kok)
            yenile()
            return
        pen = yeni_pencere("Kayıt Düzenle" if fid else "Yeni Kayıt")
        pen.grab_set()
        pen.resizable(False, False)

        v = {
            "tarih": tk.StringVar(value=tarih_yaz(mevcut["fatura_tarihi"]) if mevcut else date.today().strftime("%d.%m.%Y")),
            "no": tk.StringVar(value=mevcut["fatura_no"] if mevcut else ""),
            "firma": tk.StringVar(value=mevcut["firma_adi"] if mevcut else ""),
            "is": tk.StringVar(value=(mevcut["yapilan_is"] or "") if mevcut else ""),
            "kategori": tk.StringVar(value=(mevcut.get("kategori") or "") if mevcut else ""),
            "haric": tk.StringVar(value=tutar_yaz(mevcut["tutar_haric"]) if mevcut else ""),
            "talep": tk.StringVar(value=talep_yaz(mevcut.get("talep_no"), mevcut.get("talep_tarihi")) if mevcut else ""),
            "birim": tk.StringVar(value=(mevcut.get("talep_eden_birim") or "") if mevcut else ""),
            "dt": tk.StringVar(value=(mevcut.get("dt_no") or "") if mevcut else ""),
            "pdf": tk.StringVar(value=(mevcut["pdf_yolu"] or "") if mevcut else ""),
            "vkn": tk.StringVar(value=(mevcut.get("vkn") or "") if mevcut else ""),
            "sayfa": tk.StringVar(value=""),
        }
        v_durum = tk.StringVar(value="")
        son_metin = {"metin": ""}
        son_klasor = {"yol": ""}
        okuma_secenek = {"ayar": ayar}

        cerceve = ttk.Frame(pen, padding=14)
        cerceve.pack()
        # (etiket, anahtar, Gözat var mı)
        etiketler = [("Kategori", "kategori", False), ("Fatura Tarihi (GG.AA.YYYY)", "tarih", True),
                     ("Fatura No", "no", True), ("Firma Adı", "firma", True), ("İşin Adı", "is", True),
                     ("Tutar (KDV Hariç)", "haric", True), ("Talep No - Talep Tarihi", "talep", True),
                     ("Talep Eden Birim", "birim", True), ("DT Numarası", "dt", True)]
        girisler = {}
        gozat_dugmeleri = []

        def normal_stil(k):
            girisler[k].configure(style="TCombobox" if k in ("firma", "kategori") else "TEntry")

        def dosya_sec(baslik):
            ilk = son_klasor["yol"]
            if not ilk and v["pdf"].get() and os.path.isdir(os.path.dirname(v["pdf"].get())):
                ilk = os.path.dirname(v["pdf"].get())
            if not ilk and ayar["tarama_klasoru"] and os.path.isdir(ayar["tarama_klasoru"]):
                ilk = ayar["tarama_klasoru"]
            yol = filedialog.askopenfilename(
                parent=pen, title=baslik, initialdir=ilk or None,
                filetypes=[("PDF ve Word", "*.pdf *.docx"), ("PDF dosyaları", "*.pdf"),
                           ("Word belgeleri", "*.docx"), ("Tüm dosyalar", "*.*")])
            if yol:
                yol = os.path.normpath(yol)
                son_klasor["yol"] = os.path.dirname(yol)
            return yol

        def alan_gozat(k, et):
            """Satırın kendi Gözat'ı: seçilen PDF / DOCX dosyasından bu satırın değerini okur.
            İşin Adı / DT Numarası için DOCX seçilirse ikisi birlikte doldurulur."""
            if k in FATURA_ALANLARI:
                pdf_gozat(oku=True)   # fatura bilgileri tek bir fatura PDF'inden birlikte okunur
                return
            yol = dosya_sec("%s için dosya seçin" % et)
            if not yol:
                return
            okuma_secenek["sayfa"] = v["sayfa"].get().strip()
            adlar = {kk: ee for ee, kk, _ in etiketler}
            v_durum.set("%s okunuyor... (taranmış PDF'lerde biraz sürebilir)" % et)
            for b in gozat_dugmeleri:
                b.state(["disabled"])

            def bitti(sonuc, hata):
                try:
                    if not pen.winfo_exists():
                        return
                    for b in gozat_dugmeleri:
                        b.state(["!disabled"])
                except tk.TclError:
                    return
                if hata or not sonuc:
                    v_durum.set("%s okunamadı: %s" % (et, hata or "bilinmeyen hata"))
                    return
                dolan, eksik = [], []
                for alan, (deger, nt) in sonuc.items():
                    if deger:
                        v[alan].set(deger)
                        girisler[alan].configure(style="Oto.TEntry")
                        dolan.append(adlar[alan])
                    else:
                        eksik.append("%s: %s" % (adlar[alan], nt))
                mesaj = ""
                if dolan:
                    mesaj = "%s okundu; sarı alanları kontrol edin." % ", ".join(dolan)
                v_durum.set((mesaj + " " + " ".join(eksik)).strip() if eksik or mesaj else "")
            arka_planda(lambda: alanlari_oku(k, yol, okuma_secenek), bitti)

        for i, (et, k, gozat_var) in enumerate(etiketler):
            ttk.Label(cerceve, text=et).grid(row=i, column=0, sticky="w", pady=4)
            if k == "kategori":
                g = ttk.Combobox(cerceve, textvariable=v[k], width=52, values=KATEGORILER, state="readonly")
                g.bind("<<ComboboxSelected>>", lambda e, kk=k: normal_stil(kk))
            elif k == "firma":
                g = ttk.Combobox(cerceve, textvariable=v[k], width=52, values=db.firmalar())
                g.bind("<<ComboboxSelected>>", lambda e, kk=k: normal_stil(kk))
            else:
                g = ttk.Entry(cerceve, textvariable=v[k], width={"haric": 20, "dt": 20, "talep": 30}.get(k, 55))
            g.bind("<Key>", lambda e, kk=k: normal_stil(kk))
            g.grid(row=i, column=1, sticky="w", padx=(10, 0), pady=4)
            girisler[k] = g
            if gozat_var:
                b = ttk.Button(cerceve, text="Gözat...", command=lambda kk=k, ee=et: alan_gozat(kk, ee))
                b.grid(row=i, column=2, padx=(6, 0))
                gozat_dugmeleri.append(b)

        s_pdf = len(etiketler)
        ttk.Label(cerceve, text="PDF Dosyası").grid(row=s_pdf, column=0, sticky="w", pady=4)
        ttk.Entry(cerceve, textvariable=v["pdf"], width=55).grid(row=s_pdf, column=1, sticky="w", padx=(10, 0), pady=4)

        # ---- fatura PDF'inden tüm fatura alanlarını bir seferde doldurma (yeni kayıtta dosya seçilince)
        def pdf_oku():
            yol = v["pdf"].get().strip()
            if not yol or not os.path.exists(yol):
                return
            sayfa_spec = v["sayfa"].get().strip()
            v_durum.set("PDF okunuyor... (resim olarak taranmış çok sayfalı belgelerde 30-60 saniye sürebilir)")
            for b in gozat_dugmeleri + [btn_pdf_gozat]:
                b.state(["disabled"])

            def bitti(sonuc, hata):
                try:
                    if not pen.winfo_exists():
                        return
                    for b in gozat_dugmeleri + [btn_pdf_gozat]:
                        b.state(["!disabled"])
                except tk.TclError:
                    return
                if hata or not sonuc:
                    v_durum.set("PDF okunurken hata oluştu: %s. Bilgileri elle girebilirsiniz." % (hata,))
                    return
                son_metin["metin"] = sonuc["metin"]
                if sonuc.get("sayfa") and (sonuc.get("sayfa_sayisi") or 0) > 1:
                    v["sayfa"].set(sonuc["sayfa"])      # bulunan fatura sayfası kutuda görünür (elle değiştirilebilir)
                a = sonuc["alanlar"]
                vkn = a.get("vkn", "")
                if vkn:
                    v["vkn"].set(vkn)
                    bilinen = None
                    try:
                        bilinen = db.vkn_firma(vkn)
                    except VeritabaniHatasi:
                        pass
                    if bilinen:
                        a["firma_adi"] = bilinen     # aynı firma hep aynı yazımla gelsin
                bulunan = 0
                eslesme = [("fatura_tarihi", "tarih", tarih_yaz), ("fatura_no", "no", str),
                           ("firma_adi", "firma", str), ("tutar_haric", "haric", tutar_yaz)]
                for alan, k, bicim in eslesme:
                    if alan in a and a[alan] not in (None, ""):
                        v[k].set(bicim(a[alan]))
                        girisler[k].configure(style="Oto.TCombobox" if k == "firma" else "Oto.TEntry")
                        bulunan += 1
                v_durum.set((durum_metni(sonuc, bulunan) + "  " + sonuc.get("sayfa_bilgisi", "")).strip())
            arka_planda(lambda: pdf_bilgi_cikar(yol, ayar, sayfa_spec or None), bitti)

        def pdf_gozat(oku=None):
            """Fatura PDF'ini seçer. oku=None: yalnızca yeni kayıtta seçer seçmez fatura alanlarını okur;
            oku=True: (fatura alanlarının kendi Gözat'ı) kayıt düzenlenirken de okur."""
            yol = dosya_sec("Fatura PDF'ini seçin")
            if not yol:
                return
            if not yol.lower().endswith(".pdf"):
                v_durum.set("Fatura bilgileri yalnızca PDF'ten okunur. Fatura PDF'ini seçin.")
                return
            v["pdf"].set(yol)
            v["sayfa"].set("")                    # yeni dosyada sayfa otomatik bulunur
            if oku or (oku is None and not fid):
                pdf_oku()

        btn_pdf_gozat = ttk.Button(cerceve, text="Gözat...", command=pdf_gozat)
        btn_pdf_gozat.grid(row=s_pdf, column=2, padx=(6, 0))

        sayfa_cerceve = ttk.Frame(cerceve)
        sayfa_cerceve.grid(row=s_pdf + 1, column=1, columnspan=2, sticky="w", padx=(10, 0))
        ttk.Label(sayfa_cerceve, text="(Dosya kopyalanmaz, yalnızca konumu kaydedilir.)",
                  style="Not.TLabel").pack(side="left")
        ttk.Label(sayfa_cerceve, text="     Fatura sayfası:").pack(side="left")
        ttk.Entry(sayfa_cerceve, textvariable=v["sayfa"], width=7).pack(side="left", padx=4)
        ttk.Label(sayfa_cerceve, text="(boş = otomatik bul; örn. 7 ya da 7-8)", style="Not.TLabel").pack(side="left")
        ttk.Label(cerceve, textvariable=v_durum, style="Not.TLabel", wraplength=600, justify="left").grid(
            row=s_pdf + 2, column=0, columnspan=3, sticky="w", pady=(6, 0))
        ttk.Button(cerceve, text="Okunan metni göster",
                   command=lambda: metin_goster("PDF'den okunan ham metin", son_metin["metin"])).grid(
            row=s_pdf + 3, column=0, columnspan=2, sticky="w", pady=(4, 0))

        def kaydet():
            try:
                if v["kategori"].get() not in KATEGORILER:
                    raise ValueError("Kategori seçin.")
                try:
                    talep_no, talep_tarihi = talep_ayristir(v["talep"].get())
                except ValueError:
                    raise ValueError("Talep hücresi '3570347 - 23.01.2026' biçiminde olmalı (tarih GG.AA.YYYY).")
                if not v["firma"].get().strip():
                    raise ValueError("Firma adı boş olamaz.")
                if not v["no"].get().strip():
                    raise ValueError("Fatura numarası boş olamaz.")
                veri = {
                    "fatura_tarihi": tarih_ayristir(v["tarih"].get()),
                    "fatura_no": v["no"].get().strip(),
                    "firma_adi": v["firma"].get().strip(),
                    "yapilan_is": v["is"].get().strip(),
                    "kategori": v["kategori"].get(),
                    "talep_no": talep_no, "talep_tarihi": talep_tarihi,
                    "dt_no": v["dt"].get().strip() or None,
                    "talep_eden_birim": v["birim"].get().strip() or None,
                    "tutar_haric": tutar_ayristir(v["haric"].get()),
                    "pdf_yolu": v["pdf"].get().strip(),
                    "vkn": v["vkn"].get().strip(),
                }
                if veri["tutar_haric"] is None:
                    raise ValueError("KDV hariç tutarı girin.")
                if veri["pdf_yolu"] and not os.path.exists(veri["pdf_yolu"]):
                    if not messagebox.askyesno(UYGULAMA_ADI, "Seçilen PDF dosyası bulunamıyor.\n"
                                               "Yine de bu yolla kaydedilsin mi?", parent=pen):
                        return
                if fid:
                    db.guncelle(fid, veri, kullanici)
                else:
                    db.ekle(veri, kullanici)
            except ValueError as ex:
                messagebox.showwarning(UYGULAMA_ADI, str(ex), parent=pen)
                return
            except VeritabaniHatasi as ex:
                messagebox.showerror(UYGULAMA_ADI, str(ex), parent=pen)
                return
            pen.destroy()
            yenile()

        dugme = ttk.Frame(cerceve)
        dugme.grid(row=13, column=0, columnspan=3, pady=(14, 0), sticky="e")
        ttk.Button(dugme, text="Kaydet", command=kaydet).pack(side="left", padx=4)
        ttk.Button(dugme, text="Vazgeç", command=pen.destroy).pack(side="left", padx=4)
        if mevcut:
            ttk.Label(cerceve, style="Not.TLabel", text="Kaydeden: %s  (%s)%s" % (
                mevcut["kaydeden"] or "-", mevcut["kayit_tarihi"] or "",
                ("   |   Son düzenleyen: %s (%s)" % (mevcut["son_duzenleyen"], mevcut["duzenleme_tarihi"]))
                if mevcut["son_duzenleyen"] else "")).grid(row=14, column=0, columnspan=3, sticky="w", pady=(10, 0))
        girisler["tarih"].focus_set()
        pen.bind("<Escape>", lambda e: pen.destroy())

    # ---- PDF okuma testi (menüden) ---------------------------------------------
    def okuma_testi():
        yol = filedialog.askopenfilename(parent=kok, title="Okuma testi için bir fatura PDF'i seçin",
                                         initialdir=ayar["tarama_klasoru"] or None,
                                         filetypes=[("PDF dosyaları", "*.pdf")])
        if not yol:
            return
        bilgi = yeni_pencere("PDF okunuyor...")
        ttk.Label(bilgi, text="Okunuyor, lütfen bekleyin...", padding=20).pack()

        def bitti(sonuc, hata):
            try:
                bilgi.destroy()
            except tk.TclError:
                pass
            if hata or not sonuc:
                hata_goster("Okuma sırasında hata: %s" % hata)
                return
            a = sonuc["alanlar"]
            satirlar = ["YÖNTEM: " + {"metin": "PDF yazı katmanı", "ocr": "Windows OCR"}.get(sonuc["yontem"], "okunamadı")]
            if sonuc.get("sayfa_bilgisi"):
                satirlar.append("SAYFA: " + sonuc["sayfa_bilgisi"])
            if sonuc["hata"]:
                satirlar.append("UYARI: " + sonuc["hata"])
            satirlar += ["", "BULUNAN ALANLAR:"]
            for ad, anahtar_, bicim in [("Fatura Tarihi", "fatura_tarihi", tarih_yaz), ("Fatura No", "fatura_no", str),
                                        ("Firma Adı", "firma_adi", str), ("VKN", "vkn", str),
                                        ("Tutar (KDV Hariç)", "tutar_haric", tutar_yaz)]:
                satirlar.append("  %-18s %s" % (ad, bicim(a[anahtar_]) if anahtar_ in a else "- bulunamadı -"))
            satirlar += ["", "---------- OKUNAN HAM METİN ----------", sonuc["metin"] or "(metin yok)"]
            metin_goster("PDF Okuma Testi", "\n".join(satirlar))
        arka_planda(lambda: pdf_bilgi_cikar(yol, ayar), bitti)

    # ---- aylık özet penceresi ------------------------------------------------
    def ozet_penceresi():
        tum = db.hepsi()
        yillar = sorted({int(r["fatura_tarihi"][:4]) for r in tum} | {date.today().year}, reverse=True)
        pen = yeni_pencere("Aylık Özet")
        v_y = tk.StringVar(value=v_yil.get() if v_yil.get() != "Tümü" else str(yillar[0]))
        ust_ = ttk.Frame(pen, padding=8)
        ust_.pack(fill="x")
        ttk.Label(ust_, text="Yıl:").pack(side="left")
        cb = ttk.Combobox(ust_, textvariable=v_y, values=[str(y) for y in yillar], width=8, state="readonly")
        cb.pack(side="left", padx=6)
        kolonlar = [("ay", "Ay", 110, "w"), ("adet", "Fatura Adedi", 100, "center"),
                    ("haric", "KDV Hariç Toplam", 150, "e")]
        agac = ttk.Treeview(pen, columns=[k[0] for k in kolonlar], show="headings", height=14)
        for ad, baslik, gen, hiza in kolonlar:
            agac.heading(ad, text=baslik)
            agac.column(ad, width=gen, anchor=hiza)
        agac.tag_configure("toplam", font=("Segoe UI", 10, "bold"))
        agac.pack(fill="both", expand=True, padx=8, pady=4)
        veri_ = {"ozet": []}

        def doldur(*_):
            agac.delete(*agac.get_children())
            try:
                veri_["ozet"] = aylik_ozet(db.hepsi(), int(v_y.get()))
            except VeritabaniHatasi as e:
                hata_goster(e)
                return
            for x in veri_["ozet"]:
                agac.insert("", "end", tags=("toplam",) if x["ay"] == "TOPLAM" else (), values=(
                    x["ay"], x["adet"], tutar_yaz(x["haric"], True)))

        def aktar():
            tipler = [("Excel çalışma kitabı", "*.xlsx"), ("CSV (Excel'de açılır)", "*.csv")] if openpyxl else \
                     [("CSV (Excel'de açılır)", "*.csv")]
            yol = filedialog.asksaveasfilename(parent=pen, title="Aylık özeti kaydet", defaultextension=tipler[0][1],
                                               initialfile="Aylik_Ozet_%s" % v_y.get(), filetypes=tipler)
            if not yol:
                return
            try:
                if yol.lower().endswith(".csv") or openpyxl is None:
                    ozet_csv_yaz(yol, veri_["ozet"])
                else:
                    ozet_xlsx_yaz(yol, veri_["ozet"], v_y.get())
                messagebox.showinfo(UYGULAMA_ADI, "Aylık özet kaydedildi:\n%s" % yol, parent=pen)
            except PermissionError:
                messagebox.showerror(UYGULAMA_ADI, "Dosya yazılamadı. Excel'de açık olabilir; kapatıp tekrar deneyin.", parent=pen)
            except OSError as e:
                hata_goster(e)
        alt_ = ttk.Frame(pen, padding=8)
        alt_.pack(fill="x")
        ttk.Button(alt_, text="Excel'e Aktar", command=aktar).pack(side="left")
        ttk.Button(alt_, text="Kapat", command=pen.destroy).pack(side="right")
        cb.bind("<<ComboboxSelected>>", doldur)
        pen.bind("<Escape>", lambda e: pen.destroy())
        doldur()

    # ---- düğme işlevleri -----------------------------------------------------
    def duzenle(*_):
        fid = secili_id()
        if fid:
            form_ac(fid)

    def sil():
        if not oturum["yonetici"]:
            messagebox.showinfo(UYGULAMA_ADI, "Fatura silme yetkisi yalnızca yöneticidedir.", parent=kok)
            return
        fid = secili_id()
        if not fid:
            return
        k = db.getir(fid)
        if not k:
            yenile()
            return
        if messagebox.askyesno(UYGULAMA_ADI, "Bu fatura kalıcı olarak silinecek:\n\n%s  /  %s\n\nEmin misiniz?\n"
                               "(PDF dosyası silinmez, yalnızca kayıt silinir.)" % (k["firma_adi"], k["fatura_no"]),
                               icon="warning", parent=kok):
            try:
                db.sil(fid)
            except VeritabaniHatasi as e:
                hata_goster(e)
            yenile()

    # ---- fatura evrakları (fatura PDF'i + talep, yazı vb. ekler) ------------------
    EK_TURLERI = ["Talep", "Yazı", "Teklif", "İrsaliye", "Sözleşme", "Fatura Ön İzleme", "Diğer"]

    def evrak_penceresi(fid=None):
        fid = fid or secili_id()
        if not fid:
            return
        fatura = db.getir(fid)
        if not fatura:
            messagebox.showwarning(UYGULAMA_ADI, "Kayıt bulunamadı (başkası silmiş olabilir).", parent=kok)
            yenile()
            return
        pen = yeni_pencere("Evraklar  -  %s  /  %s" % (fatura["firma_adi"], fatura["fatura_no"]))
        pen.geometry("860x420")
        pen.grab_set()
        satirlar = {}                       # iid -> (yol, ek_id veya None)

        ttk.Label(pen, text="%s   |   %s   |   %s" % (fatura["firma_adi"], fatura["fatura_no"],
                                                       tarih_yaz(fatura["fatura_tarihi"])),
                  padding=(10, 8, 10, 2)).pack(anchor="w")
        orta2 = ttk.Frame(pen, padding=(10, 4))
        orta2.pack(fill="both", expand=True)
        kol = [("tur", "Tür", 120), ("dosya", "Dosya", 300), ("aciklama", "Açıklama", 190),
               ("ekleyen", "Ekleyen", 100), ("tarih", "Tarih", 90), ("durum", "Durum", 70)]
        agac = ttk.Treeview(orta2, columns=[k[0] for k in kol], show="headings", selectmode="browse")
        ipucu_bagla(agac)
        for ad, baslik, gen in kol:
            agac.heading(ad, text=baslik)
            agac.column(ad, width=gen, anchor="w", stretch=(ad in ("dosya", "aciklama")))
        sb3 = ttk.Scrollbar(orta2, orient="vertical", command=agac.yview)
        agac.configure(yscrollcommand=sb3.set)
        agac.pack(side="left", fill="both", expand=True)
        sb3.pack(side="right", fill="y")

        def doldur():
            agac.delete(*agac.get_children())
            satirlar.clear()
            kayitlar = [("f", "Fatura", fatura["pdf_yolu"], "", "", "", None)]
            try:
                for e in db.ekler(fid):
                    kayitlar.append(("e%d" % e["id"], e["tur"] or "", e["dosya_yolu"], e["aciklama"] or "",
                                     e["ekleyen"] or "", tarih_yaz((e["ekleme_tarihi"] or "")[:10]), e["id"]))
            except VeritabaniHatasi as ex:
                hata_goster(ex)
            for iid, tur, yol, acik, ekleyen, tarih, ek_id in kayitlar:
                if iid == "f" and not yol:
                    continue
                satirlar[iid] = (yol, ek_id)
                var = os.path.exists(yol)
                agac.insert("", "end", iid=iid, values=(tur, os.path.basename(yol) or yol, acik, ekleyen, tarih,
                                                        "var" if var else "YOK"))

        def secili():
            s = agac.selection()
            if not s:
                messagebox.showinfo(UYGULAMA_ADI, "Önce listeden bir evrak seçin.", parent=pen)
                return None
            return s[0]

        def ac(*_):
            iid = secili()
            if iid:
                try:
                    pdf_ac(satirlar[iid][0])
                except (OSError, FileNotFoundError) as ex:
                    messagebox.showerror(UYGULAMA_ADI, str(ex), parent=pen)

        def tumunu_ac():
            if not satirlar:
                return
            if len(satirlar) > 5 and not messagebox.askyesno(
                    UYGULAMA_ADI, "%d dosya birden açılacak. Devam edilsin mi?" % len(satirlar), parent=pen):
                return
            eksik = []
            for yol, _ in satirlar.values():
                try:
                    pdf_ac(yol)
                except (OSError, FileNotFoundError):
                    eksik.append(yol)
            if eksik:
                messagebox.showwarning(UYGULAMA_ADI, "Şu dosyalar açılamadı:\n\n" + "\n".join(eksik), parent=pen)

        def klasorde():
            iid = secili()
            if iid:
                try:
                    klasorde_goster(satirlar[iid][0])
                except (OSError, FileNotFoundError) as ex:
                    messagebox.showerror(UYGULAMA_ADI, str(ex), parent=pen)

        v_tur = tk.StringVar(value=EK_TURLERI[0])
        v_acik = tk.StringVar()

        def dosya_ekle():
            klasor = os.path.dirname(fatura["pdf_yolu"]) if fatura["pdf_yolu"] else ""
            ilk = klasor if os.path.isdir(klasor) else (ayar["tarama_klasoru"] or None)
            yollar = filedialog.askopenfilenames(parent=pen, title="Faturaya eklenecek evrakları seçin",
                                                 initialdir=ilk, filetypes=[("Tüm dosyalar", "*.*")])
            if not yollar:
                return
            hatalar = []
            for y in yollar:
                try:
                    db.ek_ekle(fid, v_tur.get().strip() or "Diğer", os.path.normpath(y), v_acik.get().strip(), kullanici)
                except VeritabaniHatasi as ex:
                    hatalar.append(str(ex))
            v_acik.set("")
            doldur()
            yenile()
            if hatalar:
                messagebox.showwarning(UYGULAMA_ADI, "\n\n".join(hatalar), parent=pen)

        def kaldir():
            iid = secili()
            if not iid:
                return
            if satirlar[iid][1] is None:
                messagebox.showinfo(UYGULAMA_ADI, "Fatura PDF'i buradan kaldırılamaz.\n"
                                    "Değiştirmek için faturayı düzenleyin.", parent=pen)
                return
            if messagebox.askyesno(UYGULAMA_ADI, "Bu evrak faturadan çıkarılsın mı?\n(Dosyanın kendisi silinmez.)",
                                   parent=pen):
                try:
                    db.ek_sil(satirlar[iid][1])
                except VeritabaniHatasi as ex:
                    hata_goster(ex)
                doldur()
                yenile()

        ekle_bar = ttk.Frame(pen, padding=(10, 4))
        ekle_bar.pack(fill="x")
        ttk.Label(ekle_bar, text="Tür:").pack(side="left")
        ttk.Combobox(ekle_bar, textvariable=v_tur, values=EK_TURLERI, width=16).pack(side="left", padx=(4, 10))
        ttk.Label(ekle_bar, text="Açıklama:").pack(side="left")
        ttk.Entry(ekle_bar, textvariable=v_acik, width=30).pack(side="left", padx=(4, 10))
        ttk.Button(ekle_bar, text="Dosya Ekle...", command=dosya_ekle).pack(side="left")
        alt2 = ttk.Frame(pen, padding=(10, 4, 10, 10))
        alt2.pack(fill="x")
        for metin, komut in (("Aç", ac), ("Tümünü Aç", tumunu_ac), ("Klasörde Göster", klasorde),
                             ("Evrakı Kaldır", kaldir)):
            ttk.Button(alt2, text=metin, command=komut).pack(side="left", padx=3)
        ttk.Button(alt2, text="Kapat", command=pen.destroy).pack(side="right")
        ttk.Label(pen, text="Dosyalar kopyalanmaz, yalnızca konumları kaydedilir. Bu evraklar PDF okumaya katılmaz.",
                  style="Not.TLabel", padding=(10, 0, 10, 8)).pack(anchor="w")
        agac.bind("<Double-1>", ac)
        agac.bind("<Return>", ac)
        pen.bind("<Escape>", lambda e: pen.destroy())
        doldur()
        if agac.get_children():
            agac.selection_set(agac.get_children()[0])

    def klasor_btn():
        fid = secili_id()
        if not fid:
            return
        k = db.getir(fid)
        try:
            if not k or not k["pdf_yolu"]:
                raise FileNotFoundError("Bu faturaya bağlı bir PDF yok.")
            klasorde_goster(k["pdf_yolu"])
        except (OSError, FileNotFoundError) as e:
            hata_goster(e)

    def donem_adi():
        ay = v_ay.get()
        yil = v_yil.get()
        return ((AY_ADLARI[ay - 1] + " ") if ay else "") + (yil if yil != "Tümü" else "Tüm Yıllar")

    def csv_btn():
        satirlar = durum.get("gorunen", [])
        if not satirlar:
            messagebox.showinfo(UYGULAMA_ADI, "Aktarılacak kayıt yok.", parent=kok)
            return
        donem = donem_adi()
        dosya_donem = donem.replace(" ", "_")
        tipler = [("Excel çalışma kitabı", "*.xlsx"), ("CSV (Excel'de açılır)", "*.csv")] if openpyxl else \
                 [("CSV (Excel'de açılır)", "*.csv")]
        yol = filedialog.asksaveasfilename(parent=kok, title="Excel'e aktar (%s)" % donem,
                                           defaultextension=tipler[0][1], filetypes=tipler,
                                           initialfile="Faturalar_%s" % dosya_donem)
        if not yol:
            return
        try:
            if yol.lower().endswith(".csv") or openpyxl is None:
                csv_yaz(yol, satirlar)
            else:
                xlsx_yaz(yol, satirlar, "Alış Faturaları - %s" % donem)
            if messagebox.askyesno(UYGULAMA_ADI, "%d kayıt aktarıldı:\n%s\n\nDosya şimdi açılsın mı?" % (len(satirlar), yol),
                                   parent=kok):
                pdf_ac(yol)
        except PermissionError:
            messagebox.showerror(UYGULAMA_ADI, "Dosya yazılamadı. Excel'de açık olabilir; kapatıp tekrar deneyin.", parent=kok)
        except OSError as e:
            hata_goster(e)
        except Exception as e:
            hata_kaydet(traceback.format_exc())
            messagebox.showerror(UYGULAMA_ADI, "Excel/CSV oluşturulamadı:\n%s: %s\n\nAyrıntı: %s"
                                 % (type(e).__name__, str(e)[:300], hata_gunlugu_yolu()), parent=kok)

    def satinalma_btn():
        satirlar = durum.get("gorunen", [])
        if not satirlar:
            messagebox.showinfo(UYGULAMA_ADI, "Aktarılacak kayıt yok.", parent=kok)
            return
        if openpyxl is None:
            messagebox.showerror(UYGULAMA_ADI, "Satınalma Excel'i için 'openpyxl' kütüphanesi gerekli.", parent=kok)
            return
        ay, yil = v_ay.get(), v_yil.get()
        baslik = " ".join(x for x in ((yil if yil != "Tümü" else ""), tr_buyuk(AY_ADLARI[ay - 1]) if ay else "") if x)
        baslik = (baslik + " " if baslik else "") + "HARCAMA LİSTESİ"
        yol = filedialog.asksaveasfilename(
            parent=kok, title="Satınalmaya aktar (%s)" % donem_adi(), defaultextension=".xlsx",
            filetypes=[("Excel çalışma kitabı", "*.xlsx")],
            initialfile="Makine İkmal Harcama Listesi - %s" % (AY_ADLARI[ay - 1] if ay else donem_adi()))
        if not yol:
            return
        try:
            satinalma_xlsx_yaz(yol, satirlar, baslik, yil if yil != "Tümü" else "Harcama Listesi")
            if messagebox.askyesno(UYGULAMA_ADI, "%d kayıt satınalma listesine aktarıldı:\n%s\n\nDosya şimdi açılsın mı?"
                                   % (len(satirlar), yol), parent=kok):
                pdf_ac(yol)
        except PermissionError:
            messagebox.showerror(UYGULAMA_ADI, "Dosya yazılamadı. Excel'de açık olabilir; kapatıp tekrar deneyin.", parent=kok)
        except (OSError, FileNotFoundError) as e:
            hata_goster(e)
        except Exception as e:
            hata_kaydet(traceback.format_exc())
            messagebox.showerror(UYGULAMA_ADI, "Satınalma Excel'i oluşturulamadı:\n%s: %s\n\nAyrıntı: %s"
                                 % (type(e).__name__, str(e)[:300], hata_gunlugu_yolu()), parent=kok)

    # ---- hesap menüsü eylemleri ------------------------------------------------
    def sifremi_degistir():
        if yeni_sifre_iste({"id": oturum["id"]}, False, eski_iste=True, ana=kok):
            try:
                if oturum["token"]:
                    oturum["token"] = hatirlamayi_yenile(oturum["id"])
            except VeritabaniHatasi:
                oturum["token"] = ""
            messagebox.showinfo(UYGULAMA_ADI, "Şifreniz değiştirildi.", parent=kok)

    def oturumu_kapat():
        if not messagebox.askyesno(UYGULAMA_ADI, "Oturum kapatılsın mı? Giriş ekranına dönülür.", parent=kok):
            return
        try:
            if oturum["token"]:
                db.oturum_kapat(oturum["token"])
        except VeritabaniHatasi:
            pass
        oturum_yaz(db.yol, "")                  # bu bilgisayardaki 'beni hatırla' kaydı silinir
        oturum["yeniden"] = True
        kok.destroy()

    def kullanicilar_penceresi():
        pen = yeni_pencere("Kullanıcılar")
        pen.geometry("620x380")
        pen.grab_set()
        agac = ttk.Treeview(pen, columns=("ad", "rol", "durum"), show="headings", selectmode="browse")
        for ad, baslik, gen in (("ad", "Kullanıcı Adı", 250), ("rol", "Yetki", 110), ("durum", "Durum", 190)):
            agac.heading(ad, text=baslik)
            agac.column(ad, width=gen, anchor="w")
        agac.pack(fill="both", expand=True, padx=10, pady=(10, 4))
        kayitlar = {}

        def doldur(secilecek=None):
            agac.delete(*agac.get_children())
            kayitlar.clear()
            for u in db.kullanicilar():
                kayitlar[str(u["id"])] = u
                durum = "Pasif" if not u["aktif"] else ("Şifre değiştirmeli" if u["sifre_degismeli"] else "Aktif")
                agac.insert("", "end", iid=str(u["id"]),
                            values=(u["kullanici_adi"], "Yönetici" if u["yonetici"] else "Kullanıcı", durum))
            if secilecek and agac.exists(secilecek):
                agac.selection_set(secilecek)

        def secili(baskasi=True):
            s = agac.selection()
            if not s:
                messagebox.showinfo(UYGULAMA_ADI, "Önce listeden bir kullanıcı seçin.", parent=pen)
                return None
            if baskasi and int(s[0]) == oturum["id"]:
                messagebox.showinfo(UYGULAMA_ADI, "Bu işlem kendi hesabınız için yapılamaz.\n"
                                    "(Şifrenizi Dosya > Şifremi Değiştir'den değiştirebilirsiniz.)", parent=pen)
                return None
            return kayitlar[s[0]]

        def dene(fonk):
            try:
                fonk()
            except VeritabaniHatasi as e:
                messagebox.showerror(UYGULAMA_ADI, str(e), parent=pen)
            doldur()

        def yeni():
            def kontrol(d, yonetici):
                db.kullanici_ekle(d[0], d[1], yonetici=yonetici, sifre_degismeli=True)
                return None
            sonuc = diyalog("Yeni Kullanıcı", "Geçici şifreyi kullanıcıya iletin; ilk girişte kendi şifresini belirleyecek.",
                            [("Kullanıcı adı", False, ""), ("Geçici şifre", False, gecici_sifre())], "Ekle",
                            onay="Yönetici olsun", kontrol=kontrol, ana=pen)
            if sonuc:
                messagebox.showinfo(UYGULAMA_ADI, "Hesap oluşturuldu.\n\nKullanıcı adı: %s\nGeçici şifre: %s\n\n"
                                    "Bu bilgileri kullanıcıya iletin." % (sonuc[0][0].strip(), sonuc[0][1]), parent=pen)
            doldur()

        def sifirla():
            u = secili()
            if not u:
                return

            def kontrol(d, _onay):
                db.sifre_sifirla(u["id"], d[0])
                return None
            sonuc = diyalog("Şifre Sıfırla", "'%s' için yeni geçici şifre. Kullanıcı ilk girişte değiştirecek."
                            % u["kullanici_adi"], [("Geçici şifre", False, gecici_sifre())], "Sıfırla",
                            kontrol=kontrol, ana=pen)
            if sonuc:
                messagebox.showinfo(UYGULAMA_ADI, "Şifre sıfırlandı.\n\nKullanıcı adı: %s\nGeçici şifre: %s"
                                    % (u["kullanici_adi"], sonuc[0][0]), parent=pen)
            doldur(str(u["id"]))

        def yetki():
            u = secili()
            if u:
                dene(lambda: db.kullanici_durumu(u["id"], yonetici=not u["yonetici"]))

        def aktif_pasif():
            u = secili()
            if not u:
                return
            if u["aktif"] and not messagebox.askyesno(
                    UYGULAMA_ADI, "'%s' pasife alınsın mı? Giriş yapamaz, kayıtları durur." % u["kullanici_adi"],
                    parent=pen):
                return
            dene(lambda: db.kullanici_durumu(u["id"], aktif=not u["aktif"]))

        alt_k = ttk.Frame(pen, padding=(10, 4, 10, 10))
        alt_k.pack(fill="x")
        for metin, komut in (("Yeni Kullanıcı...", yeni), ("Şifre Sıfırla...", sifirla),
                             ("Yönetici Yap / Al", yetki), ("Pasife Al / Aktif Et", aktif_pasif),
                             ("Kapat", pen.destroy)):
            ttk.Button(alt_k, text=metin, command=komut).pack(side="left", padx=3)
        pen.bind("<Escape>", lambda e: pen.destroy())
        dene(lambda: None)

    # ---- menü ----------------------------------------------------------------
    menu = tk.Menu(kok)
    m_dosya = tk.Menu(menu, tearoff=0)
    m_dosya.add_command(label="Excel'e Aktar...", command=csv_btn)
    m_dosya.add_command(label="Aylık Özet...", command=ozet_penceresi)
    m_dosya.add_command(label="PDF Okuma Testi...", command=okuma_testi)
    m_dosya.add_separator()
    m_dosya.add_command(label="Şifremi Değiştir...", command=sifremi_degistir)
    if oturum["yonetici"]:
        m_dosya.add_command(label="Kullanıcılar...", command=kullanicilar_penceresi)
    m_dosya.add_command(label="Oturumu Kapat", command=oturumu_kapat)
    m_dosya.add_separator()
    m_dosya.add_command(label="Çıkış", command=kok.destroy)
    m_gor = tk.Menu(menu, tearoff=0)
    m_tema = tk.Menu(m_gor, tearoff=0)
    for ad in TEMALAR:
        m_tema.add_radiobutton(label=ad, variable=v_tema, value=ad, command=lambda: (tema_uygula(v_tema.get(), True), yenile()))
    m_gor.add_cascade(label="Tema", menu=m_tema)
    menu.add_cascade(label="Dosya", menu=m_dosya)
    menu.add_cascade(label="Görünüm", menu=m_gor)
    m_yardim = tk.Menu(menu, tearoff=0)
    m_yardim.add_command(label="Hakkında...", command=lambda: messagebox.showinfo(
        "Hakkında", "%s\n\nDeveloped by %s" % (UYGULAMA_ADI, GELISTIRICI), parent=kok))
    menu.add_cascade(label="Yardım", menu=m_yardim)
    kok.configure(menu=menu)

    # ---- bağlantılar ---------------------------------------------------------
    btn_ekle.configure(command=lambda: form_ac())
    btn_duzenle.configure(command=duzenle)
    btn_sil.configure(command=sil)
    btn_klasor.configure(command=klasor_btn)
    btn_evrak.configure(command=lambda: evrak_penceresi())
    btn_ozet.configure(command=ozet_penceresi)
    btn_csv.configure(command=csv_btn)
    btn_satinalma.configure(command=satinalma_btn)
    btn_temizle.configure(command=temizle)
    btn_bul.configure(command=fatura_bul_penceresi)
    kok.bind("<Control-f>", fatura_bul_penceresi)
    ipucu_bagla(tree)
    tree.bind("<Double-1>", lambda e: evrak_penceresi())
    tree.bind("<Return>", duzenle)
    tree.bind("<Delete>", lambda e: sil())
    v_arama.trace_add("write", yenile)
    v_bas.trace_add("write", yenile)
    v_bit.trace_add("write", yenile)
    v_ay.trace_add("write", yenile)
    cb_firma.bind("<<ComboboxSelected>>", yenile)
    cb_yil.bind("<<ComboboxSelected>>", yenile)

    # açılışta: verisi olan en yeni yılı göster
    try:
        ilk = db.hepsi()
        v_yil.set(str(max(int(r["fatura_tarihi"][:4]) for r in ilk)) if ilk else str(date.today().year))
    except (VeritabaniHatasi, ValueError):
        v_yil.set(str(date.today().year))
    kok.title("%s   -   %s%s" % (UYGULAMA_ADI, kullanici, "  (Yönetici)" if oturum["yonetici"] else ""))
    yenile()

    def modal_one_getir(e=None):
        """Ana pencere yeniden etkin olunca (alt-tab, görev çubuğu, küçültüp açma) odağı tutan modal pencere
        (form, evrak, kullanıcılar...) arkada kalmışsa öne getirir; yoksa uygulama tıklanamaz/kilitli görünür."""
        if e is not None and e.widget is not kok:
            return
        try:
            g = kok.grab_current()
            if g is not None and g is not kok and g.winfo_exists():
                g.deiconify()
                g.lift()
                g.focus_set()
        except tk.TclError:
            pass
    kok.bind("<FocusIn>", modal_one_getir, add="+")
    kok.bind("<Map>", modal_one_getir, add="+")

    kok.deiconify()
    kok.lift()
    kok.focus_force()                  # giriş penceresinden sonra ana pencere ön planda ve odakta açılsın
    e_ara.focus_set()
    kok.mainloop()
    return oturum["yeniden"]          # 'Oturumu Kapat' seçildiyse giriş ekranıyla yeniden başlar


def yonetici_kurtarma_komutu():
    """python fatura_takip.py --yonetici-sifirla : yönetici şifresi unutulduysa komut satırından yenisini belirler
    (hesap yoksa yönetici olarak açar). Veritabanı dosyasına erişimi olan kişi yapabilir."""
    ayar = ayarlari_oku()
    try:
        db = FaturaDB(ayar["veritabani"])
        print("Veritabanı:", ayar["veritabani"])
        ad = input("Yönetici kullanıcı adı (yoksa yeni açılır): ").strip()
        sifre = getpass.getpass("Yeni şifre: ")
        if sifre != getpass.getpass("Yeni şifre (tekrar): "):
            print("Şifreler aynı değil, işlem iptal edildi.")
            return
        db.yonetici_kurtar(ad, sifre)
        print("Tamam: '%s' yönetici olarak ayarlandı, şifresi yenilendi." % ad)
    except VeritabaniHatasi as e:
        print("Hata:", e)


def kendini_test(rapor_yolu):
    """Paketlenmiş exe'nin (ya da kaynak kodun) kendi kendini sınaması: kaynaklar, kütüphaneler, pencere açma,
    veritabanı, tüm dışa aktarmalar ve PDF okuma geçici bir klasörde denenir; gerçek veriye dokunmaz.
        FaturaTakip.exe --kendini-test rapor.txt        (sonuç rapor.txt'ye yazılır; çıkış kodu 0 = hepsi tamam)"""
    satirlar, hatalar = [], []
    gecici = tempfile.mkdtemp(prefix="fatura_kendini_test_")

    def adim(ad, fonk):
        try:
            sonuc = fonk()
            satirlar.append("OK    %s%s" % (ad, (" -> %s" % sonuc) if sonuc else ""))
        except Exception:
            hatalar.append(ad)
            satirlar.append("HATA  %s\n%s" % (ad, traceback.format_exc()))

    def ortam():
        return "frozen=%s | python=%s | openpyxl=%s | pymupdf=%s | sv_ttk=%s" % (
            bool(getattr(sys, "frozen", False)), sys.version.split()[0],
            getattr(openpyxl, "__version__", None), getattr(fitz, "__doc__", "yok").split()[1] if fitz else None,
            "var" if sv_ttk else None)

    def kaynaklar():
        for ad, yol in (("ikon", IKON_YOLU), ("satınalma şablonu", SATINALMA_SABLONU)):
            if not os.path.exists(yol):
                raise FileNotFoundError("%s bulunamadı: %s" % (ad, yol))

    def pencere():
        import tkinter as tk
        kok = tk.Tk()
        try:
            kok.withdraw()
            if sv_ttk:
                sv_ttk.set_theme("light")
            return "Tk %s" % tk.TkVersion
        finally:
            kok.destroy()

    kayitlar = []

    def veritabani():
        db = FaturaDB(os.path.join(gecici, "t.db"))
        db.ilk_yonetici_olustur("deneme", "deneme-1234")
        kirli = chr(11) + "=1+1"
        for i in range(3):
            db.ekle(dict(fatura_tarihi="2026-10-0%d" % (i + 1), fatura_no="T-%d" % i, firma_adi="Firma" + kirli if i == 0 else "Firma %d" % i,
                         yapilan_is="İş" + chr(12) + " adı", tutar_haric=100000 + i, pdf_yolu="", kategori="Genel Alım",
                         talep_no="40355%02d" % i, talep_tarihi="2026-09-29", dt_no="26DT500000%d" % i,
                         talep_eden_birim="Deneme Birimi"), "deneme")
        kayitlar.extend(db.hepsi())
        return "%d kayıt, hesap girişi %s" % (len(kayitlar), "tamam" if db.giris_dene("DENEME", "deneme-1234")[0] else "BAŞARISIZ")

    def disa_aktarma():
        if openpyxl is None:
            raise RuntimeError("openpyxl yok")
        a, b, c = (os.path.join(gecici, x) for x in ("a.xlsx", "s.xlsx", "c.csv"))
        xlsx_yaz(a, kayitlar, "Deneme")
        satinalma_xlsx_yaz(b, kayitlar, "2026 EKİM HARCAMA LİSTESİ", "2026")
        csv_yaz(c, kayitlar)
        ws = openpyxl.load_workbook(b).active
        if not ws["H3"].value or ws["D3"].value != ws["F3"].value:
            raise ValueError("satınalma çıktısı beklenen biçimde değil")
        for x in (a, b):
            for satir in openpyxl.load_workbook(x).active.iter_rows(max_row=400):
                for h in satir:
                    if h.data_type == "f" and not str(h.value).startswith(("=SUBTOTAL(", "=SUM(")):
                        raise ValueError("formül enjeksiyonu: %s %r" % (h.coordinate, h.value))
        return "Excel, Satınalma, CSV"

    def pdf_okuma():
        pdf = os.path.join(gecici, "d.pdf")
        icerik = b"BT /F1 12 Tf 72 720 Td (Fatura Tarihi: 04.06.2026) Tj ET"
        satirlar_pdf = [b"%PDF-1.4",
                        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj",
                        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj",
                        b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
                        b"/Resources << /Font << /F1 5 0 R >> >> >> endobj",
                        b"4 0 obj << /Length " + str(len(icerik)).encode() + b" >>",
                        b"stream",
                        icerik,
                        b"endstream endobj",
                        b"5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj",
                        b"trailer << /Root 1 0 R >>",
                        b"%EOF"]
        govde = b"\n".join(satirlar_pdf) + b"\n"
        with open(pdf, "wb") as fh:
            fh.write(govde)
        metin = " ".join(pdf_sayfa_metinleri(pdf)[0])
        mupdf = "PyMuPDF" if fitz else "yerleşik okuyucu"
        if "04.06.2026" not in metin:
            raise ValueError("PDF metni okunamadı: %r" % metin[:80])
        pdf_bilgi_cikar(pdf, {"kurum_anahtar": "", "kurum_vkn": ""})
        return mupdf

    for ad, fonk in (("ortam", ortam), ("kaynak dosyaları (ikon, şablon)", kaynaklar), ("pencere açma ve tema", pencere),
                     ("veritabanı, hesap, kirli veri temizliği", veritabani), ("dışa aktarmalar", disa_aktarma),
                     ("PDF okuma", pdf_okuma)):
        adim(ad, fonk)
    satirlar.append("")
    satirlar.append("SONUÇ: %s" % ("HEPSİ TAMAM" if not hatalar else "BAŞARISIZ: " + ", ".join(hatalar)))
    try:
        with open(rapor_yolu, "w", encoding="utf-8") as fh:
            fh.write("\n".join(satirlar) + "\n")
    finally:
        shutil.rmtree(gecici, ignore_errors=True)
    sys.exit(1 if hatalar else 0)


if __name__ == "__main__":
    if "--kendini-test" in sys.argv:
        i = sys.argv.index("--kendini-test")
        kendini_test(sys.argv[i + 1] if i + 1 < len(sys.argv) else os.path.join(tempfile.gettempdir(), "fatura_kendini_test.txt"))
    elif "--yonetici-sifirla" in sys.argv:
        yonetici_kurtarma_komutu()
    else:
        while arayuzu_calistir():
            pass
